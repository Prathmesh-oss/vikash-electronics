import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_borders(cell, top="000000", bottom="000000", left="000000", right="000000", sz="12"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, color in borders.items():
        if color:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), sz)
            b.set(qn('w:space'), '0')
            b.set(qn('w:color'), color)
            tcBorders.append(b)
        else:
            b = OxmlElement(f'w:{border_name}')
            b.set(qn('w:val'), 'none')
            tcBorders.append(b)
    tcPr.append(tcBorders)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_form():
    doc = docx.Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Title: Project – 1 Application Form (Underlined, Centered)
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(16)
    r_title = p_title.add_run("Project – 1 Application Form")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.underline = True

    # Helper function to add a boxed table
    def add_box(content_list, height_pt=None):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.9)
        set_cell_borders(cell, sz="12")
        set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        
        for idx, item in enumerate(content_list):
            if idx > 0:
                p = cell.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(2)
                
            label, val, bold_val = item
            r_label = p.add_run(label)
            r_label.font.name = "Times New Roman"
            r_label.font.size = Pt(11.5)
            r_label.font.bold = True
            
            if val:
                r_val = p.add_run(f" {val}")
                r_val.font.name = "Times New Roman"
                r_val.font.size = Pt(11.5)
                r_val.font.bold = bold_val
                
        if height_pt:
            p_extra = cell.add_paragraph()
            p_extra.paragraph_format.space_before = Pt(height_pt)
            p_extra.paragraph_format.space_after = Pt(0)
            
        p_spacer = doc.add_paragraph()
        p_spacer.paragraph_format.space_before = Pt(0)
        p_spacer.paragraph_format.space_after = Pt(10)

    # 1. Student Division/Enrollment No:
    add_box([("Student Division/Enrollment No:", "202503103510061", False)])
    
    # 2. Student Full Name:
    add_box([("Student Full Name:", "Prathmesh Chaudhari", False)])
    
    # 3. Is live project? & Name of the client:
    add_box([
        ("Is live project?", "Yes (Industry Defined Project - IDP)", False),
        ("Name of the client:", "Vikash Electronics (Proprietor: Mr. Murlidhar Chaudhari)", False)
    ])
    
    # 4. Proposed Project Title:
    add_box([("Proposed Project Title:", "Vikash Electronics – Web-Based DJ & Electronic Equipment Repair Management and Showcase Portal", False)])
    
    # 5. Project Description (50-70 Words):
    desc = ("A specialized web application developed for Vikash Electronics, a pro-audio, DJ sound system, and stage lighting "
            "repair workshop in Surat. The portal enables digital equipment service booking, showcases repair capabilities across "
            "audio, lighting, and visual gear, offers direct WhatsApp fault consultation, and integrates geolocation routing. "
            "It eliminates third-party listing aggregator commissions and establishes direct digital presence for the workshop.")
    add_box([("Project Description (50-70 Words):\n", desc, False)], height_pt=80)
    
    # 6. Student Signature: (Blank)
    add_box([("Student Signature:", "", False)], height_pt=24)
    
    # 7. Verified by: (Blank)
    add_box([("Verified by:", "", False)], height_pt=50)
    
    output_path = r"d:\A_S_Projects\Vikas_Electronics\Project_1_Application_Form_Filled.docx"
    doc.save(output_path)
    print("Saved successfully to:", output_path)

if __name__ == "__main__":
    create_form()
