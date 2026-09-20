import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# --- COLOR PALETTE (Matched to Vikash Electronics Website) ---
BG_DARK = RGBColor(11, 17, 32)       # #0B1120 Deep Navy Slate
CARD_BG = RGBColor(19, 29, 49)       # #131D31 Card Navy
CARD_BORDER = RGBColor(30, 41, 59)   # #1E293B Subtle Border
GOLD = RGBColor(245, 158, 11)        # #F59E0B Vibrant Amber / Gold
CYAN = RGBColor(56, 189, 248)        # #38BDF8 Sky Cyan Accent
WHITE = RGBColor(255, 255, 255)      # Pure White
GRAY_LIGHT = RGBColor(226, 232, 240) # Slate 200
GRAY_MUTED = RGBColor(148, 163, 184) # Slate 400
GREEN = RGBColor(52, 211, 153)       # Emerald 400

BASE_DIR = r"d:\A_S_Projects\Vikas_Electronics"
OUTPUT_FILE = os.path.join(BASE_DIR, "Vikash_Electronics_Presentation.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def apply_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK

def add_header(slide, tag_text, title_white, title_gold=""):
    # Category Tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.4))
    tf = tag_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = tag_text.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p.font.name = "Segoe UI"

    # Slide Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    
    r1 = p_title.add_run()
    r1.text = title_white + " "
    r1.font.size = Pt(26)
    r1.font.bold = True
    r1.font.color.rgb = WHITE
    r1.font.name = "Segoe UI"

    if title_gold:
        r2 = p_title.add_run()
        r2.text = title_gold
        r2.font.size = Pt(26)
        r2.font.bold = True
        r2.font.color.rgb = GOLD
        r2.font.name = "Segoe UI"

def create_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.2)
    return shape

# ==============================================================================
# SLIDE 1: TITLE / HERO SLIDE
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
apply_background(slide1)

# Subtle top status badge
status_shape = create_card(slide1, Inches(0.8), Inches(0.7), Inches(4.5), Inches(0.4), RGBColor(15, 23, 42), CYAN)
tb = slide1.shapes.add_textbox(Inches(0.9), Inches(0.7), Inches(4.3), Inches(0.4))
p = tb.text_frame.paragraphs[0]
p.text = "⚡ SURAT'S PREMIER AUDIO & DJ REPAIR CENTER"
p.font.size = Pt(10)
p.font.bold = True
p.font.color.rgb = CYAN
p.font.name = "Segoe UI"

# Main Brand Heading
title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(7.5), Inches(2.2))
tf = title_box.text_frame
tf.word_wrap = True
p1 = tf.paragraphs[0]
r1 = p1.add_run()
r1.text = "VIKASH "
r1.font.size = Pt(44)
r1.font.bold = True
r1.font.color.rgb = WHITE
r1.font.name = "Segoe UI"

r2 = p1.add_run()
r2.text = "ELECTRONICS"
r2.font.size = Pt(44)
r2.font.bold = True
r2.font.color.rgb = GOLD
r2.font.name = "Segoe UI"

p2 = tf.add_paragraph()
p2.text = "Professional DJ & Electronic Equipment Repairing Services"
p2.font.size = Pt(20)
p2.font.bold = True
p2.font.color.rgb = GRAY_LIGHT
p2.font.name = "Segoe UI"
p2.space_before = Pt(8)

p3 = tf.add_paragraph()
p3.text = "Expert diagnosis, maintenance, and precision repair for DJ sound systems, Sharpy moving heads, stage lights, power amplifiers, and electronic equipment in Surat."
p3.font.size = Pt(13)
p3.font.color.rgb = GRAY_MUTED
p3.font.name = "Segoe UI"
p3.space_before = Pt(12)

# Quick Highlights Cards at bottom left
highlights = [
    ("Specialist", "Murlidhar Chaudhari"),
    ("Contact", "+91 98254 85520"),
    ("Location", "Limbayat, Surat, Gujarat")
]
for idx, (label, val) in enumerate(highlights):
    col_left = Inches(0.8 + idx * 2.45)
    create_card(slide1, col_left, Inches(5.6), Inches(2.3), Inches(1.2))
    hb = slide1.shapes.add_textbox(col_left + Inches(0.15), Inches(5.7), Inches(2.0), Inches(1.0))
    tf = hb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = label.upper()
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p.font.name = "Segoe UI"
    p2 = tf.add_paragraph()
    p2.text = val
    p2.font.size = Pt(12)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.font.name = "Segoe UI"
    p2.space_before = Pt(4)

# Right Side: Workshop & Logo Feature Box
logo_path = os.path.join(BASE_DIR, "public", "images", "logo", "vikash-logo.jpg")
workshop_path = os.path.join(BASE_DIR, "public", "images", "about", "about_workshop.jpg")

create_card(slide1, Inches(8.3), Inches(0.9), Inches(4.2), Inches(5.9), CARD_BG, CYAN)

if os.path.exists(workshop_path):
    slide1.shapes.add_picture(workshop_path, Inches(8.5), Inches(1.1), width=Inches(3.8))

# Overlay badge on right card
tb_overlay = slide1.shapes.add_textbox(Inches(8.5), Inches(4.2), Inches(3.8), Inches(2.3))
tf = tb_overlay.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "✓ COMPONENT-LEVEL DIAGNOSTICS"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = GREEN
p.font.name = "Segoe UI"

p2 = tf.add_paragraph()
p2.text = "Precision bench testing for event equipment, stage monitors, moving heads, and heavy PA amplifiers."
p2.font.size = Pt(11)
p2.font.color.rgb = GRAY_LIGHT
p2.font.name = "Segoe UI"
p2.space_before = Pt(6)

p3 = tf.add_paragraph()
p3.text = "★ Quality Service & Trusted Repair in Surat"
p3.font.size = Pt(12)
p3.font.bold = True
p3.font.color.rgb = GOLD
p3.font.name = "Segoe UI"
p3.space_before = Pt(10)


# ==============================================================================
# SLIDE 2: ABOUT VIKASH ELECTRONICS & OUR TECHNICIAN
# ==============================================================================
slide2 = prs.slides.add_slide(blank_layout)
apply_background(slide2)
add_header(slide2, "About Us", "WHO WE ARE &", "OUR EXPERTISE")

# Left Column: Workshop Bench Image with frame
create_card(slide2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.1), CARD_BG, CYAN)
if os.path.exists(workshop_path):
    slide2.shapes.add_picture(workshop_path, Inches(1.0), Inches(2.0), width=Inches(5.2))

# Right Column: Narrative + Technician Spotlight + Specialization Pills
right_box = slide2.shapes.add_textbox(Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.0))
tf = right_box.text_frame
tf.word_wrap = True
p1 = tf.paragraphs[0]
p1.text = "Trusted Name in DJ & Electronic Equipment Repair"
p1.font.size = Pt(18)
p1.font.bold = True
p1.font.color.rgb = WHITE
p1.font.name = "Segoe UI"

p2 = tf.add_paragraph()
p2.text = "With years of dedicated hands-on experience and expert technical know-how, Vikash Electronics provides reliable, prompt, and precision repair solutions for all types of audio, stage lighting, and visual entertainment electronics."
p2.font.size = Pt(12)
p2.font.color.rgb = GRAY_LIGHT
p2.font.name = "Segoe UI"
p2.space_before = Pt(8)

# Technician Attribution Card
create_card(slide2, Inches(6.8), Inches(3.6), Inches(5.7), Inches(1.4), RGBColor(19, 27, 46), GOLD)
tb_tech = slide2.shapes.add_textbox(Inches(7.0), Inches(3.7), Inches(5.3), Inches(1.2))
tf = tb_tech.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CHIEF REPAIR SPECIALIST"
p.font.size = Pt(10)
p.font.bold = True
p.font.color.rgb = GOLD
p.font.name = "Segoe UI"

p2 = tf.add_paragraph()
p2.text = "Murlidhar Chaudhari"
p2.font.size = Pt(17)
p2.font.bold = True
p2.font.color.rgb = WHITE
p2.font.name = "Segoe UI"

p3 = tf.add_paragraph()
p3.text = "Specializing in circuit-level diagnostics, voice coils, moving heads & high-watt PA systems."
p3.font.size = Pt(11)
p3.font.color.rgb = GRAY_MUTED
p3.font.name = "Segoe UI"

# Specialization Badges Grid
tb_spec_title = slide2.shapes.add_textbox(Inches(6.8), Inches(5.1), Inches(5.7), Inches(0.4))
p = tb_spec_title.text_frame.paragraphs[0]
p.text = "CORE REPAIR SPECTRUM:"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = CYAN
p.font.name = "Segoe UI"

specs = [
    "DJ Sound Systems", "Sharpy / Moving Heads", "Pro Audio Speakers", "High-Power Amplifiers",
    "Stage & DJ Lights", "Audio Mixers & Consoles", "LED / Smart TVs", "Home Theatres & Decoders"
]
for i, spec in enumerate(specs):
    row = i // 2
    col = i % 2
    sx = Inches(6.8 + col * 2.9)
    sy = Inches(5.6 + row * 0.6)
    create_card(slide2, sx, sy, Inches(2.75), Inches(0.45), CARD_BG, CARD_BORDER)
    tb = slide2.shapes.add_textbox(sx + Inches(0.1), sy + Inches(0.05), Inches(2.55), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = f"✔  {spec}"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = GRAY_LIGHT
    p.font.name = "Segoe UI"


# ==============================================================================
# SLIDE 3: COMPLETE REPAIR SERVICES (8 SERVICES GRID)
# ==============================================================================
slide3 = prs.slides.add_slide(blank_layout)
apply_background(slide3)
add_header(slide3, "Services Overview", "COMPREHENSIVE", "REPAIR SOLUTIONS")

services_list = [
    ("Speaker Repairing", "Voice coil rewinding, cone & surround replacement, subwoofers & line arrays.", "speaker.jpg", GOLD),
    ("Amplifier Repairing", "Power transistor fix, toroidal transformer checks, distortion & protection mode.", "amplifier.jpg", CYAN),
    ("DJ Light Repairing", "LED par cans, laser projectors, strobes, driver boards & DMX controllers.", "dj_light.jpg", GOLD),
    ("Sharpy / Moving Head", "Pan/tilt motor calibration, gobo & color wheels, ballast & beam lamp change.", "sharpy.jpg", CYAN),
    ("Sound Mixer Repairing", "Fader cleaning, potentiometer scratchiness, preamp circuits & multi-channel boards.", "sound_mixer.jpg", GOLD),
    ("Home Theatre Repairing", "AV receivers, 5.1/7.1 channel diagnosis, plate amplifiers, audio calibration.", "home_theatre.jpg", CYAN),
    ("Smart TV Repairing", "Component-level repair for LED/Smart TVs, backlight replacement, power boards.", "tv.jpg", GOLD),
    ("Other DJ Electronics", "Crossovers, equalizers, smoke/fog machines, power conditioners & custom gear.", "other_electronic.jpg", CYAN),
]

for idx, (title, desc, img_name, accent) in enumerate(services_list):
    col = idx % 4
    row = idx // 4
    left = Inches(0.8 + col * 2.95)
    top = Inches(1.8 + row * 2.6)
    width = Inches(2.8)
    height = Inches(2.4)
    
    create_card(slide3, left, top, width, height, CARD_BG, accent)
    
    img_path = os.path.join(BASE_DIR, "public", "images", "services", img_name)
    if os.path.exists(img_path):
        slide3.shapes.add_picture(img_path, left + Inches(0.15), top + Inches(0.15), width=Inches(0.9), height=Inches(0.7))
        
    tb = slide3.shapes.add_textbox(left + Inches(1.15), top + Inches(0.1), Inches(1.5), Inches(0.8))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Segoe UI"
    
    tb_desc = slide3.shapes.add_textbox(left + Inches(0.15), top + Inches(0.95), Inches(2.5), Inches(1.35))
    tf_desc = tb_desc.text_frame
    tf_desc.word_wrap = True
    p_desc = tf_desc.paragraphs[0]
    p_desc.text = desc
    p_desc.font.size = Pt(9.5)
    p_desc.font.color.rgb = GRAY_MUTED
    p_desc.font.name = "Segoe UI"


# ==============================================================================
# SLIDE 4: PRO AUDIO REPAIR FOCUS (SPEAKERS, AMPS, MIXERS)
# ==============================================================================
slide4 = prs.slides.add_slide(blank_layout)
apply_background(slide4)
add_header(slide4, "Audio Systems", "PROFESSIONAL AUDIO &", "SOUND GEAR REPAIR")

audio_cards = [
    ("Speaker & Subwoofer Repair", 
     [
         "Voice Coil Rewinding: High-temperature Kapton / CCAR wire rewinds.",
         "Cone & Surround Replacement: Re-coning 12\", 15\", 18\" high-excursion woofers.",
         "Spider Alignment: Centering and dampening for crystal clear bass response.",
         "HF Compression Drivers: Titanium diaphragm replacements for line arrays."
     ], "speaker.jpg"),
    ("Power Amplifiers (DJ & PA)",
     [
         "Output Transistors: Troubleshooting shorted 2SC5200 / 2SA1943 pairs.",
         "Toroidal Power Supply: Rectifier diode & filter capacitor bank fixes.",
         "Thermal & DC Protection: Clearing clip/protect lockups under heavy load.",
         "Channel Balancing: Eliminating hum, cross-talk, and unwanted distortion."
     ], "amplifier.jpg"),
    ("Audio Mixers & Home Theatres",
     [
         "Fader & Potentiometer: DeoxIT ultrasonic servicing and slide pot swap.",
         "Mic Preamps: Low noise op-amp replacement and phantom power 48V fix.",
         "Surround AV Receivers: 5.1/7.1 channel HDMI & optical DAC repair.",
         "Plate Amps: Active subwoofer power circuitry and crossover tuning."
     ], "sound_mixer.jpg"),
]

for idx, (title, points, img_name) in enumerate(audio_cards):
    left = Inches(0.8 + idx * 3.95)
    top = Inches(1.8)
    width = Inches(3.8)
    height = Inches(5.1)
    
    create_card(slide4, left, top, width, height, CARD_BG, GOLD)
    
    img_path = os.path.join(BASE_DIR, "public", "images", "services", img_name)
    if os.path.exists(img_path):
        slide4.shapes.add_picture(img_path, left + Inches(0.2), top + Inches(0.2), width=Inches(3.4), height=Inches(1.8))
        
    tb = slide4.shapes.add_textbox(left + Inches(0.2), top + Inches(2.1), Inches(3.4), Inches(0.6))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = GOLD
    p.font.name = "Segoe UI"
    
    tb_pts = slide4.shapes.add_textbox(left + Inches(0.2), top + Inches(2.7), Inches(3.4), Inches(2.3))
    tf_pts = tb_pts.text_frame
    tf_pts.word_wrap = True
    for p_idx, pt in enumerate(points):
        p = tf_pts.paragraphs[0] if p_idx == 0 else tf_pts.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = GRAY_LIGHT
        p.font.name = "Segoe UI"
        p.space_before = Pt(4)


# ==============================================================================
# SLIDE 5: STAGE LIGHTING & SHARPY MOVING HEADS
# ==============================================================================
slide5 = prs.slides.add_slide(blank_layout)
apply_background(slide5)
add_header(slide5, "Stage & Event Lighting", "SHARPY MOVING HEADS &", "DJ LIGHTS")

lighting_cards = [
    ("Sharpy / Beam Moving Head Lights", 
     [
         "Optical Path Alignment: Cleaning heat mirrors, lenses, and collimators.",
         "Stepper Motor Calibration: Smooth 540° pan and 270° tilt movement without slippage.",
         "Gobo & Color Wheel Repair: Sensor hall alignment, magnetic indexing, and color filters.",
         "Electronic Ballast & Igniter: High-voltage ignition troubleshooting and lamp cooling fans.",
         "Bulb Replacement: 7R, 10R, 15R high-output discharge lamps matched to ballast specs."
     ], "sharpy.jpg", CYAN),
    ("DJ Stage Lights & Effect Gear",
     [
         "DMX-512 Protocol: Fixing communication lags, signal termination, and addressing faults.",
         "RGBW LED Par Cans: Constant-current driver board repairs and LED array replacements.",
         "Laser Projectors: Galvanometer scanner alignment, laser diode cooling & TTL modulation.",
         "Strobe & Smoke Machines: High-power strobe capacitors and pump heater element repair.",
         "Event Ready Testing: 2-hour burn-in stress test before dispatch."
     ], "dj_light.jpg", GOLD),
]

for idx, (title, points, img_name, accent) in enumerate(lighting_cards):
    left = Inches(0.8 + idx * 5.95)
    top = Inches(1.8)
    width = Inches(5.8)
    height = Inches(5.1)
    
    create_card(slide5, left, top, width, height, CARD_BG, accent)
    
    img_path = os.path.join(BASE_DIR, "public", "images", "services", img_name)
    if os.path.exists(img_path):
        slide5.shapes.add_picture(img_path, left + Inches(0.25), top + Inches(0.25), width=Inches(5.3), height=Inches(2.0))
        
    tb = slide5.shapes.add_textbox(left + Inches(0.25), top + Inches(2.35), Inches(5.3), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = accent
    p.font.name = "Segoe UI"
    
    tb_pts = slide5.shapes.add_textbox(left + Inches(0.25), top + Inches(2.9), Inches(5.3), Inches(2.1))
    tf_pts = tb_pts.text_frame
    tf_pts.word_wrap = True
    for p_idx, pt in enumerate(points):
        p = tf_pts.paragraphs[0] if p_idx == 0 else tf_pts.add_paragraph()
        p.text = f"✔  {pt}"
        p.font.size = Pt(10)
        p.font.color.rgb = GRAY_LIGHT
        p.font.name = "Segoe UI"
        p.space_before = Pt(5)


# ==============================================================================
# SLIDE 6: WHY CHOOSE VIKASH ELECTRONICS (KEY ADVANTAGES)
# ==============================================================================
slide6 = prs.slides.add_slide(blank_layout)
apply_background(slide6)
add_header(slide6, "Why Choose Us", "PRECISION REPAIR &", "CUSTOMER ADVANTAGE")

usps = [
    ("Component-Level Repair", "Instead of recommending expensive whole-board replacements, we diagnose and replace individual faulty ICs, MOSFETs, capacitors, and diodes to save you money.", "💰", GOLD),
    ("Fast Turnaround for Events", "We understand that event managers, DJs, and sound engineers have tight timelines. We offer expedited emergency repairs to ensure zero show downtime.", "⏱️", CYAN),
    ("Expert Technician with Mastery", "Led by Murlidhar Chaudhari with years of specialized field experience across Indian and imported DJ gear, line arrays, and touring rigs.", "👨‍🔧", GOLD),
    ("Genuine High-Grade Spares", "We source genuine heat-tolerant transistors, original discharge lamps, and high-temp Kapton voice coils for lasting reliability.", "🛡️", CYAN),
    ("Rigorous Bench Load Testing", "Every serviced amplifier and light undergoes intensive thermal and load cycle testing on our Surat workshop bench prior to handover.", "🔬", GOLD),
    ("Transparent & Honest Pricing", "Detailed breakdown of diagnosis, required components, and labor fees before undertaking major repairs. No hidden surprise costs.", "🤝", CYAN),
]

for idx, (title, desc, icon, accent) in enumerate(usps):
    col = idx % 3
    row = idx // 3
    left = Inches(0.8 + col * 3.95)
    top = Inches(1.8 + row * 2.6)
    width = Inches(3.8)
    height = Inches(2.35)
    
    create_card(slide6, left, top, width, height, CARD_BG, accent)
    
    tb = slide6.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(3.4), Inches(0.5))
    p = tb.text_frame.paragraphs[0]
    p.text = f"{icon}  {title}"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Segoe UI"
    
    tb_desc = slide6.shapes.add_textbox(left + Inches(0.2), top + Inches(0.7), Inches(3.4), Inches(1.5))
    tf_desc = tb_desc.text_frame
    tf_desc.word_wrap = True
    p_desc = tf_desc.paragraphs[0]
    p_desc.text = desc
    p_desc.font.size = Pt(10)
    p_desc.font.color.rgb = GRAY_MUTED
    p_desc.font.name = "Segoe UI"


# ==============================================================================
# SLIDE 7: STEP-BY-STEP REPAIR WORKFLOW
# ==============================================================================
slide7 = prs.slides.add_slide(blank_layout)
apply_background(slide7)
add_header(slide7, "Quality Assurance", "OUR SYSTEMATIC", "REPAIR PROCESS")

steps = [
    ("Step 1", "Intake & Visual Inspection", "Thorough inspection for burnt traces, liquid damage, cracked voice coils, or physical impact."),
    ("Step 2", "Component Diagnostics", "Multimeter, oscilloscope, and signal generator analysis to isolate the exact micro-fault."),
    ("Step 3", "Precision Component Fix", "Professional desoldering, original replacement parts, heat-sink re-pasting, and optical alignment."),
    ("Step 4", "Load & Stress Burn-In", "Running gear at rated power output and DMX sequences to verify thermal stability."),
    ("Step 5", "Handover & Warranty", "Safe packaging, customer demonstration, maintenance guidance, and post-repair peace of mind.")
]

for idx, (step_num, step_title, step_desc) in enumerate(steps):
    left = Inches(0.8 + idx * 2.38)
    top = Inches(2.2)
    width = Inches(2.25)
    height = Inches(4.3)
    
    accent = GOLD if idx % 2 == 0 else CYAN
    create_card(slide7, left, top, width, height, CARD_BG, accent)
    
    # Step Badge
    badge = create_card(slide7, left + Inches(0.2), top + Inches(0.25), Inches(1.85), Inches(0.4), RGBColor(15, 23, 42), accent)
    tb_num = slide7.shapes.add_textbox(left + Inches(0.2), top + Inches(0.25), Inches(1.85), Inches(0.4))
    p = tb_num.text_frame.paragraphs[0]
    p.text = step_num.upper()
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = accent
    p.font.alignment = PP_ALIGN.CENTER
    p.font.name = "Segoe UI"
    
    tb_title = slide7.shapes.add_textbox(left + Inches(0.15), top + Inches(0.85), Inches(1.95), Inches(1.0))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p = tf_title.paragraphs[0]
    p.text = step_title
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = "Segoe UI"
    
    tb_desc = slide7.shapes.add_textbox(left + Inches(0.15), top + Inches(1.95), Inches(1.95), Inches(2.2))
    tf_desc = tb_desc.text_frame
    tf_desc.word_wrap = True
    p = tf_desc.paragraphs[0]
    p.text = step_desc
    p.font.size = Pt(10)
    p.font.color.rgb = GRAY_MUTED
    p.font.name = "Segoe UI"


# ==============================================================================
# SLIDE 8: WORKSHOP LOCATION & CONTACT DETAILS
# ==============================================================================
slide8 = prs.slides.add_slide(blank_layout)
apply_background(slide8)
add_header(slide8, "Get In Touch", "VISIT OUR WORKSHOP &", "CONTACT US")

# Left Column: Contact Cards
create_card(slide8, Inches(0.8), Inches(1.8), Inches(6.0), Inches(5.0), CARD_BG, GOLD)

tb_c = slide8.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.4), Inches(4.5))
tf = tb_c.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "VIKASH ELECTRONICS"
p.font.size = Pt(20)
p.font.bold = True
p.font.color.rgb = WHITE
p.font.name = "Segoe UI"

p2 = tf.add_paragraph()
p2.text = "Surat's Specialized DJ & Electronic Repair Center"
p2.font.size = Pt(12)
p2.font.color.rgb = GOLD
p2.font.name = "Segoe UI"
p2.space_before = Pt(4)

p3 = tf.add_paragraph()
p3.text = "📍 Workshop Address:"
p3.font.size = Pt(13)
p3.font.bold = True
p3.font.color.rgb = CYAN
p3.font.name = "Segoe UI"
p3.space_before = Pt(14)

p4 = tf.add_paragraph()
p4.text = "Plot No. 199, Sarve No. 1-2, Sanjay Nagar,\nNear Udhna Yard, Limbayat,\nSurat - 394210, Gujarat, India."
p4.font.size = Pt(11)
p4.font.color.rgb = GRAY_LIGHT
p4.font.name = "Segoe UI"
p4.space_before = Pt(4)

p5 = tf.add_paragraph()
p5.text = "📞 Direct Helpline / WhatsApp:"
p5.font.size = Pt(13)
p5.font.bold = True
p5.font.color.rgb = CYAN
p5.font.name = "Segoe UI"
p5.space_before = Pt(12)

p6 = tf.add_paragraph()
p6.text = "+91 98254 85520  |  Technician: Murlidhar Chaudhari"
p6.font.size = Pt(12)
p6.font.bold = True
p6.font.color.rgb = WHITE
p6.font.name = "Segoe UI"
p6.space_before = Pt(4)

p7 = tf.add_paragraph()
p7.text = "⏰ Working Hours:"
p7.font.size = Pt(13)
p7.font.bold = True
p7.font.color.rgb = CYAN
p7.font.name = "Segoe UI"
p7.space_before = Pt(12)

p8 = tf.add_paragraph()
p8.text = "Monday to Saturday: 09:30 AM – 09:00 PM\n(Sunday: Prior Appointment Only)"
p8.font.size = Pt(11)
p8.font.color.rgb = GRAY_LIGHT
p8.font.name = "Segoe UI"
p8.space_before = Pt(4)


# Right Column: Call to Action & Services Summary
create_card(slide8, Inches(7.2), Inches(1.8), Inches(5.3), Inches(5.0), CARD_BG, CYAN)

tb_cta = slide8.shapes.add_textbox(Inches(7.5), Inches(2.0), Inches(4.7), Inches(4.5))
tf_cta = tb_cta.text_frame
tf_cta.word_wrap = True

p = tf_cta.paragraphs[0]
p.text = "READY FOR YOUR NEXT EVENT?"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = GOLD
p.font.name = "Segoe UI"

p2 = tf_cta.add_paragraph()
p2.text = "Don't let faulty gear interrupt your show. Whether you are an event sound supplier, DJ, wedding organizer, or auditorium manager, Vikash Electronics guarantees quick and precision repair."
p2.font.size = Pt(11)
p2.font.color.rgb = GRAY_LIGHT
p2.font.name = "Segoe UI"
p2.space_before = Pt(8)

p3 = tf_cta.add_paragraph()
p3.text = "✔ Speakers & Heavy Subwoofers"
p3.font.size = Pt(11)
p3.font.bold = True
p3.font.color.rgb = WHITE
p3.font.name = "Segoe UI"
p3.space_before = Pt(12)

p4 = tf_cta.add_paragraph()
p4.text = "✔ High Power PA Amplifiers & Crossovers"
p4.font.size = Pt(11)
p4.font.bold = True
p4.font.color.rgb = WHITE
p4.font.name = "Segoe UI"
p4.space_before = Pt(4)

p5 = tf_cta.add_paragraph()
p5.text = "✔ Sharpy 7R/10R/15R Moving Head Lights"
p5.font.size = Pt(11)
p5.font.bold = True
p5.font.color.rgb = WHITE
p5.font.name = "Segoe UI"
p5.space_before = Pt(4)

p6 = tf_cta.add_paragraph()
p6.text = "✔ Sound Consoles, DMX Boards & Lasers"
p6.font.size = Pt(11)
p6.font.bold = True
p6.font.color.rgb = WHITE
p6.font.name = "Segoe UI"
p6.space_before = Pt(4)

p7 = tf_cta.add_paragraph()
p7.text = "🌐 Website: localhost:3000 (vikashelectronics.com)"
p7.font.size = Pt(11)
p7.font.bold = True
p7.font.color.rgb = CYAN
p7.font.name = "Segoe UI"
p7.space_before = Pt(14)

p8 = tf_cta.add_paragraph()
p8.text = "Call Now: +91 98254 85520"
p8.font.size = Pt(14)
p8.font.bold = True
p8.font.color.rgb = GOLD
p8.font.name = "Segoe UI"
p8.space_before = Pt(6)


# Save presentation
prs.save(OUTPUT_FILE)
print(f"Presentation successfully created at: {OUTPUT_FILE}")
