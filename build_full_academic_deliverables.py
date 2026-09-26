import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE
import docx
from docx import Document
from docx.shared import Inches as DocxInches, Pt as DocxPt, RGBColor as DocxRGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

BASE_DIR = r"d:\A_S_Projects\Vikas_Electronics"
PPTX_FILE_PRIMARY = os.path.join(BASE_DIR, "Vikash_Electronics_Project_Presentation.pptx")
PPTX_FILE = os.path.join(BASE_DIR, "Vikash_Electronics_Presentation.pptx")
PPTX_ACADEMIC_FILE = os.path.join(BASE_DIR, "Vikash_Electronics_Academic_Presentation.pptx")
PPTX_COMPLETE_FILE = os.path.join(BASE_DIR, "Vikash_Electronics_Complete_Academic_Presentation.pptx")
DOCX_FILE = os.path.join(BASE_DIR, "Project_1_Application_Form_Vikash_Electronics.docx")
HTML_FILE = os.path.join(BASE_DIR, "Project_1_Application_Form_Vikash_Electronics.html")

LOGO_UTU = os.path.join(BASE_DIR, "utu_flame_logo.png")
LOGO_AMTICS = os.path.join(BASE_DIR, "amtics_circle_logo.png")

# ==============================================================================
# COLOR PALETTE (Exact Match to Academic Reference PDF)
# ==============================================================================
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_BLACK = RGBColor(15, 23, 42)          # Primary Text / Heading
COLOR_RED = RGBColor(220, 38, 38)           # #DC2626 Crimson Red (subtitles, error boxes, accents)
COLOR_BLUE = RGBColor(2, 132, 199)          # #0284C7 Blue for Title
COLOR_DARK_BLUE = RGBColor(15, 76, 129)     # CSE Dept Subtitle
COLOR_GRAY_DARK = RGBColor(51, 65, 85)      # Slate 700
COLOR_GRAY_LIGHT = RGBColor(248, 250, 252)  # Slate 50
COLOR_TABLE_HEADER = RGBColor(230, 204, 204)# #E6CCCC Rosy Pink Table Header from template
COLOR_TABLE_BORDER = RGBColor(203, 213, 225)# Slate 300
COLOR_DIAGRAM_BG = RGBColor(240, 249, 255)  # Light Sky Blue Tint
COLOR_DIAGRAM_BORDER = RGBColor(14, 116, 144) # Cyan/Teal border
COLOR_ACCENT_BLUE = RGBColor(37, 99, 235)   # Arrow Blue

# ==============================================================================
# PART 1: GENERATE PPTX PRESENTATION (27 SLIDES WITH 6 UML DIAGRAMS)
# 100% VIKASH ELECTRONICS PROJECT-BASED CONTENT
# ==============================================================================
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_WHITE

def add_academic_header(slide, title_text, red_subtitle=""):
    """Adds standard academic heading matching the PDF template"""
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.3))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0)
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r_title = p.add_run()
    r_title.text = title_text
    r_title.font.size = Pt(28)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_BLACK
    r_title.font.name = "Arial"
    
    if red_subtitle:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r_sub = p2.add_run()
        r_sub.text = red_subtitle
        r_sub.font.size = Pt(22)
        r_sub.font.bold = True
        r_sub.font.color.rgb = COLOR_RED
        r_sub.font.name = "Arial"
        p2.space_before = Pt(4)

def add_diagram_box(slide, left, top, width, height, title, subtext="", border_color=COLOR_DIAGRAM_BORDER, bg_color=COLOR_WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    
    if subtext:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = subtext
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_GRAY_DARK
        r2.font.name = "Arial"
        p2.space_before = Pt(2)
    return shape

# ------------------------------------------------------------------------------
# SLIDE 1: TITLE SLIDE
# ------------------------------------------------------------------------------
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)

if os.path.exists(LOGO_UTU):
    slide1.shapes.add_picture(LOGO_UTU, Inches(0.8), Inches(0.4), width=Inches(1.2))

if os.path.exists(LOGO_AMTICS):
    slide1.shapes.add_picture(LOGO_AMTICS, Inches(11.3), Inches(0.4), width=Inches(1.2))

header_box = slide1.shapes.add_textbox(Inches(2.2), Inches(0.4), Inches(8.933), Inches(1.5))
tf = header_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Asha M. Tarsadia Institute of\nComputer Science and Technology"
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p_dept = tf.add_paragraph()
p_dept.alignment = PP_ALIGN.CENTER
r_dept = p_dept.add_run()
r_dept.text = "Computer Science and Engineering"
r_dept.font.size = Pt(20)
r_dept.font.bold = True
r_dept.font.color.rgb = COLOR_DARK_BLUE
r_dept.font.name = "Arial"
p_dept.space_before = Pt(6)

title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.25), Inches(11.333), Inches(1.3))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Vikash Electronics – Professional DJ & Electronic Equipment Repair\nShowcase & Service Management Web Application"
r.font.size = Pt(23)
r.font.bold = True
r.font.color.rgb = COLOR_BLUE
r.font.name = "Arial"

prep_box = slide1.shapes.add_textbox(Inches(2.0), Inches(4.0), Inches(9.333), Inches(1.5))
tf = prep_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Prepared by"
r.font.size = Pt(17)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = "Prathmesh Chaudhari"
r2.font.size = Pt(16)
r2.font.bold = True
r2.font.color.rgb = COLOR_BLACK
r2.font.name = "Arial"
p2.space_before = Pt(4)

p3 = tf.add_paragraph()
p3.alignment = PP_ALIGN.CENTER
r3 = p3.add_run()
r3.text = "(202503103510061)"
r3.font.size = Pt(15)
r3.font.color.rgb = COLOR_BLACK
r3.font.name = "Arial"
p3.space_before = Pt(2)

guide_box = slide1.shapes.add_textbox(Inches(2.0), Inches(5.65), Inches(9.333), Inches(1.4))
tf = guide_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Guided by"
r.font.size = Pt(17)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = "<Guide Name>"
r2.font.size = Pt(15)
r2.font.color.rgb = COLOR_BLACK
r2.font.name = "Arial"
p2.space_before = Pt(4)

# ------------------------------------------------------------------------------
# SLIDE 2: OUTLINE
# ------------------------------------------------------------------------------
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2)
add_academic_header(slide2, "Outline")

content_box = slide2.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.7), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

outline_items = [
    "Client Profile",
    "Project Introduction",
    "Objective(s)",
    "Problem Analysis",
    "Requirement Engineering",
    "Software Development Model",
    "Modeling",
    "Others"
]

for idx, item in enumerate(outline_items):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = item
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(14)

# ------------------------------------------------------------------------------
# SLIDE 3: CLIENT PROFILE
# ------------------------------------------------------------------------------
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3)
add_academic_header(slide3, "Client Profile")

content_box = slide3.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

sections = [
    ("Client name and address: ", "Vikash Electronics, a premier professional DJ sound system, stage lighting, amplifier, and electronic equipment repair center operating at Plot No. 199, Sarve No. 1-2, Sanjay Nagar, Udhna Yard, Udhna / Limbayat, Surat, Gujarat."),
    ("Your communication: ", "Requirements gathered through in-person workshop visits and WhatsApp/calls with the owner & technician (Mr. Murlidhar Chaudhari), covering speakers (voice coil rewinding), power amplifiers (output transistors), Sharpy moving head lights (optical calibration), sound mixers, diagnostic pricing, repair drop-offs, and workshop operating hours."),
    ("Contract letter: ", "Project undertaken as a client engagement to build a self-owned showcase catalog and direct customer repair appointment & inquiry platform, confirmed with the client before development.")
]

for s_idx, (sec_title, body) in enumerate(sections):
    p = tf.paragraphs[0] if s_idx == 0 else tf.add_paragraph()
    r_bold = p.add_run()
    r_bold.text = sec_title
    r_bold.font.size = Pt(16)
    r_bold.font.bold = True
    r_bold.font.color.rgb = COLOR_BLACK
    r_bold.font.name = "Arial"
    
    r_body = p.add_run()
    r_body.text = body
    r_body.font.size = Pt(15)
    r_body.font.color.rgb = COLOR_BLACK
    r_body.font.name = "Arial"
    p.space_before = Pt(20 if s_idx > 0 else 0)

# ------------------------------------------------------------------------------
# SLIDE 4: PROJECT INTRODUCTION
# ------------------------------------------------------------------------------
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4)
add_academic_header(slide4, "Project Introduction")

content_box = slide4.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

intro_paragraphs = [
    "Vikash Electronics is a prominent professional DJ and electronic repair workshop in Surat that previously depended entirely on walk-in customers and third-party aggregator directories to reach clients.",
    "This project is a dedicated repair showcase and management web application built exclusively for Vikash Electronics, listing its full service catalog (Speakers, Amplifiers, Sharpy Lights, DJ Lights, TVs, Sound Mixers) with diagnostic descriptions, bench images, and technical specifications.",
    "Customers (DJs, sound vendors, event organizers, stage technicians, and home users) can browse the collection, view equipment failure specifications, and schedule in-store repair drop-off sessions on a single website, without going through any third-party app.",
    "The site also carries the workshop's social media links, location, address, store hours, and direct contact phone (+91 98254 85520) / WhatsApp in one place.",
    "The goal is to reduce the client's dependency on aggregator platforms and save the high commission charges and listing delays typically paid to third-party marketplaces."
]

for idx, para in enumerate(intro_paragraphs):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = para
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(14)

# ------------------------------------------------------------------------------
# SLIDE 5: OBJECTIVE(S)
# ------------------------------------------------------------------------------
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5)
add_academic_header(slide5, "Objective(s)")

content_box = slide5.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

objectives_list = [
    "Build a single, self-owned website where customers can view the full DJ, stage lighting, and pro-audio electronic repair catalog with technical highlights and high-resolution workshop bench images.",
    "Allow customers to filter repair services by category (Audio, Stage Lighting, Visual, General) and submit in-store repair drop-off appointments directly on the website.",
    "Reduce Vikash Electronics's dependency on third-party listing apps and the commission charged by them.",
    "Display business information — location in Surat, address, contact number (+91 98254 85520), store hours, and direct WhatsApp links — in one place.",
    "Provide the workshop owner (Mr. Murlidhar Chaudhari) with a simple way to showcase specialized bench skills (voice coil rewinding, Sharpy optical calibration, transistor balancing) and manage incoming customer inquiries.",
    "Improve overall profit margins and brand visibility for the client by cutting out third-party delivery and aggregator charges."
]

for idx, item in enumerate(objectives_list):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = item
    r.font.size = Pt(13.5)
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(14)

# ------------------------------------------------------------------------------
# SLIDE 6: ORDER PROCESS FLOW
# ------------------------------------------------------------------------------
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6)
add_academic_header(slide6, "Order Process Flow")

flow_boxes = [
    "Browse\nMenu",
    "Select\nCategory",
    "View\nDetails",
    "Inquire /\nBook Repair",
    "Store Bench\nConsultation",
    "Repair\nConfirmed"
]

box_w = Inches(1.6)
box_h = Inches(0.9)
box_y = Inches(3.2)
start_x = Inches(0.8)
gap_x = Inches(0.4)

for idx, title in enumerate(flow_boxes):
    cur_x = start_x + idx * (box_w + gap_x)
    shape = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cur_x, box_y, box_w, box_h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_WHITE
    shape.line.color.rgb = COLOR_DIAGRAM_BORDER
    shape.line.width = Pt(1.5)
    p = shape.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    
    if idx < len(flow_boxes) - 1:
        arr = slide6.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, cur_x + box_w + Inches(0.08), box_y + Inches(0.32), Inches(0.24), Inches(0.25))
        arr.fill.solid()
        arr.fill.fore_color.rgb = COLOR_ACCENT_BLUE
        arr.line.fill.background()

foot_box = slide6.shapes.add_textbox(Inches(0.8), Inches(5.3), Inches(11.733), Inches(1.2))
tf = foot_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Customers move through the entire journey on Vikash Electronics's own\nwebsite — no aggregator hand-off, and no commission paid at any step."
r.font.size = Pt(16)
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

# ------------------------------------------------------------------------------
# SLIDE 7: PROBLEM ANALYSIS (Understand the problem)
# ------------------------------------------------------------------------------
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7)
add_academic_header(slide7, "Problem Analysis", "(Understand the problem)")

content_box = slide7.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_understand = [
    ("I. Who are the stockholders in the solution to the problem?",
     "The workshop proprietor & chief technician of Vikash Electronics, customers exploring electronic equipment repairs (DJs, sound rental vendors, event managers), senior bench testing staff, and the site administrator."),
    ("II. What are the unknowns?",
     "The customer preferred communication method, how custom component repairs (voice coil rewinding, Sharpy motors) will be managed, and how repair inquiry status will be tracked in real time."),
    ("III. Can the problem be compartmentalized?",
     "Yes — into modules: catalog showcase, category filter, equipment diagnostics modal, repair appointment booking, and WhatsApp inquiry panel."),
    ("IV. Can the problem be represented graphically?",
     "Yes, the flow can be represented through flowcharts and data flow diagrams covering browsing, filtering, and repair inquiry booking.")
]

for idx, (q, a) in enumerate(pa_understand):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(14)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = a
    r_a.font.size = Pt(13)
    r_a.font.color.rgb = COLOR_BLACK
    r_a.font.name = "Arial"
    p.space_before = Pt(12)

# ------------------------------------------------------------------------------
# SLIDE 8: PROBLEM ANALYSIS (Plan the solution)
# ------------------------------------------------------------------------------
slide8 = prs.slides.add_slide(blank_layout)
set_slide_background(slide8)
add_academic_header(slide8, "Problem Analysis", "(Plan the solution)")

content_box = slide8.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_plan = [
    ("I. Have you seen similar problem before?",
     "Yes, in platforms like electronics repair portals, pro-audio service showcases, and bespoke hardware maintenance websites."),
    ("II. Has a similar problem been solved?",
     "Yes — many specialized repair workshops now run their own direct showcase websites to avoid third-party aggregator fees and commissions."),
    ("III. Can sub-problems be defined?",
     "Yes: user registration/login, repair service catalog, category filtering, repair drop-off appointment, WhatsApp integration, and admin dashboard."),
    ("IV. Can you represent a solution in a manner that leads to effective implementation?",
     "Yes, through a modular architecture diagram and UML diagrams that lead to effective, phase-wise implementation.")
]

for idx, (q, a) in enumerate(pa_plan):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(14)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = a
    r_a.font.size = Pt(13)
    r_a.font.color.rgb = COLOR_BLACK
    r_a.font.name = "Arial"
    p.space_before = Pt(12)

# ------------------------------------------------------------------------------
# SLIDE 9: SYSTEM ARCHITECTURE
# ------------------------------------------------------------------------------
slide9 = prs.slides.add_slide(blank_layout)
set_slide_background(slide9)
add_academic_header(slide9, "System Architecture")

# Box 1: Customer (Top)
add_diagram_box(slide9, Inches(4.8), Inches(1.8), Inches(3.6), Inches(0.8), "Customer", "(Web / Mobile Browser)")

# Down arrow from Customer to Frontend
arr1 = slide9.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.5), Inches(2.65), Inches(0.2), Inches(0.35))
arr1.fill.solid(); arr1.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr1.line.fill.background()

# Box 2: Website Frontend
add_diagram_box(slide9, Inches(4.8), Inches(3.05), Inches(3.6), Inches(0.8), "Website Frontend", "(Catalog, Category Filter, Modals)")

# Down arrow from Frontend to Backend
arr2 = slide9.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.5), Inches(3.9), Inches(0.2), Inches(0.35))
arr2.fill.solid(); arr2.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr2.line.fill.background()

# Box 3: Backend Server (Center)
add_diagram_box(slide9, Inches(4.8), Inches(4.3), Inches(3.6), Inches(0.8), "Backend Server", "(Orders, Inquiries, Business Logic)")

# Left Arrow to Database
arr_l = slide9.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(4.1), Inches(4.55), Inches(0.6), Inches(0.2))
arr_l.fill.solid(); arr_l.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr_l.line.fill.background()

# Box 4: Database (Left)
add_diagram_box(slide9, Inches(1.2), Inches(4.3), Inches(2.8), Inches(0.8), "Database", "(Products, Appointments, Users)")

# Right Arrow to WhatsApp API
arr_r = slide9.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.5), Inches(4.55), Inches(0.6), Inches(0.2))
arr_r.fill.solid(); arr_r.fill.fore_color.rgb = COLOR_RED; arr_r.line.fill.background()

# Box 5: WhatsApp API (Right, red border)
add_diagram_box(slide9, Inches(9.2), Inches(4.3), Inches(2.8), Inches(0.8), "WhatsApp API", "(Direct Technician Chat)", border_color=COLOR_RED)

# Down Arrow to Admin Panel
arr_d = slide9.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.5), Inches(5.15), Inches(0.2), Inches(0.35))
arr_d.fill.solid(); arr_d.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr_d.line.fill.background()

# Box 6: Admin Panel (Bottom)
add_diagram_box(slide9, Inches(4.8), Inches(5.55), Inches(3.6), Inches(0.8), "Admin Panel", "(Owner: Catalog & Appointment Management)")

# Caption
foot_arch = slide9.shapes.add_textbox(Inches(0.8), Inches(6.5), Inches(11.733), Inches(0.8))
tf = foot_arch.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Customer requests flow from the website through the backend to the database and WhatsApp gateway;\nthe owner manages catalog and appointments through the admin panel."
r.font.size = Pt(13)
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

# ------------------------------------------------------------------------------
# SLIDE 10: PROBLEM ANALYSIS (Carry out the plan)
# ------------------------------------------------------------------------------
slide10 = prs.slides.add_slide(blank_layout)
set_slide_background(slide10)
add_academic_header(slide10, "Problem Analysis", "(Carry out the plan)")

content_box = slide10.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_carry = [
    ("I. Does the solution conform to the plan?",
     "Yes — development proceeds module by module (service catalog, category filter, appointment booking, WhatsApp routing, admin) as per the defined architecture."),
    ("II. Is each component part of the solution provably correct?",
     "Yes, each component is tested independently — for example, category filtering, detail modal triggers, phone number validations, and WhatsApp URL generation are each verified before integration.")
]

for idx, (q, a) in enumerate(pa_carry):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(15)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = a
    r_a.font.size = Pt(13.5)
    r_a.font.color.rgb = COLOR_BLACK
    r_a.font.name = "Arial"
    p.space_before = Pt(18)

# ------------------------------------------------------------------------------
# SLIDE 11: PROBLEM ANALYSIS (Examine the results)
# ------------------------------------------------------------------------------
slide11 = prs.slides.add_slide(blank_layout)
set_slide_background(slide11)
add_academic_header(slide11, "Problem Analysis", "(Examine the results)")

content_box = slide11.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_examine = [
    ("I. Is it possible to test each component part of the solution?",
     "Yes — catalog loading, category filtering, appointment submission success/failure, and Google Maps location routing are all verifiable on their own."),
    ("II. Does the solution produce results that conform to the data, functions, and features that are required?",
     "Yes — customers can view the repair catalog, explore equipment diagnostic details, book a repair appointment, and contact the workshop directly, all without a third-party app.")
]

for idx, (q, a) in enumerate(pa_examine):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(15)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = a
    r_a.font.size = Pt(13.5)
    r_a.font.color.rgb = COLOR_BLACK
    r_a.font.name = "Arial"
    p.space_before = Pt(18)

# ------------------------------------------------------------------------------
# SLIDE 12: REQUIREMENT ENGINEERING (Client-side communication)
# ------------------------------------------------------------------------------
slide12 = prs.slides.add_slide(blank_layout)
set_slide_background(slide12)
add_academic_header(slide12, "Requirement Engineering", "(Client-side communication)")

content_box = slide12.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

req_items = [
    "Complete repair catalog list with categories (Audio, Lighting, Visual, General), equipment bench images, diagnosis highlights, and component tags, matching the pro-audio positioning of the brand.",
    "Business details to be displayed: workshop address in Limbayat/Udhna, direct contact number (+91 98254 85520), and social media links (Instagram/WhatsApp/Facebook).",
    "Direct WhatsApp communication support that avoids the commission charged by third-party marketplace platforms.",
    "A simple way for the workshop staff/owner to view and manage incoming repair appointments and technical inquiries.",
    "A clean, mobile-friendly design consistent with Vikash Electronics's existing pro-audio and DJ brand identity.",
    "Easy updates to repair catalog items and pricing without needing developer help each time."
]

for idx, item in enumerate(req_items):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = item
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(14)

# ------------------------------------------------------------------------------
# SLIDES 13, 14, 15: REQUIREMENT ENGINEERING (Functional Requirements Tables)
# ------------------------------------------------------------------------------
def draw_single_fr_slide(slide, fr_id, title, desc, actor, inputs, outputs, intro_note=False):
    if intro_note:
        note_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(0.5))
        p_n = note_box.text_frame.paragraphs[0]
        p_n.text = "Enlisted below are the functional requirements for the website:"
        p_n.font.size = Pt(14)
        p_n.font.color.rgb = COLOR_BLACK
        p_n.font.name = "Arial"
        y_top = Inches(2.2)
    else:
        y_top = Inches(2.0)

    rows = 5
    cols = 2
    table_shape = slide.shapes.add_table(rows, cols, Inches(1.0), y_top, Inches(11.3), Inches(3.6))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(9.1)
    
    # Header Row with Rosy Pink Background (#E6CCCC)
    c00 = table.cell(0, 0); c00.fill.solid(); c00.fill.fore_color.rgb = COLOR_TABLE_HEADER
    c00.text = fr_id; p = c00.text_frame.paragraphs[0]; p.font.bold = True; p.font.size = Pt(13); p.font.color.rgb = COLOR_BLACK
    
    c01 = table.cell(0, 1); c01.fill.solid(); c01.fill.fore_color.rgb = COLOR_TABLE_HEADER
    c01.text = title; p = c01.text_frame.paragraphs[0]; p.font.bold = True; p.font.size = Pt(13); p.font.color.rgb = COLOR_BLACK
    
    data = [
        ("Description", desc),
        ("Actor(s)", actor),
        ("Input(s)", inputs),
        ("Output(s)", outputs)
    ]
    for r_idx, (lbl, val) in enumerate(data, start=1):
        c0 = table.cell(r_idx, 0); c0.fill.solid(); c0.fill.fore_color.rgb = COLOR_WHITE; c0.text = lbl
        p0 = c0.text_frame.paragraphs[0]; p0.font.bold = True; p0.font.size = Pt(12); p0.font.color.rgb = COLOR_BLACK
        
        c1 = table.cell(r_idx, 1); c1.fill.solid(); c1.fill.fore_color.rgb = COLOR_WHITE; c1.text = val
        p1 = c1.text_frame.paragraphs[0]; p1.font.size = Pt(12); p1.font.color.rgb = COLOR_BLACK

# Slide 13: FR-1 Registration
slide13 = prs.slides.add_slide(blank_layout)
set_slide_background(slide13)
add_academic_header(slide13, "Requirement Engineering", "(Functional Requirements)")
draw_single_fr_slide(
    slide13, "FR-1", "Registration",
    "Used to register new customers in the system.",
    "Customer",
    "Full name, Mobile number, e-mail Id, Workshop delivery address, New password, and Confirm password.",
    "Prompt dashboard with successful registration",
    intro_note=True
)

# Slide 14: FR-2 Catalog Browsing & Appointment Booking
slide14 = prs.slides.add_slide(blank_layout)
set_slide_background(slide14)
add_academic_header(slide14, "Requirement Engineering", "(Functional Requirements)")
draw_single_fr_slide(
    slide14, "FR-2", "Catalog Browsing & Appointment Booking",
    "Used to let the customer browse the collection and book a repair appointment.",
    "Customer",
    "Selected equipment item(s) (Speakers, Amplifiers, Sharpy Lights, Mixers), category filter, and repair appointment details.",
    "Equipment added to inquiry with updated details; appointment summary generated on booking"
)

# Slide 15: FR-3 Online Inquiry & Direct Contact
slide15 = prs.slides.add_slide(blank_layout)
set_slide_background(slide15)
add_academic_header(slide15, "Requirement Engineering", "(Functional Requirements)")
draw_single_fr_slide(
    slide15, "FR-3", "Online Inquiry & Direct Contact",
    "Used to let the customer contact the workshop technician directly on the website.",
    "Customer",
    "Equipment code / product ID, selected contact method (WhatsApp/Call), customer details.",
    "Inquiry confirmation and appointment status updated to 'Confirmed'"
)

# ------------------------------------------------------------------------------
# SLIDE 16: SOFTWARE DEVELOPMENT MODEL (Agile Scrum with Diagram)
# ------------------------------------------------------------------------------
slide16 = prs.slides.add_slide(blank_layout)
set_slide_background(slide16)
add_academic_header(slide16, "Software Development Model")

content_box = slide16.shapes.add_textbox(Inches(1.2), Inches(1.7), Inches(10.9), Inches(2.2))
tf = content_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
r = p.add_run()
r.text = "This project follows the Agile (Scrum) development model"
r.font.size = Pt(14)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK

p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "Work is broken into short sprints, each delivering a usable increment — e.g., catalog display first, then filtering, then inquiry, then admin panel"
r2.font.size = Pt(13)
r2.font.color.rgb = COLOR_BLACK
p2.space_before = Pt(8)

p3 = tf.add_paragraph()
r3 = p3.add_run()
r3.text = "Each sprint includes planning, development, testing, and a review with the client before moving to the next module"
r3.font.size = Pt(13)
r3.font.color.rgb = COLOR_BLACK
p3.space_before = Pt(8)

# 5 Cyclic Agile Boxes: Plan -> Design -> Develop -> Test -> Review
agile_steps = ["Plan", "Design", "Develop", "Test", "Review"]
agile_w = Inches(1.8)
agile_h = Inches(0.9)
agile_gap = Inches(0.4)
agile_start = Inches(1.4)
agile_y = Inches(4.3)

for idx, step in enumerate(agile_steps):
    bx = agile_start + idx * (agile_w + agile_gap)
    add_diagram_box(slide16, bx, agile_y, agile_w, agile_h, step)
    
    if idx < len(agile_steps) - 1:
        arr = slide16.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, bx + agile_w + Inches(0.08), agile_y + Inches(0.3), Inches(0.24), Inches(0.3))
        arr.fill.solid(); arr.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr.line.fill.background()

# Feedback loop line beneath
loop_line = slide16.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.3), Inches(5.65), Inches(8.8), Inches(0.04))
loop_line.fill.solid(); loop_line.fill.fore_color.rgb = COLOR_ACCENT_BLUE; loop_line.line.fill.background()

arr_up = slide16.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(2.2), Inches(5.25), Inches(0.2), Inches(0.4))
arr_up.fill.solid(); arr_up.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr_up.line.fill.background()

line_down = slide16.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(11.1), Inches(5.2), Inches(0.04), Inches(0.45))
line_down.fill.solid(); line_down.fill.fore_color.rgb = COLOR_ACCENT_BLUE; line_down.line.fill.background()

foot_agile = slide16.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.6))
p = foot_agile.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "repeats every sprint, with client review at the end of each cycle"
r.font.size = Pt(12)
r.font.color.rgb = COLOR_BLACK

# ------------------------------------------------------------------------------
# SLIDE 17: MODELING (Developing the use cases)
# ------------------------------------------------------------------------------
slide17 = prs.slides.add_slide(blank_layout)
set_slide_background(slide17)
add_academic_header(slide17, "Modeling", "(Developing the use cases)")

content_box = slide17.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

use_case_questions = [
    "I. Who are the primary and secondary actors?",
    "II. What are the actor's goal(s)?",
    "III. What preconditions should exist before the story begins?",
    "IV. What main tasks or functions are performed by the actor?",
    "V. What exceptions might be considered as the story is described?",
    "VI. What variations in the actor's interaction are possible?",
    "VII. What system information will the actor acquire, produce, or change?",
    "VIII. Will the actor have to inform the system about changes in the external environment?",
    "IX. What information does the actor desire from the system?",
    "X. Does the actor wish to be informed about unexpected changes?"
]

for idx, q in enumerate(use_case_questions):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = q
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(8)

# ------------------------------------------------------------------------------
# SLIDE 18: MODELING (USE CASE DIAGRAM) - 1st UML DIAGRAM
# ------------------------------------------------------------------------------
slide18 = prs.slides.add_slide(blank_layout)
set_slide_background(slide18)
add_academic_header(slide18, "Modeling", "(Use Case Diagram)")

# System boundary box
sys_box = slide18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(1.6), Inches(5.2), Inches(5.3))
sys_box.fill.background()
sys_box.line.color.rgb = COLOR_BLACK
sys_box.line.width = Pt(1.5)

sys_lbl = slide18.shapes.add_textbox(Inches(4.5), Inches(6.5), Inches(5.2), Inches(0.4))
p = sys_lbl.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Vikash Electronics System"
r.font.size = Pt(12)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK

# 6 Use Cases (Red-bordered ellipses matching reference PDF layout)
uc_items = [
    ("Register / Login", Inches(1.85)),
    ("Browse Repair Catalog", Inches(2.6)),
    ("Filter Equipment Categories", Inches(3.35)),
    ("Book Repair Appointment", Inches(4.1)),
    ("Make WhatsApp Inquiry", Inches(4.85)),
    ("Manage Services & Pricing", Inches(5.6))
]

for uc_text, y_pos in uc_items:
    oval = slide18.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.0), y_pos, Inches(4.2), Inches(0.6))
    oval.fill.solid()
    oval.fill.fore_color.rgb = COLOR_WHITE
    oval.line.color.rgb = COLOR_RED
    oval.line.width = Pt(1.5)
    p = oval.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = uc_text
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLACK

# Stick Actor function
def draw_stick_actor(slide, center_x, top_y, label_text):
    # Head
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, center_x - Inches(0.2), top_y, Inches(0.4), Inches(0.4))
    head.fill.background()
    head.line.color.rgb = COLOR_BLACK
    head.line.width = Pt(1.5)
    # Body
    body = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, center_x - Inches(0.01), top_y + Inches(0.4), Inches(0.02), Inches(0.55))
    body.fill.solid(); body.fill.fore_color.rgb = COLOR_BLACK; body.line.fill.background()
    # Arms
    arms = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, center_x - Inches(0.3), top_y + Inches(0.55), Inches(0.6), Inches(0.02))
    arms.fill.solid(); arms.fill.fore_color.rgb = COLOR_BLACK; arms.line.fill.background()
    # Legs
    leg_l = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, center_x - Inches(0.2), top_y + Inches(0.95), Inches(0.02), Inches(0.45))
    leg_l.fill.solid(); leg_l.fill.fore_color.rgb = COLOR_BLACK; leg_l.line.fill.background()
    leg_r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, center_x + Inches(0.2), top_y + Inches(0.95), Inches(0.02), Inches(0.45))
    leg_r.fill.solid(); leg_r.fill.fore_color.rgb = COLOR_BLACK; leg_r.line.fill.background()
    # Label
    tb = slide.shapes.add_textbox(center_x - Inches(1.0), top_y + Inches(1.45), Inches(2.0), Inches(0.4))
    p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label_text; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = COLOR_BLACK

# Customer Actor on Left
draw_stick_actor(slide18, Inches(3.0), Inches(3.0), "Customer")

# Association Lines from Customer to Use Cases
for _, y_pos in uc_items[:5]:
    line = slide18.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.3), y_pos + Inches(0.28), Inches(1.7), Inches(0.02))
    line.fill.solid(); line.fill.fore_color.rgb = COLOR_GRAY_DARK; line.line.fill.background()

# Owner / Admin Actor on Left/Bottom
draw_stick_actor(slide18, Inches(3.0), Inches(5.0), "Owner /\nTechnician")

# Lines from Admin to Manage Services & Bookings
for y_pos in [Inches(4.1), Inches(4.85), Inches(5.6)]:
    line = slide18.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(3.3), y_pos + Inches(0.3), Inches(1.7), Inches(0.02))
    line.fill.solid(); line.fill.fore_color.rgb = COLOR_GRAY_DARK; line.line.fill.background()

# ------------------------------------------------------------------------------
# SLIDE 19: MODELING (CLASS DIAGRAM) - 2nd UML DIAGRAM
# ------------------------------------------------------------------------------
slide19 = prs.slides.add_slide(blank_layout)
set_slide_background(slide19)
add_academic_header(slide19, "Modeling", "(Class Diagram)")

def draw_uml_class(slide, left, top, width, height, class_name, attributes, methods):
    # Outer rectangle
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid(); shape.fill.fore_color.rgb = COLOR_WHITE
    shape.line.color.rgb = COLOR_BLACK; shape.line.width = Pt(1.2)
    
    # Header box
    hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.4))
    hdr.fill.solid(); hdr.fill.fore_color.rgb = COLOR_DIAGRAM_BG
    hdr.line.color.rgb = COLOR_BLACK; hdr.line.width = Pt(1.2)
    p = hdr.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = class_name; r.font.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = COLOR_BLACK
    
    # Body
    tb = slide.shapes.add_textbox(left + Inches(0.05), top + Inches(0.42), width - Inches(0.1), height - Inches(0.45))
    tf = tb.text_frame; tf.word_wrap = True
    
    for idx, attr in enumerate(attributes):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = f"- {attr}"; r.font.size = Pt(8.5); r.font.color.rgb = COLOR_BLACK
    
    p_sep = tf.add_paragraph()
    r = p_sep.add_run(); r.text = "--------------------------------------"; r.font.size = Pt(6); r.font.color.rgb = COLOR_TABLE_BORDER
    
    for meth in methods:
        p = tf.add_paragraph()
        r = p.add_run(); r.text = f"+ {meth}()"; r.font.size = Pt(8.5); r.font.color.rgb = COLOR_BLACK

# Row 1 Classes
draw_uml_class(slide19, Inches(1.5), Inches(1.8), Inches(2.8), Inches(1.9),
               "Customer",
               ["customerId", "name", "email", "phone"],
               ["register", "login", "bookRepair"])

draw_uml_class(slide19, Inches(5.3), Inches(1.8), Inches(2.8), Inches(1.9),
               "RepairAppointment",
               ["appointmentId", "date", "status", "serviceType"],
               ["addEquipment", "confirmBooking"])

draw_uml_class(slide19, Inches(9.1), Inches(1.8), Inches(2.8), Inches(1.9),
               "Inquiry",
               ["inquiryId", "equipmentCode", "status", "channel"],
               ["sendWhatsApp", "resolveQuery"])

# Row 2 Classes
draw_uml_class(slide19, Inches(1.5), Inches(4.5), Inches(2.8), Inches(1.9),
               "Technician",
               ["technicianId", "name", "specialization"],
               ["manageCatalog", "viewAppointments"])

draw_uml_class(slide19, Inches(5.3), Inches(4.5), Inches(2.8), Inches(1.9),
               "EquipmentItem",
               ["equipmentCode", "faultType", "brand"],
               ["getSpecs", "diagnoseFault"])

draw_uml_class(slide19, Inches(9.1), Inches(4.5), Inches(2.8), Inches(1.9),
               "ServiceCategory",
               ["categoryId", "name", "estimatePrice", "category"],
               ["updatePrice", "getServices"])

# Associations: Horizontal Line between Customer & RepairAppointment
l1 = slide19.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.3), Inches(2.7), Inches(1.0), Inches(0.02))
l1.fill.solid(); l1.fill.fore_color.rgb = COLOR_BLACK; l1.line.fill.background()
tb = slide19.shapes.add_textbox(Inches(4.3), Inches(2.35), Inches(1.0), Inches(0.3))
tb.text_frame.paragraphs[0].text = "places\n1      *"; tb.text_frame.paragraphs[0].font.size = Pt(7.5); tb.text_frame.paragraphs[0].font.color.rgb = COLOR_BLACK

# Line between RepairAppointment & Inquiry
l2 = slide19.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.1), Inches(2.7), Inches(1.0), Inches(0.02))
l2.fill.solid(); l2.fill.fore_color.rgb = COLOR_BLACK; l2.line.fill.background()
tb = slide19.shapes.add_textbox(Inches(8.1), Inches(2.35), Inches(1.0), Inches(0.3))
tb.text_frame.paragraphs[0].text = "has\n1      1"; tb.text_frame.paragraphs[0].font.size = Pt(7.5); tb.text_frame.paragraphs[0].font.color.rgb = COLOR_BLACK

# Vertical Line between RepairAppointment & EquipmentItem
l3 = slide19.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.7), Inches(3.7), Inches(0.02), Inches(0.8))
l3.fill.solid(); l3.fill.fore_color.rgb = COLOR_BLACK; l3.line.fill.background()
tb = slide19.shapes.add_textbox(Inches(5.7), Inches(3.9), Inches(1.0), Inches(0.3))
tb.text_frame.paragraphs[0].text = "contains\n1\n*"; tb.text_frame.paragraphs[0].font.size = Pt(7.5); tb.text_frame.paragraphs[0].font.color.rgb = COLOR_BLACK

# Horizontal Line between EquipmentItem & ServiceCategory
l4 = slide19.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.1), Inches(5.4), Inches(1.0), Inches(0.02))
l4.fill.solid(); l4.fill.fore_color.rgb = COLOR_BLACK; l4.line.fill.background()
tb = slide19.shapes.add_textbox(Inches(8.1), Inches(5.05), Inches(1.0), Inches(0.3))
tb.text_frame.paragraphs[0].text = "references\n*      1"; tb.text_frame.paragraphs[0].font.size = Pt(7.5); tb.text_frame.paragraphs[0].font.color.rgb = COLOR_BLACK

# Line between Technician and ServiceCategory (Underneath)
l5_down = slide19.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.9), Inches(6.4), Inches(0.02), Inches(0.4))
l5_down.fill.solid(); l5_down.fill.fore_color.rgb = COLOR_BLACK; l5_down.line.fill.background()

l5_hor = slide19.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.9), Inches(6.8), Inches(7.6), Inches(0.02))
l5_hor.fill.solid(); l5_hor.fill.fore_color.rgb = COLOR_BLACK; l5_hor.line.fill.background()

l5_up = slide19.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10.5), Inches(6.4), Inches(0.02), Inches(0.4))
l5_up.fill.solid(); l5_up.fill.fore_color.rgb = COLOR_BLACK; l5_up.line.fill.background()

tb = slide19.shapes.add_textbox(Inches(5.5), Inches(6.55), Inches(2.5), Inches(0.3))
tb.text_frame.paragraphs[0].text = "manages 1  *"; tb.text_frame.paragraphs[0].font.size = Pt(8); tb.text_frame.paragraphs[0].font.color.rgb = COLOR_BLACK

# ------------------------------------------------------------------------------
# SLIDE 20: MODELING (SEQUENCE DIAGRAM) - 3rd UML DIAGRAM
# ------------------------------------------------------------------------------
slide20 = prs.slides.add_slide(blank_layout)
set_slide_background(slide20)
add_academic_header(slide20, "Modeling", "(Sequence Diagram)")

# Subtitle
sub_box = slide20.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.4))
p = sub_box.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Use case: Browse Repair Catalog & Inquire"; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = COLOR_BLACK

lifelines = [
    ("Customer", Inches(1.6)),
    ("Website\n(Frontend)", Inches(4.0)),
    ("Backend\nServer", Inches(6.6)),
    ("WhatsApp\nGateway", Inches(9.0)),
    ("Database", Inches(11.4))
]

for name, x in lifelines:
    add_diagram_box(slide20, x - Inches(0.8), Inches(2.0), Inches(1.6), Inches(0.6), name)
    # Vertical line
    line = slide20.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.6), Inches(0.02), Inches(4.4))
    line.fill.solid(); line.fill.fore_color.rgb = COLOR_TABLE_BORDER; line.line.fill.background()

seq_messages = [
    (Inches(1.6), Inches(4.0), Inches(2.9), "Browse Catalog", True),
    (Inches(4.0), Inches(6.6), Inches(3.4), "Get Equipment Services", True),
    (Inches(6.6), Inches(11.4), Inches(3.9), "Query Services", True),
    (Inches(11.4), Inches(6.6), Inches(4.4), "Return Catalog Data", False),
    (Inches(6.6), Inches(4.0), Inches(4.9), "Catalog Data", False),
    (Inches(1.6), Inches(4.0), Inches(5.4), "Select Equipment / Book Repair", True),
    (Inches(4.0), Inches(6.6), Inches(5.8), "Create Booking", True),
    (Inches(6.6), Inches(9.0), Inches(6.1), "Send Inquiry", True),
    (Inches(9.0), Inches(6.6), Inches(6.35), "Inquiry Sent", False),
    (Inches(6.6), Inches(11.4), Inches(6.55), "Save Appointment", True),
    (Inches(6.6), Inches(4.0), Inches(6.75), "Booking Confirmed", False),
    (Inches(4.0), Inches(1.6), Inches(6.95), "Show Confirmation", False)
]

for x1, x2, y, text, is_forward in seq_messages:
    left = min(x1, x2)
    w = abs(x2 - x1)
    
    m_line = slide20.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, y, w, Inches(0.02))
    m_line.fill.solid(); m_line.fill.fore_color.rgb = COLOR_BLACK; m_line.line.fill.background()
    
    if x2 > x1:
        arr = slide20.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x2 - Inches(0.1), y - Inches(0.04), Inches(0.1), Inches(0.1))
    else:
        arr = slide20.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, x2, y - Inches(0.04), Inches(0.1), Inches(0.1))
    arr.fill.solid(); arr.fill.fore_color.rgb = COLOR_BLACK; arr.line.fill.background()
    
    lbl = slide20.shapes.add_textbox(left, y - Inches(0.24), w, Inches(0.25))
    p = lbl.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(8); r.font.bold = True; r.font.color.rgb = COLOR_BLACK

# ------------------------------------------------------------------------------
# SLIDE 21: MODELING (ACTIVITY DIAGRAM) - 4th UML DIAGRAM
# ------------------------------------------------------------------------------
slide21 = prs.slides.add_slide(blank_layout)
set_slide_background(slide21)
add_academic_header(slide21, "Modeling", "(Activity Diagram)")

# Initial Black Circle
c_init = slide21.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.45), Inches(1.7), Inches(0.35), Inches(0.35))
c_init.fill.solid(); c_init.fill.fore_color.rgb = COLOR_BLACK; c_init.line.fill.background()

arr_d1 = slide21.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.55), Inches(2.05), Inches(0.15), Inches(0.25))
arr_d1.fill.solid(); arr_d1.fill.fore_color.rgb = COLOR_BLACK; arr_d1.line.fill.background()

# Step 1: Browse Repair Catalog
add_diagram_box(slide21, Inches(5.1), Inches(2.3), Inches(3.0), Inches(0.55), "Browse Repair Catalog")

arr_d2 = slide21.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.55), Inches(2.85), Inches(0.15), Inches(0.25))
arr_d2.fill.solid(); arr_d2.fill.fore_color.rgb = COLOR_BLACK; arr_d2.line.fill.background()

# Step 2: Select Equipment / Category
add_diagram_box(slide21, Inches(5.1), Inches(3.1), Inches(3.0), Inches(0.55), "Select Equipment / Category")

arr_d3 = slide21.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.55), Inches(3.65), Inches(0.15), Inches(0.25))
arr_d3.fill.solid(); arr_d3.fill.fore_color.rgb = COLOR_BLACK; arr_d3.line.fill.background()

# Step 3: Proceed to Inquire / Book
add_diagram_box(slide21, Inches(5.1), Inches(3.9), Inches(3.0), Inches(0.55), "Proceed to Inquire / Book")

arr_d4 = slide21.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.55), Inches(4.45), Inches(0.15), Inches(0.25))
arr_d4.fill.solid(); arr_d4.fill.fore_color.rgb = COLOR_BLACK; arr_d4.line.fill.background()

# Decision Diamond: Inquiry Valid?
diamond = slide21.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(5.5), Inches(4.7), Inches(2.2), Inches(0.95))
diamond.fill.solid(); diamond.fill.fore_color.rgb = COLOR_WHITE
diamond.line.color.rgb = COLOR_RED; diamond.line.width = Pt(1.5)
p = diamond.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Inquiry\nValid?"; r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = COLOR_RED

# Branch: No -> Show Validation Error (Right)
arr_no = slide21.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(7.7), Inches(5.08), Inches(0.8), Inches(0.15))
arr_no.fill.solid(); arr_no.fill.fore_color.rgb = COLOR_BLACK; arr_no.line.fill.background()

tb_no = slide21.shapes.add_textbox(Inches(7.8), Inches(4.8), Inches(0.6), Inches(0.3))
tb_no.text_frame.paragraphs[0].text = "No"; tb_no.text_frame.paragraphs[0].font.size = Pt(9); tb_no.text_frame.paragraphs[0].font.bold = True

box_err = slide21.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.5), Inches(4.85), Inches(2.0), Inches(0.6))
box_err.fill.solid(); box_err.fill.fore_color.rgb = COLOR_WHITE
box_err.line.color.rgb = COLOR_RED; box_err.line.width = Pt(1.5)
p = box_err.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Show Validation\nError"; r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = COLOR_RED

# Retry Line back up to Proceed
l_up = slide21.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9.5), Inches(4.15), Inches(0.02), Inches(0.7))
l_up.fill.solid(); l_up.fill.fore_color.rgb = COLOR_RED; l_up.line.fill.background()

l_left = slide21.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.1), Inches(4.15), Inches(1.4), Inches(0.02))
l_left.fill.solid(); l_left.fill.fore_color.rgb = COLOR_RED; l_left.line.fill.background()

arr_retry = slide21.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(8.1), Inches(4.08), Inches(0.15), Inches(0.15))
arr_retry.fill.solid(); arr_retry.fill.fore_color.rgb = COLOR_RED; arr_retry.line.fill.background()

tb_ret = slide21.shapes.add_textbox(Inches(8.8), Inches(3.9), Inches(0.8), Inches(0.3))
tb_ret.text_frame.paragraphs[0].text = "Retry"; tb_ret.text_frame.paragraphs[0].font.size = Pt(8.5); tb_ret.text_frame.paragraphs[0].font.bold = True; tb_ret.text_frame.paragraphs[0].font.color.rgb = COLOR_RED

# Branch: Yes (Down)
arr_yes = slide21.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.55), Inches(5.65), Inches(0.15), Inches(0.25))
arr_yes.fill.solid(); arr_yes.fill.fore_color.rgb = COLOR_BLACK; arr_yes.line.fill.background()

tb_yes = slide21.shapes.add_textbox(Inches(6.7), Inches(5.65), Inches(0.6), Inches(0.3))
tb_yes.text_frame.paragraphs[0].text = "Yes"; tb_yes.text_frame.paragraphs[0].font.size = Pt(9); tb_yes.text_frame.paragraphs[0].font.bold = True

# Step 4: Confirm Appointment
add_diagram_box(slide21, Inches(5.1), Inches(5.9), Inches(3.0), Inches(0.55), "Confirm Repair Appointment")

arr_d5 = slide21.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.55), Inches(6.45), Inches(0.15), Inches(0.2))
arr_d5.fill.solid(); arr_d5.fill.fore_color.rgb = COLOR_BLACK; arr_d5.line.fill.background()

# Step 5: Notify Technician & Customer
add_diagram_box(slide21, Inches(5.1), Inches(6.65), Inches(3.0), Inches(0.5), "Notify Technician & Customer")

# Final Bullseye Circle
c_out = slide21.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.45), Inches(7.2), Inches(0.35), Inches(0.35))
c_out.fill.background(); c_out.line.color.rgb = COLOR_BLACK; c_out.line.width = Pt(1.5)

c_in = slide21.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.5), Inches(7.25), Inches(0.25), Inches(0.25))
c_in.fill.solid(); c_in.fill.fore_color.rgb = COLOR_BLACK; c_in.line.fill.background()

# ------------------------------------------------------------------------------
# SLIDE 22: MODELING (STATE DIAGRAM) - 5th UML DIAGRAM
# ------------------------------------------------------------------------------
slide22 = prs.slides.add_slide(blank_layout)
set_slide_background(slide22)
add_academic_header(slide22, "Modeling", "(State Diagram)")

sub_box = slide22.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.733), Inches(0.4))
p = sub_box.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Inquiry & Repair Appointment lifecycle"; r.font.size = Pt(16); r.font.bold = True; r.font.color.rgb = COLOR_BLACK

# Initial Node (Black circle)
c_init = slide22.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.2), Inches(3.4), Inches(0.35), Inches(0.35))
c_init.fill.solid(); c_init.fill.fore_color.rgb = COLOR_BLACK; c_init.line.fill.background()

arr_init = slide22.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(1.55), Inches(3.5), Inches(0.3), Inches(0.15))
arr_init.fill.solid(); arr_init.fill.fore_color.rgb = COLOR_BLACK; arr_init.line.fill.background()

state_boxes = [
    ("Submitted", Inches(1.9)),
    ("Confirmed", Inches(4.3)),
    ("Repair\nScheduled", Inches(6.7)),
    ("Bench Testing\nIn-Store", Inches(9.1)),
    ("Delivered", Inches(11.0))
]

for name, x in state_boxes:
    w = Inches(1.4) if "Delivered" in name else Inches(1.6)
    add_diagram_box(slide22, x, Inches(3.2), w, Inches(0.75), name)

transitions = [
    (Inches(3.5), Inches(4.3), Inches(3.57), "verified"),
    (Inches(5.9), Inches(6.7), Inches(3.57), "accepted"),
    (Inches(8.3), Inches(9.1), Inches(3.57), "ready"),
    (Inches(10.7), Inches(11.0), Inches(3.57), "delivered")
]

for x1, x2, y, lbl_text in transitions:
    arr = slide22.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x1 + Inches(0.05), y, x2 - x1 - Inches(0.1), Inches(0.15))
    arr.fill.solid(); arr.fill.fore_color.rgb = COLOR_BLACK; arr.line.fill.background()
    
    tb = slide22.shapes.add_textbox(x1, y - Inches(0.3), x2 - x1, Inches(0.25))
    p = tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = lbl_text; r.font.size = Pt(8); r.font.color.rgb = COLOR_BLACK

# Final Bullseye after Delivered
fin_out = slide22.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12.6), Inches(3.4), Inches(0.35), Inches(0.35))
fin_out.fill.background(); fin_out.line.color.rgb = COLOR_BLACK; fin_out.line.width = Pt(1.5)

fin_in = slide22.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12.65), Inches(3.45), Inches(0.25), Inches(0.25))
fin_in.fill.solid(); fin_in.fill.fore_color.rgb = COLOR_BLACK; fin_in.line.fill.background()

arr_fin = slide22.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(12.4), Inches(3.5), Inches(0.2), Inches(0.15))
arr_fin.fill.solid(); arr_fin.fill.fore_color.rgb = COLOR_BLACK; arr_fin.line.fill.background()

# Cancelled State Box (Red, below)
box_cancel = slide22.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.3), Inches(5.2), Inches(1.6), Inches(0.75))
box_cancel.fill.solid(); box_cancel.fill.fore_color.rgb = COLOR_WHITE
box_cancel.line.color.rgb = COLOR_RED; box_cancel.line.width = Pt(1.5)
p = box_cancel.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Cancelled"; r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = COLOR_RED

# Transition 1 to Cancelled: From Submitted
l_sub_down = slide22.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.7), Inches(3.95), Inches(0.02), Inches(1.55))
l_sub_down.fill.solid(); l_sub_down.fill.fore_color.rgb = COLOR_BLACK; l_sub_down.line.fill.background()

l_sub_right = slide22.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(2.7), Inches(5.5), Inches(1.6), Inches(0.15))
l_sub_right.fill.solid(); l_sub_right.fill.fore_color.rgb = COLOR_BLACK; l_sub_right.line.fill.background()

tb_inv = slide22.shapes.add_textbox(Inches(2.8), Inches(4.5), Inches(1.4), Inches(0.3))
tb_inv.text_frame.paragraphs[0].text = "invalid\ndetails"; tb_inv.text_frame.paragraphs[0].font.size = Pt(8); tb_inv.text_frame.paragraphs[0].font.color.rgb = COLOR_BLACK

# Transition 2 to Cancelled: From Confirmed
arr_conf_down = slide22.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(5.0), Inches(3.95), Inches(0.15), Inches(1.25))
arr_conf_down.fill.solid(); arr_conf_down.fill.fore_color.rgb = COLOR_BLACK; arr_conf_down.line.fill.background()

tb_canc = slide22.shapes.add_textbox(Inches(4.2), Inches(4.4), Inches(1.5), Inches(0.3))
tb_canc.text_frame.paragraphs[0].text = "customer\ncancels"; tb_canc.text_frame.paragraphs[0].font.size = Pt(8); tb_canc.text_frame.paragraphs[0].font.color.rgb = COLOR_BLACK

# Final Bullseye after Cancelled
arr_c_fin = slide22.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.9), Inches(5.5), Inches(0.4), Inches(0.15))
arr_c_fin.fill.solid(); arr_c_fin.fill.fore_color.rgb = COLOR_BLACK; arr_c_fin.line.fill.background()

c_c_out = slide22.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.3), Inches(5.4), Inches(0.35), Inches(0.35))
c_c_out.fill.background(); c_c_out.line.color.rgb = COLOR_BLACK; c_c_out.line.width = Pt(1.5)

c_c_in = slide22.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.35), Inches(5.45), Inches(0.25), Inches(0.25))
c_c_in.fill.solid(); c_c_in.fill.fore_color.rgb = COLOR_BLACK; c_c_in.line.fill.background()

# ------------------------------------------------------------------------------
# SLIDE 23: DEPLOYMENT DIAGRAM - 6th UML DIAGRAM
# ------------------------------------------------------------------------------
slide23 = prs.slides.add_slide(blank_layout)
set_slide_background(slide23)
add_academic_header(slide23, "Deployment Diagram")

# 3D-styled Nodes
def draw_cube_node(slide, left, top, width, height, stereotype, node_title, sub_desc=""):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid(); shape.fill.fore_color.rgb = COLOR_WHITE
    shape.line.color.rgb = COLOR_BLACK; shape.line.width = Pt(1.5)
    
    tf = shape.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = f"<<{stereotype}>>\n"; r.font.size = Pt(9); r.font.color.rgb = COLOR_GRAY_DARK
    
    r_t = p.add_run(); r_t.text = f"{node_title}\n"; r_t.font.size = Pt(11); r_t.font.bold = True; r_t.font.color.rgb = COLOR_BLACK
    
    if sub_desc:
        r_d = p.add_run(); r_d.text = sub_desc; r_d.font.size = Pt(8.5); r_d.font.color.rgb = COLOR_BLACK
    
    top_edge = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, left + Inches(0.1), top - Inches(0.18), width, Inches(0.18))
    top_edge.fill.solid(); top_edge.fill.fore_color.rgb = COLOR_GRAY_LIGHT
    top_edge.line.color.rgb = COLOR_BLACK; top_edge.line.width = Pt(1.2)
    
    side_edge = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + width, top - Inches(0.08), Inches(0.15), height)
    side_edge.fill.solid(); side_edge.fill.fore_color.rgb = COLOR_TABLE_BORDER
    side_edge.line.color.rgb = COLOR_BLACK; side_edge.line.width = Pt(1.2)
    return shape

# Node 1: Customer Device
draw_cube_node(slide23, Inches(1.0), Inches(2.2), Inches(3.0), Inches(1.3), "device", "Customer Device", "Web / Mobile Browser")

# Node 2: Web Server
draw_cube_node(slide23, Inches(5.1), Inches(2.2), Inches(3.0), Inches(1.3), "server", "Web Server", "Frontend + Backend")

# Node 3: WhatsApp Gateway
draw_cube_node(slide23, Inches(9.2), Inches(2.2), Inches(3.0), Inches(1.3), "external system", "WhatsApp Gateway", "Direct Technician Chat")

# Node 4: Database Server
draw_cube_node(slide23, Inches(5.1), Inches(4.7), Inches(3.0), Inches(1.3), "server", "Database Server", "Catalog, Inquiries, Users")

# Connector: Customer Device <-> Web Server (HTTPS)
c1 = slide23.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.15), Inches(2.8), Inches(0.95), Inches(0.02))
c1.fill.solid(); c1.fill.fore_color.rgb = COLOR_BLACK; c1.line.fill.background()

tb_https = slide23.shapes.add_textbox(Inches(4.15), Inches(2.5), Inches(0.95), Inches(0.3))
p = tb_https.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
p.text = "HTTPS"; p.font.size = Pt(8.5); p.font.bold = True

# Connector: Web Server <-> WhatsApp Gateway (HTTPS / API)
c2 = slide23.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.25), Inches(2.8), Inches(0.95), Inches(0.02))
c2.fill.solid(); c2.fill.fore_color.rgb = COLOR_BLACK; c2.line.fill.background()

tb_api = slide23.shapes.add_textbox(Inches(8.25), Inches(2.45), Inches(0.95), Inches(0.35))
p = tb_api.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
p.text = "HTTPS /\nAPI"; p.font.size = Pt(8); p.font.bold = True

# Connector: Web Server <-> Database Server (TCP / SQL)
c3 = slide23.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6.6), Inches(3.5), Inches(0.02), Inches(1.2))
c3.fill.solid(); c3.fill.fore_color.rgb = COLOR_BLACK; c3.line.fill.background()

tb_sql = slide23.shapes.add_textbox(Inches(6.65), Inches(3.9), Inches(1.2), Inches(0.3))
p = tb_sql.text_frame.paragraphs[0]; p.text = "TCP / SQL"; p.font.size = Pt(8.5); p.font.bold = True

# Bottom caption
foot_dep = slide23.shapes.add_textbox(Inches(0.8), Inches(6.4), Inches(11.733), Inches(0.8))
tf = foot_dep.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "The customer's browser communicates with the web server over HTTPS;\nthe server queries the database and calls the WhatsApp gateway's API to process inquiries."
r.font.size = Pt(14)
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

# ------------------------------------------------------------------------------
# SLIDES 24 & 25: MODELING (USE CASES TABLES)
# ------------------------------------------------------------------------------
def draw_detailed_uc_table(slide, uc_title, actor, goal, precond, trigger):
    rows = 5
    cols = 2
    table_shape = slide.shapes.add_table(rows, cols, Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.8))
    table = table_shape.table
    table.columns[0].width = Inches(2.4)
    table.columns[1].width = Inches(8.9)
    
    c00 = table.cell(0, 0); c00.fill.solid(); c00.fill.fore_color.rgb = COLOR_TABLE_HEADER
    c00.text = uc_title; p = c00.text_frame.paragraphs[0]; p.font.bold = True; p.font.size = Pt(13); p.font.color.rgb = COLOR_BLACK
    
    c01 = table.cell(0, 1); c01.fill.solid(); c01.fill.fore_color.rgb = COLOR_TABLE_HEADER
    c01.text = ""; p = c01.text_frame.paragraphs[0]; p.font.bold = True; p.font.size = Pt(13); p.font.color.rgb = COLOR_BLACK
    
    data = [
        ("Primary Actor(s)", actor),
        ("Goal of the use case", goal),
        ("Precondition(s)", precond),
        ("Trigger", trigger)
    ]
    for r_idx, (lbl, val) in enumerate(data, start=1):
        c0 = table.cell(r_idx, 0); c0.fill.solid(); c0.fill.fore_color.rgb = COLOR_WHITE; c0.text = lbl
        p0 = c0.text_frame.paragraphs[0]; p0.font.bold = True; p0.font.size = Pt(12); p0.font.color.rgb = COLOR_BLACK
        
        c1 = table.cell(r_idx, 1); c1.fill.solid(); c1.fill.fore_color.rgb = COLOR_WHITE; c1.text = val
        p1 = c1.text_frame.paragraphs[0]; p1.font.size = Pt(12); p1.font.color.rgb = COLOR_BLACK

# Slide 24: Browse & Inquire Equipment Repair
slide24 = prs.slides.add_slide(blank_layout)
set_slide_background(slide24)
add_academic_header(slide24, "Modeling", "(Use cases)")
draw_detailed_uc_table(
    slide24, "Browse & Inquire Equipment Repair",
    "Customer",
    "To browse the repair catalog and successfully inquire or book a diagnostic appointment with workshop technicians",
    "Customer has an active account, a stable internet connection, and the repair services are available",
    "Customer selects equipment item(s) from catalog and clicks 'Book Repair' or 'Inquire Workshop'"
)

# Slide 25: Direct WhatsApp Contact
slide25 = prs.slides.add_slide(blank_layout)
set_slide_background(slide25)
add_academic_header(slide25, "Modeling", "(Use cases)")
draw_detailed_uc_table(
    slide25, "Direct WhatsApp\nContact",
    "Customer",
    "To initiate direct WhatsApp inquiry for selected electronic equipment without a third-party app",
    "Customer has reviewed the equipment details and clicked the WhatsApp contact button",
    "Customer selects contact method (WhatsApp) and sends the pre-filled equipment inquiry code"
)

# ------------------------------------------------------------------------------
# SLIDE 26: OTHERS
# ------------------------------------------------------------------------------
slide26 = prs.slides.add_slide(blank_layout)
set_slide_background(slide26)
add_academic_header(slide26, "Others")

content_box = slide26.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

others_items = [
    "Technology stack under consideration: Next.js (React) front-end, Node.js back-end framework with database, integrated with a WhatsApp communication gateway for customer inquiries",
    "Additional planned features: repair appointment history for customers, an admin view for the workshop to manage catalog items and incoming repair jobs, and direct links to the workshop's social media and YouTube pages",
    "Screenshots of development will be added here as the project progresses"
]

for idx, item in enumerate(others_items):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = item
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(18)

# ------------------------------------------------------------------------------
# SLIDE 27: THANK YOU
# ------------------------------------------------------------------------------
slide27 = prs.slides.add_slide(blank_layout)
set_slide_background(slide27)

if os.path.exists(LOGO_UTU):
    slide27.shapes.add_picture(LOGO_UTU, Inches(0.8), Inches(0.4), width=Inches(1.2))
if os.path.exists(LOGO_AMTICS):
    slide27.shapes.add_picture(LOGO_AMTICS, Inches(11.3), Inches(0.4), width=Inches(1.2))

ty_box = slide27.shapes.add_textbox(Inches(2.0), Inches(2.2), Inches(9.333), Inches(2.0))
tf = ty_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Thank You"
r.font.size = Pt(50)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

# Crimson red divider line beneath Thank You
line_ty = slide27.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.1), Inches(3.6), Inches(3.1), Inches(0.05))
line_ty.fill.solid(); line_ty.fill.fore_color.rgb = COLOR_RED; line_ty.line.fill.background()

cred_box = slide27.shapes.add_textbox(Inches(2.0), Inches(4.3), Inches(9.333), Inches(1.8))
tf = cred_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Prathmesh Chaudhari (202503103510061)"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = "DJ & Electronic Equipment Repair Showcase & Management Web Application — Vikash Electronics"
r2.font.size = Pt(16)
r2.font.color.rgb = COLOR_GRAY_DARK
r2.font.name = "Arial"
p2.space_before = Pt(8)

# Save Presentation to all target filenames
prs.save(PPTX_FILE_PRIMARY)
print(f"[1/3] Successfully generated all 27 slides at: {PPTX_FILE_PRIMARY}")

for path in [PPTX_FILE, PPTX_ACADEMIC_FILE, PPTX_COMPLETE_FILE]:
    try:
        prs.save(path)
        print(f"  -> Also updated: {os.path.basename(path)}")
    except Exception as e:
        print(f"  -> Note: {os.path.basename(path)} is currently open in PowerPoint: {e}")


# ==============================================================================
# PART 2: GENERATE PROJECT – 1 APPLICATION FORM (WORD .DOCX & HTML)
# ==============================================================================
doc = Document()

# Set standard page margins (0.75 in)
for section in doc.sections:
    section.top_margin = DocxInches(0.75)
    section.bottom_margin = DocxInches(0.75)
    section.left_margin = DocxInches(0.75)
    section.right_margin = DocxInches(0.75)

def set_cell_background(cell, fill_color_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

# --- Institute Header ---
p_inst = doc.add_paragraph()
p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p_inst.add_run("UKA TARSADIA UNIVERSITY\n")
r.font.name = "Arial"; r.font.size = DocxPt(16); r.font.bold = True; r.font.color.rgb = DocxRGBColor(15, 23, 42)

r2 = p_inst.add_run("Asha M. Tarsadia Institute of Computer Science and Technology (AMTICS)\n")
r2.font.name = "Arial"; r2.font.size = DocxPt(13); r2.font.bold = True; r2.font.color.rgb = DocxRGBColor(15, 76, 129)

r3 = p_inst.add_run("Department of Computer Science and Engineering\n")
r3.font.name = "Arial"; r3.font.size = DocxPt(12); r3.font.bold = True; r3.font.color.rgb = DocxRGBColor(51, 65, 85)

r4 = p_inst.add_run("ACADEMIC YEAR: 2025 – 2026 | B.TECH CSE (SEMESTER – VII)\n")
r4.font.name = "Arial"; r4.font.size = DocxPt(10.5); r4.font.bold = True; r4.font.color.rgb = DocxRGBColor(100, 116, 139)

# Document Title Box
p_title = doc.add_paragraph()
p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_t = p_title.add_run("PROJECT – 1 APPLICATION FORM (IDP / UDP)")
r_t.font.name = "Arial"; r_t.font.size = DocxPt(14); r_t.font.bold = True; r_t.font.color.rgb = DocxRGBColor(220, 38, 38)
p_title.paragraph_format.space_before = DocxPt(4)
p_title.paragraph_format.space_after = DocxPt(12)

# --- Table 1: Basic Project & Student Particulars ---
tbl1 = doc.add_table(rows=8, cols=2)
tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl1.autofit = False

tbl1_data = [
    ("Project Title", "Vikash Electronics – Web-Based DJ & Electronic Equipment Repair Management and Showcase Portal"),
    ("Project Type", "[X] Industry Defined Project (IDP)     [  ] User Defined Project (UDP)"),
    ("Student Name", "Prathmesh Chaudhari"),
    ("Enrollment Number", "202503103510061"),
    ("Degree / Department", "B.Tech in Computer Science and Engineering (AMTICS, UTU)"),
    ("Internal Project Guide", "Prof. <Guide Name> (Assistant Professor, Dept. of CSE, AMTICS)"),
    ("External Industry / Client", "Vikash Electronics (Proprietor: Mr. Murlidhar Chaudhari)"),
    ("Industry Address & Contact", "Plot No. 199, Sarve No. 1-2, Sanjay Nagar, Udhna Yard, Limbayat / Udhna, Surat - 394210\nDirect Helpline: +91 98254 85520")
]

for row_idx, (label, val) in enumerate(tbl1_data):
    row = tbl1.rows[row_idx]
    c0, c1 = row.cells[0], row.cells[1]
    c0.width = DocxInches(2.4)
    c1.width = DocxInches(4.8)
    
    set_cell_background(c0, "F1F5F9")
    set_cell_margins(c0, 100, 100, 140, 140)
    set_cell_margins(c1, 100, 100, 140, 140)
    
    p0 = c0.paragraphs[0]; p0.paragraph_format.space_after = DocxPt(0)
    r0 = p0.add_run(label); r0.font.name = "Arial"; r0.font.size = DocxPt(10); r0.font.bold = True
    
    p1 = c1.paragraphs[0]; p1.paragraph_format.space_after = DocxPt(0)
    r1 = p1.add_run(val); r1.font.name = "Arial"; r1.font.size = DocxPt(10)

doc.add_paragraph().paragraph_format.space_after = DocxPt(8)

# --- Section 2: Project Abstract & Problem Definition ---
h2 = doc.add_paragraph()
r_h2 = h2.add_run("1. Project Abstract & Problem Definition")
r_h2.font.name = "Arial"; r_h2.font.size = DocxPt(12); r_h2.font.bold = True; r_h2.font.color.rgb = DocxRGBColor(15, 76, 129)
h2.paragraph_format.space_after = DocxPt(4)

p_abs = doc.add_paragraph()
p_abs.paragraph_format.space_after = DocxPt(8)
r = p_abs.add_run(
    "Vikash Electronics is a well-established professional DJ sound system, stage lighting, and pro-audio electronic equipment repair center operating in Surat. "
    "The business specializes in precision bench testing, voice coil rewinding, amplifier output transistor balancing, Sharpy moving head stage lighting optical repair, and digital sound console troubleshooting. "
    "Historically, the workshop depended strictly on walk-in visits and fragmented third-party local directories (e.g. Justdial/Indiamart), which impose heavy listing commissions, route distorted leads, and fail to exhibit genuine technician credentials.\n\n"
    "The proposed project is a high-performance, self-owned web application built exclusively for Vikash Electronics. "
    "The portal provides an interactive showcase of repair capabilities across Audio, Stage Lighting, Visual, and General categories, allows customers (DJs, sound vendors, event planners) to book in-store diagnostic appointments, triggers direct zero-commission WhatsApp consultations with pre-filled fault parameters, and provides GPS turn-by-turn navigation to the Limbayat workshop."
)
r.font.name = "Arial"; r.font.size = DocxPt(10)

# --- Section 3: Project Objectives ---
h3 = doc.add_paragraph()
r_h3 = h3.add_run("2. Project Objectives")
r_h3.font.name = "Arial"; r_h3.font.size = DocxPt(12); r_h3.font.bold = True; r_h3.font.color.rgb = DocxRGBColor(15, 76, 129)
h3.paragraph_format.space_after = DocxPt(4)

objs = [
    "To design and deploy a modern, responsive web application for Vikash Electronics to eliminate third-party middleman aggregator commissions.",
    "To implement an interactive equipment repair catalog categorized by Audio, Stage Lighting, Visual Equipment, and General Electronics.",
    "To incorporate a direct emergency contact panel enabling one-tap native mobile phone calling (+91 98254 85520) and pre-filled WhatsApp fault inquiries.",
    "To provide workshop geolocation transparency through embedded interactive Google Maps routing to Plot 199, Sanjay Nagar, Udhna Yard, Surat.",
    "To establish technical authority and customer trust by featuring authentic workshop bench testing photos and fault diagnostic guides.",
    "To integrate Schema.org LocalBusiness structured data and search-engine optimized metadata to rank high on Google for local Surat electronic repairs."
]

for obj in objs:
    p_obj = doc.add_paragraph(style='List Bullet')
    p_obj.paragraph_format.space_after = DocxPt(2)
    r = p_obj.add_run(obj)
    r.font.name = "Arial"; r.font.size = DocxPt(10)

doc.add_paragraph().paragraph_format.space_after = DocxPt(6)

# --- Section 4: Functional Modules & UML Modeling ---
h4 = doc.add_paragraph()
r_h4 = h4.add_run("3. System Modules & UML Modeling Specifications")
r_h4.font.name = "Arial"; r_h4.font.size = DocxPt(12); r_h4.font.bold = True; r_h4.font.color.rgb = DocxRGBColor(15, 76, 129)
h4.paragraph_format.space_after = DocxPt(4)

tbl_mod = doc.add_table(rows=6, cols=3)
tbl_mod.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_mod.autofit = False

headers = ["Module / Diagram", "Description & Scope", "Deliverable Type"]
for i, h in enumerate(headers):
    c = tbl_mod.cell(0, i)
    set_cell_background(c, "E6CCCC") # Matching Rosy Pink from template
    set_cell_margins(c, 100, 100, 120, 120)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h); r.font.name = "Arial"; r.font.size = DocxPt(9.5); r.font.bold = True

modules_data = [
    ("FR-1: Service Catalog Showcase", "Category-based filtering (Audio, Lighting, Visual, General) with fault symptoms and diagnosis specs.", "Frontend UI Component"),
    ("FR-2: Direct Communication Bridge", "One-touch native dialer and WhatsApp API deep-linking with auto-filled inquiry syntax.", "Integration API"),
    ("FR-3: Geolocation & Navigation", "Embedded Google Maps iframe centered on Limbayat workshop with GPS routing.", "Location Service"),
    ("UML Diagrams (6 Diagrams)", "Complete UML suite: Use Case, Class, Sequence, Activity, State, and Deployment Diagrams.", "System Architecture"),
    ("Local SEO & Performance", "Schema.org LocalBusiness JSON-LD markup, OpenGraph cards, Next.js server-side optimization.", "SEO & Deployment")
]

for row_idx, (m_name, m_desc, m_type) in enumerate(modules_data, start=1):
    row = tbl_mod.rows[row_idx]
    row.cells[0].width = DocxInches(2.2)
    row.cells[1].width = DocxInches(3.4)
    row.cells[2].width = DocxInches(1.6)
    
    set_cell_margins(row.cells[0], 80, 80, 100, 100)
    set_cell_margins(row.cells[1], 80, 80, 100, 100)
    set_cell_margins(row.cells[2], 80, 80, 100, 100)
    
    p0 = row.cells[0].paragraphs[0]; r0 = p0.add_run(m_name); r0.font.name = "Arial"; r0.font.size = DocxPt(9); r0.font.bold = True
    p1 = row.cells[1].paragraphs[0]; r1 = p1.add_run(m_desc); r1.font.name = "Arial"; r1.font.size = DocxPt(9)
    p2 = row.cells[2].paragraphs[0]; r2 = p2.add_run(m_type); r2.font.name = "Arial"; r2.font.size = DocxPt(9)

doc.add_paragraph().paragraph_format.space_after = DocxPt(8)

# --- Section 5: Hardware & Software Requirements ---
h5 = doc.add_paragraph()
r_h5 = h5.add_run("4. Tools, Technologies & Environmental Requirements")
r_h5.font.name = "Arial"; r_h5.font.size = DocxPt(12); r_h5.font.bold = True; r_h5.font.color.rgb = DocxRGBColor(15, 76, 129)
h5.paragraph_format.space_after = DocxPt(4)

tech_specs = [
    ("Frontend Framework", "Next.js 14 / React 18, TypeScript for robust static type checking"),
    ("Styling & Icons", "Tailwind CSS (utility-first dark slate palette), Lucide-React vector icon set"),
    ("State & Data Architecture", "Centralized TypeScript data stores (business.ts, services.ts), SSR / SSG"),
    ("Gateways & External APIs", "WhatsApp Business Deep-link API (wa.me), Google Maps Embed API, RFC 3966 tel protocol"),
    ("Development Methodology", "Agile (Scrum) Sprint Cycle: Plan -> Design -> Develop -> Test -> Review"),
    ("Client Hosting Environment", "Vercel Edge Network / Node.js Runtime with HTTPS SSL encryption")
]

tbl_tech = doc.add_table(rows=len(tech_specs), cols=2)
tbl_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
for idx, (k, v) in enumerate(tech_specs):
    row = tbl_tech.rows[idx]
    row.cells[0].width = DocxInches(2.5)
    row.cells[1].width = DocxInches(4.7)
    set_cell_background(row.cells[0], "F8FAFC")
    set_cell_margins(row.cells[0], 60, 60, 100, 100)
    set_cell_margins(row.cells[1], 60, 60, 100, 100)
    
    p0 = row.cells[0].paragraphs[0]; r0 = p0.add_run(k); r0.font.name = "Arial"; r0.font.size = DocxPt(9); r0.font.bold = True
    p1 = row.cells[1].paragraphs[0]; r1 = p1.add_run(v); r1.font.name = "Arial"; r1.font.size = DocxPt(9)

doc.add_paragraph().paragraph_format.space_after = DocxPt(16)

# --- Section 6: Endorsement & Signature Block ---
h6 = doc.add_paragraph()
r_h6 = h6.add_run("5. Signatures & Academic / Industry Endorsements")
r_h6.font.name = "Arial"; r_h6.font.size = DocxPt(12); r_h6.font.bold = True; r_h6.font.color.rgb = DocxRGBColor(15, 76, 129)
h6.paragraph_format.space_after = DocxPt(8)

tbl_sign = doc.add_table(rows=2, cols=4)
tbl_sign.alignment = WD_TABLE_ALIGNMENT.CENTER

sig_headers = ["Student Signature", "Industry Mentor / Client", "Internal Project Guide", "Project Coordinator / HOD"]
for i, h in enumerate(sig_headers):
    c = tbl_sign.cell(0, i)
    set_cell_background(c, "F1F5F9")
    set_cell_margins(c, 80, 80, 100, 100)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(h); r.font.name = "Arial"; r.font.size = DocxPt(9); r.font.bold = True

sig_bodies = [
    "\n\n\n_______________________\nPrathmesh Chaudhari\n(202503103510061)\nDate: ____/____/2026",
    "\n\n\n_______________________\nMr. Murlidhar Chaudhari\nProprietor, Vikash Electronics\nSeal & Signature",
    "\n\n\n_______________________\nProf. <Guide Name>\nAssistant Professor, CSE\nAMTICS, UTU",
    "\n\n\n_______________________\nDr. / Prof. [HOD Name]\nHead of Department, CSE\nAMTICS, UTU"
]

for i, b in enumerate(sig_bodies):
    c = tbl_sign.cell(1, i)
    c.width = DocxInches(1.8)
    set_cell_margins(c, 100, 100, 80, 80)
    p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(b); r.font.name = "Arial"; r.font.size = DocxPt(8.5); r.font.color.rgb = DocxRGBColor(51, 65, 85)

doc.save(DOCX_FILE)
print(f"[2/3] Successfully generated Project – 1 Application Form (Word) at: {DOCX_FILE}")


# ==============================================================================
# PART 3: GENERATE PRINTABLE HTML APPLICATION FORM
# ==============================================================================
html_content = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Project – 1 Application Form | Vikash Electronics</title>
    <style>
        @page {{
            size: A4;
            margin: 1.5cm;
        }}
        body {{
            font-family: 'Segoe UI', Arial, sans-serif;
            color: #0F172A;
            line-height: 1.5;
            background: #fff;
            max-width: 850px;
            margin: 0 auto;
            padding: 20px;
        }}
        .header {{
            text-align: center;
            border-bottom: 2px solid #0F4C81;
            padding-bottom: 12px;
            margin-bottom: 20px;
        }}
        .header h1 {{
            font-size: 20px;
            margin: 0;
            color: #0F172A;
            text-transform: uppercase;
            letter-spacing: 0.5px;
        }}
        .header h2 {{
            font-size: 15px;
            margin: 4px 0;
            color: #0F4C81;
        }}
        .header h3 {{
            font-size: 13px;
            margin: 2px 0;
            color: #334155;
            font-weight: 600;
        }}
        .form-badge {{
            display: inline-block;
            background: #FEE2E2;
            color: #DC2626;
            font-weight: bold;
            font-size: 14px;
            padding: 4px 16px;
            border-radius: 4px;
            margin-top: 8px;
            border: 1px solid #FCA5A5;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin: 12px 0 20px 0;
            font-size: 12.5px;
        }}
        th, td {{
            border: 1px solid #CBD5E1;
            padding: 8px 10px;
            text-align: left;
            vertical-align: top;
        }}
        th {{
            background: #E6CCCC;
            color: #0F172A;
            font-weight: bold;
        }}
        td.label {{
            background: #F1F5F9;
            font-weight: bold;
            width: 28%;
        }}
        h4.section-title {{
            font-size: 14px;
            color: #0F4C81;
            margin: 18px 0 6px 0;
            border-left: 4px solid #0F4C81;
            padding-left: 8px;
        }}
        p {{
            font-size: 12.5px;
            margin: 6px 0;
            text-align: justify;
        }}
        ul {{
            margin: 6px 0 14px 20px;
            padding: 0;
            font-size: 12.5px;
        }}
        li {{
            margin-bottom: 4px;
        }}
        .signatures-table td {{
            text-align: center;
            height: 90px;
            vertical-align: bottom;
            font-size: 11px;
            background: #FAFAFA;
        }}
        .signatures-table th {{
            background: #F1F5F9;
            text-align: center;
            font-size: 11.5px;
        }}
        @media print {{
            body {{ padding: 0; }}
            .no-print {{ display: none; }}
        }}
    </style>
</head>
<body>
    <div class="header">
        <h1>UKA TARSADIA UNIVERSITY</h1>
        <h2>Asha M. Tarsadia Institute of Computer Science and Technology (AMTICS)</h2>
        <h3>Department of Computer Science and Engineering</h3>
        <p style="text-align: center; margin: 4px 0; font-size: 11.5px; color: #64748B;">ACADEMIC YEAR: 2025 – 2026 | B.TECH CSE (SEMESTER – VII)</p>
        <div class="form-badge">PROJECT – 1 APPLICATION FORM (IDP / UDP)</div>
    </div>

    <h4 class="section-title">1. Project & Candidate Particulars</h4>
    <table>
        <tr>
            <td class="label">Project Title</td>
            <td><strong>Vikash Electronics – Web-Based DJ & Electronic Equipment Repair Management and Showcase Portal</strong></td>
        </tr>
        <tr>
            <td class="label">Project Classification</td>
            <td><strong>[✓] Industry Defined Project (IDP)</strong> &nbsp;&nbsp;&nbsp;&nbsp; [ ] User Defined Project (UDP)</td>
        </tr>
        <tr>
            <td class="label">Student Name</td>
            <td><strong>Prathmesh Chaudhari</strong></td>
        </tr>
        <tr>
            <td class="label">Enrollment Number</td>
            <td><strong>202503103510061</strong></td>
        </tr>
        <tr>
            <td class="label">Degree & Semester</td>
            <td>B.Tech in Computer Science and Engineering (7th Semester)</td>
        </tr>
        <tr>
            <td class="label">Internal Faculty Guide</td>
            <td>Prof. &lt;Guide Name&gt; (Assistant Professor, Dept. of CSE, AMTICS)</td>
        </tr>
        <tr>
            <td class="label">Industry / Client Name</td>
            <td><strong>Vikash Electronics</strong> (Proprietor: Mr. Murlidhar Chaudhari)</td>
        </tr>
        <tr>
            <td class="label">Workshop Address & Phone</td>
            <td>Plot No. 199, Sarve No. 1-2, Sanjay Nagar, Udhna Yard, Limbayat / Udhna, Surat - 394210<br>Contact: +91 98254 85520 | Email: contact@vikashelectronics.in</td>
        </tr>
    </table>

    <h4 class="section-title">2. Project Abstract & Problem Definition</h4>
    <p>
        Vikash Electronics is a well-established professional DJ sound system, stage lighting, and pro-audio electronic equipment repair center operating in Surat. The business specializes in precision bench testing, voice coil rewinding, amplifier output transistor balancing, Sharpy moving head stage lighting optical repair, and digital sound console troubleshooting. Historically, the workshop depended strictly on walk-in visits and fragmented third-party local directories (e.g. Justdial/Indiamart), which impose heavy listing commissions, route distorted leads, and fail to exhibit genuine technician credentials.
    </p>
    <p>
        The proposed project is a high-performance, self-owned web application built exclusively for Vikash Electronics. The portal provides an interactive showcase of repair capabilities across Audio, Stage Lighting, Visual, and General categories, allows customers (DJs, sound vendors, event planners) to book in-store diagnostic appointments, triggers direct zero-commission WhatsApp consultations with pre-filled fault parameters, and provides GPS turn-by-turn navigation to the Limbayat workshop.
    </p>

    <h4 class="section-title">3. Project Objectives</h4>
    <ul>
        <li>To design and deploy a modern, responsive web application for Vikash Electronics to eliminate third-party middleman aggregator commissions.</li>
        <li>To implement an interactive equipment repair catalog categorized by Audio, Stage Lighting, Visual Equipment, and General Electronics.</li>
        <li>To incorporate a direct emergency contact panel enabling one-tap native mobile phone calling (+91 98254 85520) and pre-filled WhatsApp fault inquiries.</li>
        <li>To provide workshop geolocation transparency through embedded interactive Google Maps routing to Plot 199, Sanjay Nagar, Udhna Yard, Surat.</li>
        <li>To establish technical authority and customer trust by featuring authentic workshop bench testing photos and fault diagnostic guides.</li>
        <li>To integrate Schema.org LocalBusiness structured data and search-engine optimized metadata to rank high on Google for local Surat electronic repairs.</li>
    </ul>

    <h4 class="section-title">4. System Modules & UML Diagrams Scope</h4>
    <table>
        <tr>
            <th style="width: 30%;">Module / UML Diagram</th>
            <th style="width: 50%;">Description & Functional Coverage</th>
            <th style="width: 20%;">Deliverable Type</th>
        </tr>
        <tr>
            <td><strong>FR-1: Service Catalog Showcase</strong></td>
            <td>Multi-category filtering (Audio, Lighting, Visual, General) with fault symptoms and diagnosis specs.</td>
            <td>Frontend Component</td>
        </tr>
        <tr>
            <td><strong>FR-2: Direct Communication Bridge</strong></td>
            <td>One-touch native dialer and WhatsApp API deep-linking with auto-filled inquiry syntax.</td>
            <td>Integration API</td>
        </tr>
        <tr>
            <td><strong>FR-3: Geolocation & Navigation</strong></td>
            <td>Embedded Google Maps iframe centered on Limbayat workshop with GPS routing.</td>
            <td>Location Service</td>
        </tr>
        <tr>
            <td><strong>6 UML Diagrams Suite</strong></td>
            <td>Use Case Diagram, Class Diagram, Sequence Diagram, Activity Diagram, State Diagram, Deployment Diagram.</td>
            <td>System Architecture</td>
        </tr>
        <tr>
            <td><strong>Local SEO & SSR Optimization</strong></td>
            <td>Schema.org LocalBusiness JSON-LD markup, OpenGraph cards, Next.js server-side optimization.</td>
            <td>SEO & Performance</td>
        </tr>
    </table>

    <h4 class="section-title">5. Technical Architecture & Environment</h4>
    <table>
        <tr>
            <td class="label">Frontend Framework</td>
            <td>Next.js 14 / React 18, TypeScript for robust static type safety</td>
        </tr>
        <tr>
            <td class="label">Styling & Icons</td>
            <td>Tailwind CSS (utility-first dark slate palette), Lucide-React vector icon set</td>
        </tr>
        <tr>
            <td class="label">State & Data Store</td>
            <td>Centralized TypeScript data stores (business.ts, services.ts), SSR / SSG</td>
        </tr>
        <tr>
            <td class="label">APIs & Protocols</td>
            <td>WhatsApp Business API (wa.me), Google Maps Embed API, RFC 3966 tel protocol</td>
        </tr>
        <tr>
            <td class="label">Software Process Model</td>
            <td>Agile (Scrum) Sprint Cycle: Plan → Design → Develop → Test → Review</td>
        </tr>
    </table>

    <h4 class="section-title">6. Signatures & Academic / Industry Endorsements</h4>
    <table class="signatures-table">
        <tr>
            <th style="width: 25%;">Student Candidate</th>
            <th style="width: 25%;">Industry Mentor / Client</th>
            <th style="width: 25%;">Internal Guide</th>
            <th style="width: 25%;">Project Coordinator / HOD</th>
        </tr>
        <tr>
            <td>
                <br><br>
                _______________________<br>
                <strong>Prathmesh Chaudhari</strong><br>
                (202503103510061)<br>
                Date: ____/____/2026
            </td>
            <td>
                <br><br>
                _______________________<br>
                <strong>Mr. Murlidhar Chaudhari</strong><br>
                Proprietor, Vikash Electronics<br>
                Seal & Signature
            </td>
            <td>
                <br><br>
                _______________________<br>
                <strong>Prof. &lt;Guide Name&gt;</strong><br>
                Assistant Professor, CSE<br>
                AMTICS, UTU
            </td>
            <td>
                <br><br>
                _______________________<br>
                <strong>Dr. / Prof. [HOD Name]</strong><br>
                Head of Department, CSE<br>
                AMTICS, UTU
            </td>
        </tr>
    </table>
</body>
</html>
"""

with open(HTML_FILE, "w", encoding="utf-8") as f:
    f.write(html_content)
print(f"[3/3] Successfully generated printable HTML Application Form at: {HTML_FILE}")
