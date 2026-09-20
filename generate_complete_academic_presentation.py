import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE

# --- COLOR PALETTE MATCHING ACADEMIC REFERENCE SLIDES ---
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_BLACK = RGBColor(15, 23, 42)        # Black / Deep Slate
COLOR_RED = RGBColor(220, 38, 38)         # #DC2626 Crimson Red (for subtitles in parentheses & special boxes)
COLOR_BLUE = RGBColor(2, 132, 199)        # #0284C7 Blue for Title
COLOR_DARK_BLUE = RGBColor(15, 76, 129)   # Department Subtitle
COLOR_GRAY_DARK = RGBColor(51, 65, 85)    # Slate 700
COLOR_GRAY_LIGHT = RGBColor(248, 250, 252)# Slate 50
COLOR_TABLE_HEADER = RGBColor(230, 204, 204) # #E6CCCC Rosy Pink Table Header from template
COLOR_TABLE_BORDER = RGBColor(203, 213, 225) # Slate 300
COLOR_DIAGRAM_BG = RGBColor(240, 249, 255)   # Light Sky Blue Tint for boxes
COLOR_DIAGRAM_BORDER = RGBColor(14, 116, 144)# Cyan/Teal border for diagrams
COLOR_ACCENT_BLUE = RGBColor(37, 99, 235)    # Arrow & Flow blue

BASE_DIR = r"d:\A_S_Projects\Vikas_Electronics"
OUTPUT_FILE = os.path.join(BASE_DIR, "Vikash_Electronics_Complete_Academic_Presentation.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

LOGO_UTU = os.path.join(BASE_DIR, "utu_flame_logo.png")
LOGO_AMTICS = os.path.join(BASE_DIR, "amtics_circle_logo.png")

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_WHITE

def add_academic_header(slide, title_text, red_subtitle=""):
    """Adds standard academic heading matching the PDF template"""
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.3))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0)
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r_title = p.add_run()
    r_title.text = title_text
    r_title.font.size = Pt(28)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_BLACK
    r_title.font.name = "Arial"
    
    if red_subtitle:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r_sub = p2.add_run()
        r_sub.text = red_subtitle
        r_sub.font.size = Pt(22)
        r_sub.font.bold = True
        r_sub.font.color.rgb = COLOR_RED
        r_sub.font.name = "Arial"
        p2.space_before = Pt(3)

def add_diagram_box(slide, left, top, width, height, text_title, text_sub="", bg_col=COLOR_DIAGRAM_BG, border_col=COLOR_DIAGRAM_BORDER, text_col=COLOR_BLACK):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_col
    shape.line.color.rgb = border_col
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text_title
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = text_col
    r.font.name = "Arial"
    
    if text_sub:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run()
        r2.text = text_sub
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = COLOR_GRAY_DARK
        r2.font.name = "Arial"
        p2.space_before = Pt(2)
    return shape

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)

if os.path.exists(LOGO_UTU):
    slide1.shapes.add_picture(LOGO_UTU, Inches(0.8), Inches(0.4), width=Inches(1.2))
if os.path.exists(LOGO_AMTICS):
    slide1.shapes.add_picture(LOGO_AMTICS, Inches(11.3), Inches(0.4), width=Inches(1.2))

header_box = slide1.shapes.add_textbox(Inches(2.2), Inches(0.4), Inches(8.933), Inches(1.5))
tf = header_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Asha M. Tarsadia Institute of\nComputer Science and Technology"
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p_dept = tf.add_paragraph()
p_dept.alignment = PP_ALIGN.CENTER
r_dept = p_dept.add_run()
r_dept.text = "Computer Science and Engineering"
r_dept.font.size = Pt(20)
r_dept.font.bold = True
r_dept.font.color.rgb = COLOR_DARK_BLUE
r_dept.font.name = "Arial"
p_dept.space_before = Pt(6)

title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(1.5))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Vikash Electronics – Web-Based DJ & Electronic Equipment Repair Management and Showcase Portal"
r.font.size = Pt(25)
r.font.bold = True
r.font.color.rgb = COLOR_BLUE
r.font.name = "Arial"

prep_box = slide1.shapes.add_textbox(Inches(2.0), Inches(4.1), Inches(9.333), Inches(1.4))
tf = prep_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Prepared by"
r.font.size = Pt(17)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = "Prathmesh Chaudhari"
r2.font.size = Pt(16)
r2.font.color.rgb = COLOR_BLACK
r2.font.name = "Arial"
p2.space_before = Pt(4)

p3 = tf.add_paragraph()
p3.alignment = PP_ALIGN.CENTER
r3 = p3.add_run()
r3.text = "(202503103510061)"
r3.font.size = Pt(15)
r3.font.color.rgb = COLOR_BLACK
r3.font.name = "Arial"
p3.space_before = Pt(2)

guide_box = slide1.shapes.add_textbox(Inches(2.0), Inches(5.7), Inches(9.333), Inches(1.3))
tf = guide_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Guided by"
r.font.size = Pt(17)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = "<Guide Name>"
r2.font.size = Pt(15)
r2.font.color.rgb = COLOR_BLACK
r2.font.name = "Arial"
p2.space_before = Pt(4)


# ==============================================================================
# SLIDE 2: OUTLINE
# ==============================================================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2)
add_academic_header(slide2, "Outline")

content_box = slide2.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.7), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

outline_items = [
    "Client Profile",
    "Project Introduction",
    "Objective(s)",
    "Problem Analysis",
    "Requirement Engineering",
    "Software Development Model",
    "Modeling",
    "Others"
]
for idx, item in enumerate(outline_items):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = f"•  {item}"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(14)


# ==============================================================================
# SLIDE 3: CLIENT PROFILE
# ==============================================================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3)
add_academic_header(slide3, "Client Profile")

content_box = slide3.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

sections = [
    ("Client name and address: Vikash Electronics, a premier specialized DJ & electronic equipment repair workshop operating in Surat, Gujarat", [
        "Proprietor & Chief Technician: Mr. Murlidhar Chaudhari",
        "Workshop Address: Plot No. 199, Sarve No. 1-2, Sanjay Nagar, Near Udhna Yard, Limbayat, Surat, Gujarat - 394210",
        "Direct Contact: +91 98254 85520 | Domain: Pro Audio, Stage Lighting & Moving Heads"
    ]),
    ("Your communication: Requirements gathered through in-person workshop meetings and WhatsApp/calls with the owner, covering repair categories, fault diagnostics, branding, and direct contact needs", [
        "Conducted on-site inspection of live repair benches, diagnostic multimeters, oscilloscopes, and coil winding fixtures.",
        "Demonstrated responsive prototypes across mobile viewports to ensure fast one-tap direct calling."
    ]),
    ("Contract letter: Project undertaken as a client engagement with Vikash Electronics, with scope, photography, and deliverables confirmed with the client before development", [
        "Client authorization received from Mr. Murlidhar Chaudhari endorsing the digital showcase and web launch."
    ])
]

for s_idx, (sec_title, points) in enumerate(sections):
    p_sec = tf.paragraphs[0] if s_idx == 0 else tf.add_paragraph()
    r_sec = p_sec.add_run()
    r_sec.text = sec_title
    r_sec.font.size = Pt(15)
    r_sec.font.bold = True
    r_sec.font.color.rgb = COLOR_BLACK
    r_sec.font.name = "Arial"
    p_sec.space_before = Pt(16 if s_idx > 0 else 0)

    for pt in points:
        p_pt = tf.add_paragraph()
        r_pt = p_pt.add_run()
        r_pt.text = f"    • {pt}"
        r_pt.font.size = Pt(11.5)
        r_pt.font.color.rgb = COLOR_GRAY_DARK
        r_pt.font.name = "Arial"
        p_pt.space_before = Pt(3)


# ==============================================================================
# SLIDE 4: PROJECT INTRODUCTION
# ==============================================================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4)
add_academic_header(slide4, "Project Introduction")

content_box = slide4.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

intro_paragraphs = [
    "Vikash Electronics is a specialized electronic repair workshop that currently depends on offline word-of-mouth and localized walk-in customers across Surat.",
    "This project is a dedicated web-based electronic repair portal built exclusively for Vikash Electronics, listing its comprehensive repair capabilities with fault descriptions, turnaround details, and authentic workshop bench imagery.",
    "Customers (DJs, sound rental providers, event organizers, and consumers) can browse repair categories (Speakers, Power Amplifiers, Sharpy Moving Head Lights, Sound Consoles, and TVs) and contact the technician directly through one-click phone calls or WhatsApp without third-party commission middleman apps.",
    "The site also features the workshop's verified physical address in Limbayat, interactive Google Maps routing, operating hours, and schema-structured local SEO.",
    "The goal is to provide a self-owned authoritative digital presence for Vikash Electronics, eliminate reliance on fragmented generic business directories, and streamline emergency repair inquiries."
]

for idx, para in enumerate(intro_paragraphs):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = para
    r.font.size = Pt(14)
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(14)


# ==============================================================================
# SLIDE 5: OBJECTIVE(S)
# ==============================================================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5)
add_academic_header(slide5, "Objective(s)")

content_box = slide5.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

objectives_list = [
    "Build a single, self-owned website where customers can view the full spectrum of electronic and DJ equipment repair services with descriptions and technical highlights",
    "Allow customers to initiate direct phone inquiries and instant WhatsApp consultations with the chief technician directly through the website",
    "Reduce reliance on third-party aggregators and directories (e.g. Justdial/Indiamart) and eliminate commission charges or ad clutter",
    "Display authentic business information — verified workshop bench testing images, Limbayat location, contact numbers, and operating hours — in one place",
    "Provide technical credibility by highlighting component-level diagnostics (voice coil rewinding, Sharpy optical alignment, transistor pair matching) performed by Murlidhar Chaudhari",
    "Improve customer discovery and business inquiries across Surat and South Gujarat by embedding structured Schema.org LocalBusiness SEO data"
]

for idx, item in enumerate(objectives_list):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = item
    r.font.size = Pt(13.5)
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(14)


# ==============================================================================
# SLIDE 6: SERVICE PROCESS FLOW (FLOWCHART DIAGRAM)
# ==============================================================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6)
add_academic_header(slide6, "Service & Repair Process Flow")

flow_steps = [
    ("Browse\nServices", "Audio, Lighting,\nTV & Visual"),
    ("Select\nEquipment", "Speakers, Amps,\nSharpy Lights"),
    ("Inspect\nDiagnostics", "Fault Symptoms &\nComponent Fixes"),
    ("Initiate\nContact", "One-Tap Call or\nWhatsApp Chat"),
    ("Workshop\nIntake", "Physical Drop-off\n& Bench Check"),
    ("Precision\nRepair", "Component Fix &\nStress Load Test"),
    ("Stage Ready\nDelivery", "Verified Testing &\nClient Handover")
]

step_width = Inches(1.4)
step_height = Inches(1.6)
spacing = Inches(0.3)
start_x = Inches(0.8)
y_pos = Inches(2.6)

for idx, (title, sub) in enumerate(flow_steps):
    x = start_x + idx * (step_width + spacing)
    add_diagram_box(slide6, x, y_pos, step_width, step_height, title, sub)
    
    # Arrow to next step
    if idx < len(flow_steps) - 1:
        arrow_x = x + step_width + Inches(0.04)
        arrow = slide6.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arrow_x, y_pos + Inches(0.65), Inches(0.22), Inches(0.3))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = COLOR_ACCENT_BLUE
        arrow.line.fill.background()

# Bottom explanation banner
foot_box = slide6.shapes.add_textbox(Inches(0.8), Inches(5.2), Inches(11.733), Inches(1.2))
tf = foot_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Customers move through the entire journey on Vikash Electronics's own website — no third-party aggregator hand-off, direct technician consultation, and zero commission paid at any step."
r.font.size = Pt(14)
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"


# ==============================================================================
# SLIDE 7: PROBLEM ANALYSIS (Understand the problem)
# ==============================================================================
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7)
add_academic_header(slide7, "Problem Analysis", "(Understand the problem)")

content_box = slide7.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_understand = [
    ("I. Who are the stockholders in the solution to the problem?",
     "The workshop owner & chief technician (Mr. Murlidhar Chaudhari), customers needing urgent repairs (event managers, DJs, sound rental vendors, wedding planners), and the web portal administrator"),
     
    ("II. What are the unknowns?",
     "The dynamic variety of hardware circuit topologies across imported and local brands, component replacement availability, fluctuating spare parts costs, and how repair status will be communicated in real time"),
     
    ("III. Can the problem be compartmentalized?",
     "Yes — into modules: equipment service catalog, fault diagnostic showcase, direct telephonic/WhatsApp communication bridge, workshop location navigation, and local search engine optimization"),
     
    ("IV. Can the problem be represented graphically?",
     "Yes, the flow can be represented through flowcharts, system architecture diagrams, and UML diagrams covering equipment browsing, inquiry routing, and repair lifecycle")
]

for idx, (q, a) in enumerate(pa_understand):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(14)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = a
    r_a.font.size = Pt(12)
    r_a.font.color.rgb = COLOR_GRAY_DARK
    r_a.font.name = "Arial"
    p.space_before = Pt(12)


# ==============================================================================
# SLIDE 8: PROBLEM ANALYSIS (Plan the solution)
# ==============================================================================
slide8 = prs.slides.add_slide(blank_layout)
set_slide_background(slide8)
add_academic_header(slide8, "Problem Analysis", "(Plan the solution)")

content_box = slide8.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_plan = [
    ("I. Have you seen similar problem before?",
     "Yes, in generic local service platforms, Justdial, Indiamart, and various electronic repair shop portals"),
     
    ("II. Has a similar problem been solved?",
     "Yes — many specialized businesses now run their own direct-discovery websites to establish authoritative branding and avoid aggregator fees and delays"),
     
    ("III. Can sub-problems be defined?",
     "Yes: service catalog presentation, technical fault breakdown, direct contact routing (Call/WhatsApp), workshop location navigation, and search engine discoverability"),
     
    ("IV. Can you represent a solution in a manner that leads to effective implementation?",
     "Yes, through a modular Next.js architecture diagram, component hierarchy, and UML modeling that leads to effective, phase-wise implementation")
]

for idx, (q, a) in enumerate(pa_plan):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(14)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = a
    r_a.font.size = Pt(12)
    r_a.font.color.rgb = COLOR_GRAY_DARK
    r_a.font.name = "Arial"
    p.space_before = Pt(12)


# ==============================================================================
# SLIDE 9: SYSTEM ARCHITECTURE (DIAGRAM)
# ==============================================================================
slide9 = prs.slides.add_slide(blank_layout)
set_slide_background(slide9)
add_academic_header(slide9, "System Architecture")

# Box 1: Customer (Top)
add_diagram_box(slide9, Inches(4.8), Inches(1.7), Inches(3.7), Inches(0.8), "Customer", "(Web / Mobile Browser)")

# Down arrow
arr1 = slide9.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.5), Inches(2.55), Inches(0.3), Inches(0.35))
arr1.fill.solid(); arr1.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr1.line.fill.background()

# Box 2: Website Frontend (Next.js)
add_diagram_box(slide9, Inches(4.4), Inches(2.95), Inches(4.5), Inches(0.9), "Website Frontend (Next.js 14)", "(Hero, Service Catalog, About, Contact, Map)")

# Down arrow
arr2 = slide9.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.5), Inches(3.9), Inches(0.3), Inches(0.35))
arr2.fill.solid(); arr2.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr2.line.fill.background()

# Box 3: Application Server / API Layer (Center)
add_diagram_box(slide9, Inches(4.4), Inches(4.3), Inches(4.5), Inches(0.9), "Application Logic / Server", "(Routing, Service Data, Schema.org SEO)")

# Left Box: Data Store
add_diagram_box(slide9, Inches(0.8), Inches(4.3), Inches(3.0), Inches(0.9), "Data Repository", "(business.ts, Repair Specs, Assets)")

# Right Box: External Integrations (WhatsApp & Maps)
add_diagram_box(slide9, Inches(9.5), Inches(4.3), Inches(3.0), Inches(0.9), "Communication Gateway", "(Direct Dialer & WhatsApp API)", bg_col=COLOR_WHITE, border_col=COLOR_RED, text_col=COLOR_RED)

# Horizontal Arrows to Left & Right
arr_left = slide9.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, Inches(3.9), Inches(4.6), Inches(0.4), Inches(0.25))
arr_left.fill.solid(); arr_left.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr_left.line.fill.background()

arr_right = slide9.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(9.0), Inches(4.6), Inches(0.4), Inches(0.25))
arr_right.fill.solid(); arr_right.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr_right.line.fill.background()

# Down arrow to Technician / Admin
arr3 = slide9.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.5), Inches(5.25), Inches(0.3), Inches(0.3))
arr3.fill.solid(); arr3.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr3.line.fill.background()

# Box 4: Technician / Workshop Management (Bottom)
add_diagram_box(slide9, Inches(4.4), Inches(5.6), Inches(4.5), Inches(0.8), "Workshop Management", "(Owner: Inquiry Intake & Repair Processing)")

# Footnote
foot_box = slide9.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.733), Inches(0.6))
tf = foot_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Customer requests flow from the website frontend through Next.js server components to the data layer and WhatsApp gateway; the owner manages services and inquiries directly."
r.font.size = Pt(12)
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"


# ==============================================================================
# SLIDE 10: PROBLEM ANALYSIS (Carry out the plan)
# ==============================================================================
slide10 = prs.slides.add_slide(blank_layout)
set_slide_background(slide10)
add_academic_header(slide10, "Problem Analysis", "(Carry out the plan)")

content_box = slide10.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_carry = [
    ("I. Does the solution conform to the plan?",
     "Yes — development proceeds module by module (Hero section, categorized service cards, about workshop, interactive Google Map, direct call/WhatsApp integration) as per the defined architecture"),
     
    ("II. Is each component part of the solution provably correct?",
     "Yes, each component is tested independently — for example, service card rendering, click-to-call link protocols, WhatsApp pre-populated query strings, and responsive mobile drawer state are each verified before integration")
]

for idx, (q, a) in enumerate(pa_carry):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(14)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = a
    r_a.font.size = Pt(12)
    r_a.font.color.rgb = COLOR_GRAY_DARK
    r_a.font.name = "Arial"
    p.space_before = Pt(14)


# ==============================================================================
# SLIDE 11: PROBLEM ANALYSIS (Examine the results)
# ==============================================================================
slide11 = prs.slides.add_slide(blank_layout)
set_slide_background(slide11)
add_academic_header(slide11, "Problem Analysis", "(Examine the results)")

content_box = slide11.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_examine = [
    ("I. Is it possible to test each component part of the solution?",
     "Yes — local server page loads, interactive category filtering, one-tap mobile calling, WhatsApp URL generation, and Google Maps iframe loading are all verifiable on their own"),
     
    ("II. Does the solution produce results that conform to the data, functions, and features that are required?",
     "Yes — customers can view specialized repair services, inspect genuine workshop bench credentials, and initiate direct repair booking with the technician, all without a third-party app")
]

for idx, (q, a) in enumerate(pa_examine):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(14)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = a
    r_a.font.size = Pt(12)
    r_a.font.color.rgb = COLOR_GRAY_DARK
    r_a.font.name = "Arial"
    p.space_before = Pt(14)


# ==============================================================================
# SLIDE 12: REQUIREMENT ENGINEERING (Client-side communication)
# ==============================================================================
slide12 = prs.slides.add_slide(blank_layout)
set_slide_background(slide12)
add_academic_header(slide12, "Requirement Engineering", "(Client-side communication)")

content_box = slide12.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

req_items = [
    "Complete repair services list with categories (Audio, Lighting, Visual, General), item images, descriptions, and technical highlight tags, matching the pro-audio positioning of the brand",
    "Business details to be displayed: physical workshop address in Limbayat, Surat, direct contact helpline (+91 98254 85520), and operating hours (Mon-Sat 9:30 AM – 9:00 PM)",
    "Instant communication support (Direct Phone Call & WhatsApp chat) that avoids aggregator delays and middleman commissions",
    "A simple way for the technician (Murlidhar Chaudhari) to showcase specialized skills (voice coil rewinding, Sharpy moving head calibration, transistor balancing)",
    "A clean, mobile-friendly design consistent with the pro-audio stage concert and DJ lighting industry brand identity",
    "Easy updates to service details and pricing in a centralized data file without needing complex developer intervention each time"
]

for idx, item in enumerate(req_items):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = item
    r.font.size = Pt(13)
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(12)


# ==============================================================================
# SLIDES 13, 14, 15: REQUIREMENT ENGINEERING (Functional Requirements Tables)
# ==============================================================================
def draw_single_fr_slide(slide, fr_id, title, desc, actor, inputs, outputs, intro_note=False):
    if intro_note:
        note_box = slide.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(11.3), Inches(0.5))
        p_n = note_box.text_frame.paragraphs[0]
        p_n.text = "Enlisted below are the functional requirements for the website:"
        p_n.font.size = Pt(14)
        p_n.font.color.rgb = COLOR_BLACK
        p_n.font.name = "Arial"
        y_top = Inches(2.2)
    else:
        y_top = Inches(2.0)

    rows = 5
    cols = 2
    table_shape = slide.shapes.add_table(rows, cols, Inches(1.0), y_top, Inches(11.3), Inches(3.6))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(9.1)
    
    # Header Row
    c00 = table.cell(0, 0); c00.fill.solid(); c00.fill.fore_color.rgb = COLOR_TABLE_HEADER
    c00.text = fr_id; p = c00.text_frame.paragraphs[0]; p.font.bold = True; p.font.size = Pt(13); p.font.color.rgb = COLOR_BLACK
    
    c01 = table.cell(0, 1); c01.fill.solid(); c01.fill.fore_color.rgb = COLOR_TABLE_HEADER
    c01.text = title; p = c01.text_frame.paragraphs[0]; p.font.bold = True; p.font.size = Pt(13); p.font.color.rgb = COLOR_BLACK
    
    data = [
        ("Description", desc),
        ("Actor(s)", actor),
        ("Input(s)", inputs),
        ("Output(s)", outputs)
    ]
    for r_idx, (lbl, val) in enumerate(data, start=1):
        c0 = table.cell(r_idx, 0); c0.fill.solid(); c0.fill.fore_color.rgb = COLOR_WHITE; c0.text = lbl
        p0 = c0.text_frame.paragraphs[0]; p0.font.bold = True; p0.font.size = Pt(12); p0.font.color.rgb = COLOR_BLACK
        
        c1 = table.cell(r_idx, 1); c1.fill.solid(); c1.fill.fore_color.rgb = COLOR_WHITE; c1.text = val
        p1 = c1.text_frame.paragraphs[0]; p1.font.size = Pt(12); p1.font.color.rgb = COLOR_GRAY_DARK

# Slide 13: FR-1
slide13 = prs.slides.add_slide(blank_layout)
set_slide_background(slide13)
add_academic_header(slide13, "Requirement Engineering", "(Functional Requirements)")
draw_single_fr_slide(
    slide13, "FR-1", "Service Browsing & Equipment Selection",
    "Used to let the customer browse the specialized electronic repair services and inspect fault symptoms.",
    "Customer (DJ, Event Manager, Consumer)",
    "Category selection (Audio, Lighting, Visual), selected equipment item (e.g., Speakers, Amplifiers, Sharpy).",
    "Interactive service cards rendered with technical failure details, diagnostic highlights, and repair specifications.",
    intro_note=True
)

# Slide 14: FR-2
slide14 = prs.slides.add_slide(blank_layout)
set_slide_background(slide14)
add_academic_header(slide14, "Requirement Engineering", "(Functional Requirements)")
draw_single_fr_slide(
    slide14, "FR-2", "Direct Emergency Call & WhatsApp Consultation",
    "Used to let the customer initiate instant telephonic or WhatsApp consultation directly with the technician.",
    "Customer, Chief Technician (Murlidhar Chaudhari)",
    "Customer clicks 'Call Now' (+91 98254 85520) or 'Inquire on WhatsApp' action button.",
    "Native phone dialer prompted with business number, or WhatsApp window launched with pre-filled inquiry text."
)

# Slide 15: FR-3
slide15 = prs.slides.add_slide(blank_layout)
set_slide_background(slide15)
add_academic_header(slide15, "Requirement Engineering", "(Functional Requirements)")
draw_single_fr_slide(
    slide15, "FR-3", "Workshop Geolocation & Turn-by-Turn Navigation",
    "Used to display the exact workshop location in Limbayat, Surat and provide GPS routing for equipment drop-off.",
    "Customer",
    "Customer scrolls to Contact section or selects 'Get Directions' on Google Map.",
    "Embedded interactive Google Map centered on Plot 199, Sanjay Nagar, Limbayat, with GPS navigation route generated."
)


# ==============================================================================
# SLIDE 16: SOFTWARE DEVELOPMENT MODEL (AGILE SCRUM WITH DIAGRAM)
# ==============================================================================
slide16 = prs.slides.add_slide(blank_layout)
set_slide_background(slide16)
add_academic_header(slide16, "Software Development Model")

content_box = slide16.shapes.add_textbox(Inches(1.2), Inches(1.7), Inches(10.9), Inches(2.2))
tf = content_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
r = p.add_run()
r.text = "This project follows the Agile (Scrum) development model"
r.font.size = Pt(14)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK

p2 = tf.add_paragraph()
r2 = p2.add_run()
r2.text = "Work is broken into short sprints, each delivering a usable increment — e.g., service catalog first, then direct calling integration, then interactive maps, then local SEO optimization"
r2.font.size = Pt(13)
r2.font.color.rgb = COLOR_BLACK
p2.space_before = Pt(8)

p3 = tf.add_paragraph()
r3 = p3.add_run()
r3.text = "Each sprint includes planning, development, testing, and a review with the client (Mr. Murlidhar Chaudhari) before moving to the next module"
r3.font.size = Pt(13)
r3.font.color.rgb = COLOR_BLACK
p3.space_before = Pt(8)

# 5 Cyclic Agile Boxes: Plan -> Design -> Develop -> Test -> Review
agile_steps = ["Plan", "Design", "Develop", "Test", "Review"]
agile_w = Inches(1.8)
agile_h = Inches(1.0)
agile_gap = Inches(0.4)
agile_start = Inches(1.4)
agile_y = Inches(4.3)

for idx, step in enumerate(agile_steps):
    bx = agile_start + idx * (agile_w + agile_gap)
    add_diagram_box(slide16, bx, agile_y, agile_w, agile_h, step)
    
    if idx < len(agile_steps) - 1:
        arr = slide16.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, bx + agile_w + Inches(0.08), agile_y + Inches(0.35), Inches(0.24), Inches(0.3))
        arr.fill.solid(); arr.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr.line.fill.background()

# Feedback loop line beneath
loop_line = slide16.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.4), Inches(5.65), Inches(10.6), Inches(0.08))
loop_line.fill.solid(); loop_line.fill.fore_color.rgb = COLOR_ACCENT_BLUE; loop_line.line.fill.background()

arr_up = slide16.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(1.3), Inches(5.35), Inches(0.25), Inches(0.35))
arr_up.fill.solid(); arr_up.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr_up.line.fill.background()

foot_agile = slide16.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.6))
p = foot_agile.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "repeats every sprint, with client review at the end of each cycle"
r.font.size = Pt(13)
r.font.color.rgb = COLOR_BLACK


# ==============================================================================
# SLIDE 17: MODELING (Developing the use cases)
# ==============================================================================
slide17 = prs.slides.add_slide(blank_layout)
set_slide_background(slide17)
add_academic_header(slide17, "Modeling", "(Developing the use cases)")

content_box = slide17.shapes.add_textbox(Inches(1.0), Inches(1.7), Inches(11.333), Inches(5.4))
tf = content_box.text_frame
tf.word_wrap = True

use_case_questions = [
    ("I. Who are the primary and secondary actors?", "Primary: Event Organizers, DJs, Sound Rental Vendors, Household Consumers. Secondary: Chief Technician (Murlidhar Chaudhari)."),
    ("II. What are the actor's goal(s)?", "Quickly find reliable repair capabilities for faulty equipment, verify technician credentials, and initiate direct contact."),
    ("III. What preconditions should exist before the story begins?", "Customer has access to an active internet connection and a web browser on mobile or desktop."),
    ("IV. What main tasks or functions are performed by the actor?", "Browse repair services, inspect diagnostic highlights, trigger direct phone calls, initiate WhatsApp messages, and locate the workshop on map."),
    ("V. What exceptions might be considered as the story is described?", "Customer device lacks cellular calling capability (system falls back to copy-to-clipboard or WhatsApp)."),
    ("VI. What variations in the actor's interaction are possible?", "Direct mobile call vs desktop WhatsApp Web redirect; browsing by category vs scrolling full page."),
    ("VII. What system information will the actor acquire, produce, or change?", "Actor acquires repair specializations, diagnostic highlights, operating hours, and location coordinates."),
    ("VIII. Will the actor have to inform the system about changes in the external environment?", "Owner updates phone numbers or adds new equipment specializations via centralized data file."),
    ("IX. What information does the actor desire from the system?", "Turnaround feasibility, technician competence, exact workshop address in Limbayat."),
    ("X. Does the actor wish to be informed about unexpected changes?", "Yes, holiday workshop closures or emergency on-call notices displayed prominently in banner.")
]

for idx, (q, a) in enumerate(use_case_questions):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q} "
    r_q.font.size = Pt(11)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    
    r_a = p.add_run()
    r_a.text = a
    r_a.font.size = Pt(10.5)
    r_a.font.color.rgb = COLOR_GRAY_DARK
    p.space_before = Pt(4)


# ==============================================================================
# SLIDE 18: MODELING (USE CASE DIAGRAM)
# ==============================================================================
slide18 = prs.slides.add_slide(blank_layout)
set_slide_background(slide18)
add_academic_header(slide18, "Modeling", "(Use Case Diagram)")

# System Boundary Box
sys_box = slide18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(1.8), Inches(4.8), Inches(4.8))
sys_box.fill.background()
sys_box.line.color.rgb = COLOR_BLACK
sys_box.line.width = Pt(2.0)

# System Label
lbl_box = slide18.shapes.add_textbox(Inches(4.5), Inches(6.0), Inches(4.8), Inches(0.5))
p = lbl_box.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Vikash Electronics Web Platform"
r.font.size = Pt(13)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK

# Use cases (Ovals with red border matching reference slide 18)
use_cases_list = [
    ("Browse Repair Services", Inches(2.1)),
    ("View Diagnostic Details", Inches(2.9)),
    ("Call Technician Directly", Inches(3.7)),
    ("Inquire via WhatsApp", Inches(4.5)),
    ("Locate Workshop & Navigate", Inches(5.3))
]

for title, y_pos in use_cases_list:
    oval = slide18.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.9), y_pos, Inches(4.0), Inches(0.65))
    oval.fill.solid(); oval.fill.fore_color.rgb = COLOR_WHITE
    oval.line.color.rgb = COLOR_RED; oval.line.width = Pt(2.0)
    p = oval.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_RED

# Actor 1: Customer (Left)
act1 = slide18.shapes.add_textbox(Inches(1.2), Inches(3.2), Inches(2.0), Inches(1.2))
p = act1.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "👤\nCustomer\n(DJ / Event Organizer)"
r.font.size = Pt(12)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK

# Connect lines from Customer to Use Cases
for _, y_pos in use_cases_list:
    line = slide18.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.8), y_pos + Inches(0.3), Inches(2.1), Inches(0.02))
    line.fill.solid(); line.fill.fore_color.rgb = COLOR_GRAY_DARK; line.line.fill.background()

# Actor 2: Owner / Admin (Right)
act2 = slide18.shapes.add_textbox(Inches(10.0), Inches(4.2), Inches(2.2), Inches(1.2))
p = act2.text_frame.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "🔧\nOwner / Technician\n(Murlidhar Chaudhari)"
r.font.size = Pt(12)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK

# Connect line from Owner to Contact/Locate
for y_pos in [Inches(3.7), Inches(4.5), Inches(5.3)]:
    line = slide18.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(8.9), y_pos + Inches(0.3), Inches(1.5), Inches(0.02))
    line.fill.solid(); line.fill.fore_color.rgb = COLOR_GRAY_DARK; line.line.fill.background()


# ==============================================================================
# SLIDE 19: MODELING (CLASS DIAGRAM)
# ==============================================================================
slide19 = prs.slides.add_slide(blank_layout)
set_slide_background(slide19)
add_academic_header(slide19, "Modeling", "(Class Diagram)")

def draw_uml_class(slide, left, top, width, height, class_name, attributes, methods):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid(); shape.fill.fore_color.rgb = COLOR_WHITE
    shape.line.color.rgb = COLOR_BLACK; shape.line.width = Pt(1.5)
    
    # Header strip
    hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.4))
    hdr.fill.solid(); hdr.fill.fore_color.rgb = COLOR_DIAGRAM_BG
    hdr.line.color.rgb = COLOR_BLACK; hdr.line.width = Pt(1.5)
    p = hdr.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = class_name; r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = COLOR_BLACK
    
    # Body text (Attributes + Methods)
    tb = slide.shapes.add_textbox(left + Inches(0.05), top + Inches(0.4), width - Inches(0.1), height - Inches(0.45))
    tf = tb.text_frame; tf.word_wrap = True
    
    for idx, attr in enumerate(attributes):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        r = p.add_run(); r.text = f"- {attr}"; r.font.size = Pt(8.5); r.font.color.rgb = COLOR_GRAY_DARK
    
    p_sep = tf.add_paragraph()
    r = p_sep.add_run(); r.text = "---------------------------------"; r.font.size = Pt(7); r.font.color.rgb = COLOR_TABLE_BORDER
    
    for meth in methods:
        p = tf.add_paragraph()
        r = p.add_run(); r.text = f"+ {meth}()"; r.font.size = Pt(8.5); r.font.color.rgb = COLOR_BLACK

# Class 1: Customer
draw_uml_class(slide19, Inches(0.8), Inches(1.8), Inches(3.2), Inches(2.2), "Customer",
               ["customerId: int", "name: string", "phone: string", "eventType: string"],
               ["browseServices", "viewFaultDiagnostics", "callWorkshop", "sendWhatsAppInquiry"])

# Class 2: RepairService
draw_uml_class(slide19, Inches(5.0), Inches(1.8), Inches(3.4), Inches(2.2), "RepairService",
               ["serviceId: string", "title: string", "category: string", "description: string", "highlights: string[]"],
               ["getDetails", "filterByCategory", "getFaultGuide"])

# Class 3: Technician
draw_uml_class(slide19, Inches(9.3), Inches(1.8), Inches(3.2), Inches(2.2), "Technician",
               ["technicianId: int", "name: string", "phone: string", "experienceYears: int"],
               ["receiveCall", "respondWhatsApp", "inspectEquipment", "executeRepair"])

# Class 4: RepairInquiry
draw_uml_class(slide19, Inches(2.8), Inches(4.6), Inches(3.4), Inches(2.2), "RepairInquiry",
               ["inquiryId: string", "equipmentType: string", "faultDescription: string", "contactChannel: enum", "status: string"],
               ["createInquiry", "updateStatus", "estimateTurnaround"])

# Class 5: WorkshopLocation
draw_uml_class(slide19, Inches(7.2), Inches(4.6), Inches(3.4), Inches(2.2), "WorkshopLocation",
               ["plotNo: string", "area: string", "landmark: string", "city: string", "latitude: float", "longitude: float"],
               ["getCoordinates", "openGoogleMapsRoute", "getOperatingHours"])

# Relationship Labels
rl_tb = slide19.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.5))
p = rl_tb.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Customer (1) places (*) RepairInquiry; RepairInquiry (*) references (1) RepairService managed by (1) Technician at (1) WorkshopLocation"; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = COLOR_BLACK


# ==============================================================================
# SLIDE 20: MODELING (SEQUENCE DIAGRAM)
# ==============================================================================
slide20 = prs.slides.add_slide(blank_layout)
set_slide_background(slide20)
add_academic_header(slide20, "Modeling", "(Sequence Diagram)")

# Subtitle
sub_box = slide20.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.4))
p = sub_box.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Use case: Equipment Service Browsing & Direct WhatsApp Consultation"; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = COLOR_BLACK

participants = [
    ("Customer", Inches(1.2)),
    ("Website\n(Frontend)", Inches(3.8)),
    ("Server Logic\n(Next.js)", Inches(6.4)),
    ("Data Store\n(business.ts)", Inches(8.8)),
    ("WhatsApp / Dialer\nGateway", Inches(11.2))
]

for name, x in participants:
    # Lifeline box
    add_diagram_box(slide20, x - Inches(0.9), Inches(2.0), Inches(1.8), Inches(0.6), name)
    # Vertical lifeline
    line = slide20.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, Inches(2.6), Inches(0.02), Inches(4.3))
    line.fill.solid(); line.fill.fore_color.rgb = COLOR_TABLE_BORDER; line.line.fill.background()

messages = [
    (Inches(1.2), Inches(3.8), Inches(2.9), "1. Browse Repair Catalog", True),
    (Inches(3.8), Inches(6.4), Inches(3.4), "2. Request Service Data", True),
    (Inches(6.4), Inches(8.8), Inches(3.9), "3. Query Categories & Specs", True),
    (Inches(8.8), Inches(6.4), Inches(4.4), "4. Return Equipment Specs", False),
    (Inches(6.4), Inches(3.8), Inches(4.9), "5. Render Service Cards & Bench Photos", False),
    (Inches(1.2), Inches(3.8), Inches(5.4), "6. Click 'Inquire on WhatsApp'", True),
    (Inches(3.8), Inches(11.2), Inches(5.9), "7. Launch WhatsApp with Pre-filled Fault Query", True),
    (Inches(11.2), Inches(1.2), Inches(6.4), "8. Direct Chat Connected with Technician", False)
]

for x1, x2, y, text, is_solid in messages:
    left = min(x1, x2)
    w = abs(x2 - x1)
    
    # Message line
    m_line = slide20.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, y, w, Inches(0.02))
    m_line.fill.solid(); m_line.fill.fore_color.rgb = COLOR_BLACK; m_line.line.fill.background()
    
    # Arrow head
    if x2 > x1:
        arr = slide20.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x2 - Inches(0.12), y - Inches(0.05), Inches(0.12), Inches(0.12))
    else:
        arr = slide20.shapes.add_shape(MSO_SHAPE.LEFT_ARROW, x2, y - Inches(0.05), Inches(0.12), Inches(0.12))
    arr.fill.solid(); arr.fill.fore_color.rgb = COLOR_BLACK; arr.line.fill.background()
    
    # Text label above line
    lbl = slide20.shapes.add_textbox(left, y - Inches(0.28), w, Inches(0.3))
    p = lbl.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text; r.font.size = Pt(8.5); r.font.bold = True; r.font.color.rgb = COLOR_BLACK


# ==============================================================================
# SLIDE 21: MODELING (ACTIVITY DIAGRAM)
# ==============================================================================
slide21 = prs.slides.add_slide(blank_layout)
set_slide_background(slide21)
add_academic_header(slide21, "Modeling", "(Activity Diagram)")

# Initial Black Circle
c_init = slide21.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.3), Inches(1.7), Inches(0.35), Inches(0.35))
c_init.fill.solid(); c_init.fill.fore_color.rgb = COLOR_BLACK; c_init.line.fill.background()

# Action 1: Browse Services
add_diagram_box(slide21, Inches(5.0), Inches(2.2), Inches(2.9), Inches(0.6), "Browse Repair Services Catalog")

# Action 2: Inspect Fault & Diagnostics
add_diagram_box(slide21, Inches(5.0), Inches(3.0), Inches(2.9), Inches(0.6), "Inspect Technical Fault Details")

# Decision Diamond: Contact via Phone or WhatsApp?
diamond = slide21.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(5.4), Inches(3.8), Inches(2.1), Inches(1.1))
diamond.fill.solid(); diamond.fill.fore_color.rgb = COLOR_WHITE
diamond.line.color.rgb = COLOR_RED; diamond.line.width = Pt(1.5)
p = diamond.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Choose\nChannel?"; r.font.size = Pt(9.5); r.font.bold = True; r.font.color.rgb = COLOR_RED

# Branch Left: Phone Call
add_diagram_box(slide21, Inches(1.8), Inches(4.05), Inches(2.6), Inches(0.6), "Launch Native Phone Dialer\n(+91 98254 85520)")

# Branch Right: WhatsApp Chat
add_diagram_box(slide21, Inches(8.5), Inches(4.05), Inches(2.6), Inches(0.6), "Open Pre-filled WhatsApp Chat\n(Service & Fault Query)")

# Action 4: Technician Consultation
add_diagram_box(slide21, Inches(5.0), Inches(5.2), Inches(2.9), Inches(0.6), "Consult with Technician (Murlidhar C.)")

# Action 5: Drop-off at Workshop
add_diagram_box(slide21, Inches(5.0), Inches(6.0), Inches(2.9), Inches(0.6), "Drop-off Gear at Limbayat Workshop")

# Final Bullseye Circle
c_final_out = slide21.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.3), Inches(6.8), Inches(0.4), Inches(0.4))
c_final_out.fill.background(); c_final_out.line.color.rgb = COLOR_BLACK; c_final_out.line.width = Pt(2.0)

c_final_in = slide21.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6.35), Inches(6.85), Inches(0.3), Inches(0.3))
c_final_in.fill.solid(); c_final_in.fill.fore_color.rgb = COLOR_BLACK; c_final_in.line.fill.background()


# ==============================================================================
# SLIDE 22: MODELING (STATE DIAGRAM - REPAIR JOB LIFECYCLE)
# ==============================================================================
slide22 = prs.slides.add_slide(blank_layout)
set_slide_background(slide22)
add_academic_header(slide22, "Modeling", "(State Diagram)")

sub_box = slide22.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.4))
p = sub_box.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Repair Job Lifecycle (Equipment Intake to Delivery)"; r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = COLOR_BLACK

# Initial State
init_dot = slide22.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), Inches(3.2), Inches(0.35), Inches(0.35))
init_dot.fill.solid(); init_dot.fill.fore_color.rgb = COLOR_BLACK; init_dot.line.fill.background()

states = [
    ("Inquiry\nInitiated", Inches(1.5)),
    ("Inspected &\nEstimated", Inches(3.6)),
    ("Repair In\nProgress", Inches(5.7)),
    ("Stress Load\nTesting", Inches(7.8)),
    ("Stage Ready\nDelivered", Inches(9.9))
]

for name, x in states:
    add_diagram_box(slide22, x, Inches(2.9), Inches(1.7), Inches(0.9), name)

# Arrows between states
transitions = ["online inquiry", "equipment intake", "quote approved", "bench burn-in"]
for idx in range(4):
    x_from = Inches(1.5) + idx * Inches(2.1) + Inches(1.7)
    arr = slide22.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x_from + Inches(0.05), Inches(3.25), Inches(0.3), Inches(0.2))
    arr.fill.solid(); arr.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr.line.fill.background()

# Final State Circle on far right
fin_out = slide22.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12.0), Inches(3.15), Inches(0.4), Inches(0.4))
fin_out.fill.background(); fin_out.line.color.rgb = COLOR_BLACK; fin_out.line.width = Pt(2.0)
fin_in = slide22.shapes.add_shape(MSO_SHAPE.OVAL, Inches(12.05), Inches(3.2), Inches(0.3), Inches(0.3))
fin_in.fill.solid(); fin_in.fill.fore_color.rgb = COLOR_BLACK; fin_in.line.fill.background()

# Cancelled State Box (Red, below)
box_cancel = slide22.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(4.8), Inches(2.5), Inches(0.9))
box_cancel.fill.solid(); box_cancel.fill.fore_color.rgb = COLOR_WHITE
box_cancel.line.color.rgb = COLOR_RED; box_cancel.line.width = Pt(2.0)
p = box_cancel.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run(); r.text = "Inquiry Cancelled /\nNot Repairable"; r.font.bold = True; r.font.size = Pt(11); r.font.color.rgb = COLOR_RED

# Transition to Cancelled
arr_c = slide22.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(4.3), Inches(3.9), Inches(0.2), Inches(0.8))
arr_c.fill.solid(); arr_c.fill.fore_color.rgb = COLOR_RED; arr_c.line.fill.background()


# ==============================================================================
# SLIDE 23: DEPLOYMENT DIAGRAM
# ==============================================================================
slide23 = prs.slides.add_slide(blank_layout)
set_slide_background(slide23)
add_academic_header(slide23, "Deployment Diagram")

# Node 1: Customer Device
add_diagram_box(slide23, Inches(0.8), Inches(2.2), Inches(3.4), Inches(2.2), "Customer Device\n<<device>>", "\nWeb / Mobile Browser\n(Chrome, Safari, Edge)\nHTTP/2 & Service Worker")

# Node 2: Web Server (Center)
add_diagram_box(slide23, Inches(5.0), Inches(2.2), Inches(3.4), Inches(2.2), "Web Server\n<<server>>", "\nNext.js 14 Node.js Runtime\nServer-Side Rendering (SSR)\nFrontend UI + API Handlers")

# Node 3: External Gateway (Right)
add_diagram_box(slide23, Inches(9.2), Inches(2.2), Inches(3.4), Inches(2.2), "Communication Services\n<<external system>>", "\nWhatsApp Business API\nNative Telephony Protocol\nGoogle Maps Routing API")

# Node 4: Static Asset & Data Store (Bottom Center)
add_diagram_box(slide23, Inches(5.0), Inches(5.0), Inches(3.4), Inches(1.5), "Static Data & Asset Store\n<<storage>>", "\nbusiness.ts, SEO Metadata,\nWorkshop & Service Images")

# Arrows between Nodes
arr_https = slide23.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(4.3), Inches(3.2), Inches(0.6), Inches(0.2))
arr_https.fill.solid(); arr_https.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr_https.line.fill.background()
tb = slide23.shapes.add_textbox(Inches(4.2), Inches(2.8), Inches(0.8), Inches(0.3))
tb.text_frame.paragraphs[0].text = "HTTPS"; tb.text_frame.paragraphs[0].font.size = Pt(9); tb.text_frame.paragraphs[0].font.bold = True

arr_ext = slide23.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(8.5), Inches(3.2), Inches(0.6), Inches(0.2))
arr_ext.fill.solid(); arr_ext.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr_ext.line.fill.background()
tb = slide23.shapes.add_textbox(Inches(8.4), Inches(2.8), Inches(0.8), Inches(0.3))
tb.text_frame.paragraphs[0].text = "API / Tel"; tb.text_frame.paragraphs[0].font.size = Pt(9); tb.text_frame.paragraphs[0].font.bold = True

arr_db = slide23.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.6), Inches(4.5), Inches(0.2), Inches(0.4))
arr_db.fill.solid(); arr_db.fill.fore_color.rgb = COLOR_ACCENT_BLUE; arr_db.line.fill.background()

foot_dep = slide23.shapes.add_textbox(Inches(0.8), Inches(6.7), Inches(11.733), Inches(0.5))
p = foot_dep.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "The customer's browser communicates with the Next.js web server over HTTPS; the server queries the structured data store and redirects to WhatsApp/Google Maps API."
r.font.size = Pt(12); r.font.color.rgb = COLOR_BLACK


# ==============================================================================
# SLIDES 24 & 25: MODELING (USE CASES TABLES)
# ==============================================================================
def draw_detailed_uc_table(slide, uc_title, actor, goal, precond, trigger):
    rows = 5
    cols = 2
    table_shape = slide.shapes.add_table(rows, cols, Inches(1.0), Inches(2.0), Inches(11.3), Inches(3.8))
    table = table_shape.table
    table.columns[0].width = Inches(2.4)
    table.columns[1].width = Inches(8.9)
    
    c00 = table.cell(0, 0); c00.fill.solid(); c00.fill.fore_color.rgb = COLOR_TABLE_HEADER
    c00.text = "Use Case Name"; p = c00.text_frame.paragraphs[0]; p.font.bold = True; p.font.size = Pt(13); p.font.color.rgb = COLOR_BLACK
    
    c01 = table.cell(0, 1); c01.fill.solid(); c01.fill.fore_color.rgb = COLOR_TABLE_HEADER
    c01.text = uc_title; p = c01.text_frame.paragraphs[0]; p.font.bold = True; p.font.size = Pt(13); p.font.color.rgb = COLOR_BLACK
    
    data = [
        ("Primary Actor(s)", actor),
        ("Goal of the use case", goal),
        ("Precondition(s)", precond),
        ("Trigger", trigger)
    ]
    for r_idx, (lbl, val) in enumerate(data, start=1):
        c0 = table.cell(r_idx, 0); c0.fill.solid(); c0.fill.fore_color.rgb = COLOR_WHITE; c0.text = lbl
        p0 = c0.text_frame.paragraphs[0]; p0.font.bold = True; p0.font.size = Pt(12); p0.font.color.rgb = COLOR_BLACK
        
        c1 = table.cell(r_idx, 1); c1.fill.solid(); c1.fill.fore_color.rgb = COLOR_WHITE; c1.text = val
        p1 = c1.text_frame.paragraphs[0]; p1.font.size = Pt(12); p1.font.color.rgb = COLOR_GRAY_DARK

# Slide 24: UC-1
slide24 = prs.slides.add_slide(blank_layout)
set_slide_background(slide24)
add_academic_header(slide24, "Modeling", "(Use cases)")
draw_detailed_uc_table(
    slide24, "Inquire for Equipment Repair",
    "Customer (DJ, Event Organizer, Consumer)",
    "To browse the specialized repair services and successfully initiate a direct repair consultation with the chief technician",
    "Customer has an active internet connection, and the Vikash Electronics website is loaded on mobile or desktop",
    "Customer selects a specific equipment service (e.g. Speakers, Amplifiers, Sharpy Lights) and taps 'Call Now' or 'Inquire on WhatsApp'"
)

# Slide 25: UC-2
slide25 = prs.slides.add_slide(blank_layout)
set_slide_background(slide25)
add_academic_header(slide25, "Modeling", "(Use cases)")
draw_detailed_uc_table(
    slide25, "Locate Workshop & Navigate",
    "Customer (Equipment Owner)",
    "To locate the physical workshop address in Limbayat, Surat and obtain GPS turn-by-turn navigation for equipment drop-off",
    "Customer has identified faulty electronic gear requiring physical diagnosis and reached the Contact section",
    "Customer clicks 'Get Directions' or interacts with the embedded Google Maps module"
)


# ==============================================================================
# SLIDE 26: OTHERS
# ==============================================================================
slide26 = prs.slides.add_slide(blank_layout)
set_slide_background(slide26)
add_academic_header(slide26, "Others")

content_box = slide26.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

others_items = [
    "Technology stack implemented: Next.js 14 (React 18), TypeScript for strict type contracts, Tailwind CSS for dark-mode utility styling, Framer Motion for UI micro-animations, and Lucide Icons",
    "SEO & Search Discovery: Schema.org LocalBusiness JSON-LD markup, OpenGraph social cards, geo-targeted meta tags for Surat pro-audio repair rankings",
    "Additional planned features: Online repair ticket tracking system where customers can check job status by mobile number, an admin view for the technician to log incoming equipment repairs, and direct links to YouTube repair tutorials",
    "Screenshots of development: Live responsive portal running in development mode on localhost:3000 featuring authentic workshop bench testing photos and full service catalog"
]

for idx, item in enumerate(others_items):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = item
    r.font.size = Pt(13.5)
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(14)


# ==============================================================================
# SLIDE 27: THANK YOU SLIDE
# ==============================================================================
slide27 = prs.slides.add_slide(blank_layout)
set_slide_background(slide27)

if os.path.exists(LOGO_UTU):
    slide27.shapes.add_picture(LOGO_UTU, Inches(0.8), Inches(0.4), width=Inches(1.2))
if os.path.exists(LOGO_AMTICS):
    slide27.shapes.add_picture(LOGO_AMTICS, Inches(11.3), Inches(0.4), width=Inches(1.2))

ty_box = slide27.shapes.add_textbox(Inches(2.0), Inches(2.2), Inches(9.333), Inches(2.0))
tf = ty_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Thank You"
r.font.size = Pt(52)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

# Red underline beneath Thank You
line_ty = slide27.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.1), Inches(3.6), Inches(3.1), Inches(0.05))
line_ty.fill.solid(); line_ty.fill.fore_color.rgb = COLOR_RED; line_ty.line.fill.background()

cred_box = slide27.shapes.add_textbox(Inches(2.0), Inches(4.3), Inches(9.333), Inches(1.8))
tf = cred_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Prathmesh Chaudhari (202503103510061)"
r.font.size = Pt(18)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = "Vikash Electronics – Web-Based DJ & Electronic Equipment Repair Portal"
r2.font.size = Pt(16)
r2.font.color.rgb = COLOR_GRAY_DARK
r2.font.name = "Arial"
p2.space_before = Pt(8)


# Save presentation
prs.save(OUTPUT_FILE)
print(f"Successfully generated all 27 slides at: {OUTPUT_FILE}")
