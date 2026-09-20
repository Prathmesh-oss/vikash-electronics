import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# --- COLOR PALETTE MATCHING ACADEMIC TEMPLATE ---
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_BLACK = RGBColor(15, 23, 42)      # Deep Black / Dark Slate #0F172A
COLOR_RED = RGBColor(220, 38, 38)       # Crimson Red for Subtitles in Parentheses #DC2626
COLOR_BLUE = RGBColor(2, 132, 199)      # Cyan Blue for Project Title #0284C7
COLOR_DARK_BLUE = RGBColor(15, 76, 129) # CSE Department Slate Blue
COLOR_GRAY_DARK = RGBColor(51, 65, 85)  # Slate 700
COLOR_GRAY_LIGHT = RGBColor(241, 245, 249) # Slate 100
COLOR_TABLE_HEADER = RGBColor(230, 204, 204) # Soft Rosy Pink Header from Template #E6CCCC
COLOR_TABLE_BORDER = RGBColor(203, 213, 225) # Slate 300
COLOR_GOLD = RGBColor(217, 119, 6)      # Amber 600

BASE_DIR = r"d:\A_S_Projects\Vikas_Electronics"
OUTPUT_FILE = os.path.join(BASE_DIR, "Vikash_Electronics_Academic_Presentation.pptx")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

LOGO_UTU = os.path.join(BASE_DIR, "utu_flame_logo.png")
LOGO_AMTICS = os.path.join(BASE_DIR, "amtics_circle_logo.png")

def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLOR_WHITE

def add_academic_header(slide, title_text, red_subtitle=""):
    """Adds the standard academic heading matching the PDF template"""
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.3))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0)
    tf.margin_top = Inches(0)
    
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r_title = p.add_run()
    r_title.text = title_text
    r_title.font.size = Pt(28)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_BLACK
    r_title.font.name = "Arial"
    
    if red_subtitle:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        r_sub = p2.add_run()
        r_sub.text = red_subtitle
        r_sub.font.size = Pt(22)
        r_sub.font.bold = True
        r_sub.font.color.rgb = COLOR_RED
        r_sub.font.name = "Arial"
        p2.space_before = Pt(4)

# ==============================================================================
# SLIDE 1: TITLE SLIDE (Asha M. Tarsadia Institute of Computer Science and Technology)
# ==============================================================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)

# Top Logos
if os.path.exists(LOGO_UTU):
    slide1.shapes.add_picture(LOGO_UTU, Inches(0.8), Inches(0.4), width=Inches(1.2))

if os.path.exists(LOGO_AMTICS):
    slide1.shapes.add_picture(LOGO_AMTICS, Inches(11.3), Inches(0.4), width=Inches(1.2))

# University & College Header
header_box = slide1.shapes.add_textbox(Inches(2.2), Inches(0.4), Inches(8.933), Inches(1.5))
tf = header_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Asha M. Tarsadia Institute of\nComputer Science and Technology"
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p_dept = tf.add_paragraph()
p_dept.alignment = PP_ALIGN.CENTER
r_dept = p_dept.add_run()
r_dept.text = "Computer Science and Engineering"
r_dept.font.size = Pt(20)
r_dept.font.bold = True
r_dept.font.color.rgb = COLOR_DARK_BLUE
r_dept.font.name = "Arial"
p_dept.space_before = Pt(6)

# Project Title in Blue
title_box = slide1.shapes.add_textbox(Inches(1.0), Inches(2.3), Inches(11.333), Inches(1.3))
tf = title_box.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Vikash Electronics – Web-Based DJ & Electronic Equipment Repair Management and Showcase Portal"
r.font.size = Pt(24)
r.font.bold = True
r.font.color.rgb = COLOR_BLUE
r.font.name = "Arial"

# Prepared by
prep_box = slide1.shapes.add_textbox(Inches(2.0), Inches(4.0), Inches(9.333), Inches(1.5))
tf = prep_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Prepared by"
r.font.size = Pt(17)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = "Prathmesh Chaudhari"
r2.font.size = Pt(16)
r2.font.color.rgb = COLOR_BLACK
r2.font.name = "Arial"
p2.space_before = Pt(4)

p3 = tf.add_paragraph()
p3.alignment = PP_ALIGN.CENTER
r3 = p3.add_run()
r3.text = "(Enrollment No: 202503103510061)"
r3.font.size = Pt(15)
r3.font.color.rgb = COLOR_BLACK
r3.font.name = "Arial"
p3.space_before = Pt(2)

# Guided by
guide_box = slide1.shapes.add_textbox(Inches(2.0), Inches(5.6), Inches(9.333), Inches(1.4))
tf = guide_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
r = p.add_run()
r.text = "Guided by"
r.font.size = Pt(17)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
r2 = p2.add_run()
r2.text = "Prof. [Guide Name / Faculty Advisor]"
r2.font.size = Pt(15)
r2.font.color.rgb = COLOR_BLACK
r2.font.name = "Arial"
p2.space_before = Pt(4)


# ==============================================================================
# SLIDE 2: OUTLINE
# ==============================================================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2)
add_academic_header(slide2, "Outline")

content_box = slide2.shapes.add_textbox(Inches(1.8), Inches(1.8), Inches(9.7), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

outline_items = [
    "Client Profile",
    "Project Introduction",
    "Objective(s)",
    "Problem Analysis",
    "Requirement Engineering",
    "Software Development Model",
    "Modeling",
    "Others"
]

for idx, item in enumerate(outline_items):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = f"•  {item}"
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = COLOR_BLACK
    r.font.name = "Arial"
    p.space_before = Pt(14)


# ==============================================================================
# SLIDE 3: CLIENT PROFILE
# ==============================================================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3)
add_academic_header(slide3, "Client Profile")

content_box = slide3.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

sections = [
    ("Client name and address:", [
        "Client / Business Name: Vikash Electronics",
        "Proprietor & Chief Technician: Mr. Murlidhar Chaudhari",
        "Workshop Address: Plot No. 199, Sarve No. 1-2, Sanjay Nagar, Near Udhna Yard, Limbayat, Surat - 394210, Gujarat, India.",
        "Direct Contact: +91 98254 85520 | Domain: Pro-Audio & Stage Lighting Electronic Repairs"
    ]),
    ("Your communication:", [
        "Conducted in-person field visits to the Surat workshop to observe live repair processes (voice coil rewinding, transistor testing, and Sharpy optical calibration).",
        "Engaged in iterative telephonic & WhatsApp sessions to prioritize essential customer touchpoints (instant phone dialer, emergency WhatsApp chat, equipment categories).",
        "Demonstrated responsive prototypes on mobile and desktop devices to validate usability and local SEO requirements."
    ]),
    ("Contract letter:", [
        "Obtained formal Project Authorization & Client Consent from Mr. Murlidhar Chaudhari (Owner, Vikash Electronics) endorsing the design, development, and digital launch of the web portal."
    ])
]

for s_idx, (sec_title, points) in enumerate(sections):
    p_sec = tf.paragraphs[0] if s_idx == 0 else tf.add_paragraph()
    r_sec = p_sec.add_run()
    r_sec.text = sec_title
    r_sec.font.size = Pt(17)
    r_sec.font.bold = True
    r_sec.font.color.rgb = COLOR_BLACK
    r_sec.font.name = "Arial"
    p_sec.space_before = Pt(16 if s_idx > 0 else 0)

    for pt in points:
        p_pt = tf.add_paragraph()
        r_pt = p_pt.add_run()
        r_pt.text = f"    • {pt}"
        r_pt.font.size = Pt(12)
        r_pt.font.color.rgb = COLOR_GRAY_DARK
        r_pt.font.name = "Arial"
        p_pt.space_before = Pt(4)


# ==============================================================================
# SLIDE 4: PROJECT INTRODUCTION
# ==============================================================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4)
add_academic_header(slide4, "Project Introduction")

content_box = slide4.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

intro_points = [
    ("Background of the Business:", 
     "Surat hosts thousands of grand events, weddings, stage performances, and DJ concerts annually, creating massive demand for reliable high-voltage audio and stage lighting gear. Vikash Electronics is a premier specialized repair center handling professional power amplifiers, 18-inch subwoofers, line arrays, Sharpy moving head lights, DMX consoles, and smart TVs."),
    
    ("The Core Problem & Gap:", 
     "Historically, specialized repair services operated strictly through word-of-mouth with zero digital presence. Event managers, DJ operators, and sound suppliers facing urgent on-site equipment breakdowns lacked immediate access to service specializations, workshop credibility, exact location, or direct emergency contact."),
     
    ("Proposed Web-Based Solution:", 
     "A modern, highly performant web application engineered using Next.js 14, TypeScript, and TailwindCSS. The platform serves as a complete digital service catalog, local SEO engine with Schema.org LocalBusiness structured data, authentic workshop credentialing showcase, and direct one-click communication bridge.")
]

for idx, (head, body) in enumerate(intro_points):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r1 = p.add_run()
    r1.text = f"•  {head}\n"
    r1.font.size = Pt(15)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_BLACK
    r1.font.name = "Arial"
    
    r2 = p.add_run()
    r2.text = f"    {body}"
    r2.font.size = Pt(12.5)
    r2.font.color.rgb = COLOR_GRAY_DARK
    r2.font.name = "Arial"
    p.space_before = Pt(14)


# ==============================================================================
# SLIDE 5: OBJECTIVE(S)
# ==============================================================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5)
add_academic_header(slide5, "Objective(s)")

content_box = slide5.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

objectives = [
    ("1. Establish Comprehensive Online Service Showcase:",
     "Present an organized, technically detailed catalog of all 8 repair specializations (Speakers, Amplifiers, Sharpy Lights, DJ Lights, Mixers, Home Theatres, TVs, and Custom Electronics)."),
     
    ("2. Enable Rapid One-Touch Customer Acquisition:",
     "Implement frictionless direct calling (`tel:`) and instant pre-filled WhatsApp inquiry bridges (`wa.me`) optimized for mobile sound engineers needing urgent event repairs."),
     
    ("3. Promote Authenticity & Technical Credibility:",
     "Showcase verified workshop testing instruments (oscilloscopes, power variacs, signal generators) and chief technician profile (Murlidhar Chaudhari) to inspire customer confidence."),
     
    ("4. Dominate Local Search Indexing (Local SEO):",
     "Embed Schema.org LocalBusiness JSON-LD structured data and geo-targeted keywords to rank on Google Search for DJ & electronic equipment repairs in Surat and South Gujarat."),
     
    ("5. High Performance, Accessibility & Responsiveness:",
     "Achieve sub-second load times, smooth responsive UI across devices, and dark-themed aesthetics aligned with the entertainment and stage lighting industry.")
]

for idx, (head, desc) in enumerate(objectives):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r1 = p.add_run()
    r1.text = f"•  {head} "
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_BLACK
    r1.font.name = "Arial"
    
    r2 = p.add_run()
    r2.text = desc
    r2.font.size = Pt(12)
    r2.font.color.rgb = COLOR_GRAY_DARK
    r2.font.name = "Arial"
    p.space_before = Pt(12)


# ==============================================================================
# SLIDE 6: PROBLEM ANALYSIS (Understand the problem)
# ==============================================================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6)
add_academic_header(slide6, "Problem Analysis", "(Understand the problem)")

content_box = slide6.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_items = [
    ("I. Who are the stakeholders in the solution to the problem?",
     "• Primary Stakeholders: Workshop Owner & Chief Technician (Murlidhar Chaudhari), Professional DJs, Event Management Companies, Stage Lighting Operators, Sound System Rental Providers.\n• Secondary Stakeholders: Marriage Band Crews, Auditorium Managers, and Household Consumers requiring AV Receiver / Smart TV repairs."),
     
    ("II. What are the unknowns?",
     "• Diverse hardware architectures across domestic and imported brands (Yamaha, Pioneer, Studio Master, Clay Paky, JBL).\n• Component availability, varying degree of burn/thermal damage, and non-standard turnaround times per fault."),
     
    ("III. Can the problem be compartmentalized?",
     "• Yes, into distinct functional subsystems: (1) Service Catalog & Hardware Diagnostic Showcase, (2) Direct Emergency Customer Outreach (Call/WhatsApp), (3) Workshop Credibility & Bench Verification, (4) Geolocation Routing (Google Maps), and (5) Local SEO Rich Snippets."),
     
    ("IV. Can the problem be represented graphically?",
     "• Yes, via System Architecture Diagrams, User Interaction Flowcharts, Component Hierarchies, and Use Case Diagrams.")
]

for idx, (q, a) in enumerate(pa_items):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(13)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = f"    {a}"
    r_a.font.size = Pt(11)
    r_a.font.color.rgb = COLOR_GRAY_DARK
    r_a.font.name = "Arial"
    p.space_before = Pt(10)


# ==============================================================================
# SLIDE 7: PROBLEM ANALYSIS (Plan the solution)
# ==============================================================================
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7)
add_academic_header(slide7, "Problem Analysis", "(Plan the solution)")

content_box = slide7.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_plan = [
    ("I. Have you seen similar problem before?",
     "• Yes; local specialized engineering and electronics repair workshops frequently face digital invisibility, relying heavily on traditional word-of-mouth without structured online branding."),
     
    ("II. Has a similar problem been solved?",
     "• Generic business directories (e.g., Justdial, IndiaMART) list repairers, but they lack technical depth (voice coil rewinding, 2SC5200 transistor matching, DMX-512 protocols), exhibit heavy ad clutter, and do not build individual workshop brand equity."),
     
    ("III. Can sub-problems be defined?",
     "• Sub-problem 1: Technical Presentation — Explaining complex electronic repair procedures clearly to both seasoned audio engineers and non-technical clients.\n• Sub-problem 2: Frictionless Engagement — Providing direct calling without burdensome account registrations.\n• Sub-problem 3: Precise Geolocation — Navigating clients to the workshop located in industrial Limbayat, Surat."),
     
    ("IV. Can you represent a solution in a manner that leads to effective implementation?",
     "• Yes; by designing a decoupled component-driven web application using Next.js 14 App Router, structured TypeScript data models (`business.ts`), Tailwind CSS design tokens, and LocalBusiness schema markup.")
]

for idx, (q, a) in enumerate(pa_plan):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(13)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = f"    {a}"
    r_a.font.size = Pt(11)
    r_a.font.color.rgb = COLOR_GRAY_DARK
    r_a.font.name = "Arial"
    p.space_before = Pt(10)


# ==============================================================================
# SLIDE 8: PROBLEM ANALYSIS (Carry out the plan)
# ==============================================================================
slide8 = prs.slides.add_slide(blank_layout)
set_slide_background(slide8)
add_academic_header(slide8, "Problem Analysis", "(Carry out the plan)")

content_box = slide8.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_carry = [
    ("I. Does the solution conform to the plan?",
     "• Yes; the project strictly conforms to planned engineering architecture:\n"
     "  - Framework: Next.js 14 (React 18) leveraging Server Components and fast Client hydration.\n"
     "  - Layout System: Modular component hierarchy (`Navbar`, `Hero`, `Services`, `About`, `Contact`, `Footer`).\n"
     "  - Data Abstraction: Centralized data repository (`src/data/business.ts`) allowing rapid maintenance.\n"
     "  - Responsive Design: Tailwind CSS mobile-first breakpoint grid (`sm`, `md`, `lg`, `xl`).\n"
     "  - SEO Compliance: Semantic HTML5, metadata objects, OpenGraph tags, and Schema.org JSON-LD."),
     
    ("II. Is each component part of the solution provably correct?",
     "• Verified through rigorous unit & functional inspection:\n"
     "  - Navigation (`Navbar.tsx`): Correct anchor routing `#home`, `#services`, `#about`, `#contact` with mobile menu toggle.\n"
     "  - Service Cards (`ServiceCard.tsx`): Type-safe prop rendering with icon bindings and technical highlight tags.\n"
     "  - Communication Bridges: Validated tel-links (`tel:9825485520`) and WhatsApp API endpoints.\n"
     "  - Map Component: Correct Google Maps iframe integration targeting Plot 199, Sanjay Nagar, Limbayat coordinates.\n"
     "  - Build Integrity: Clean compilation with zero TypeScript errors under `tsc --noEmit`.")
]

for idx, (q, a) in enumerate(pa_carry):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(14)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = f"    {a}"
    r_a.font.size = Pt(11.5)
    r_a.font.color.rgb = COLOR_GRAY_DARK
    r_a.font.name = "Arial"
    p.space_before = Pt(14)


# ==============================================================================
# SLIDE 9: PROBLEM ANALYSIS (Examine the results)
# ==============================================================================
slide9 = prs.slides.add_slide(blank_layout)
set_slide_background(slide9)
add_academic_header(slide9, "Problem Analysis", "(Examine the results)")

content_box = slide9.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

pa_results = [
    ("I. Is it possible to test each component part of the solution?",
     "• Yes; tested comprehensively using multiple verification methods:\n"
     "  - Local Server Verification: Next.js dev server running on `http://localhost:3000` with sub-100ms response times.\n"
     "  - Responsive Layout Testing: Verified using Chrome DevTools across Mobile (375px), Tablet (768px), and Desktop (1280px+).\n"
     "  - Link & Action Verification: Validated calling functions, WhatsApp redirects, and external map routing.\n"
     "  - Static Type Checking: Validated with TypeScript 5 compiler ensuring 100% type soundness."),
     
    ("II. Does the solution produce results that conform to the data, functions, and features that are required?",
     "• Fully verified and conforming to all required business criteria:\n"
     "  - Displays authentic business identity, technician name (Murlidhar Chaudhari), and Limbayat address.\n"
     "  - Accurately details 8 distinct repair domains with high-resolution photographic evidence.\n"
     "  - Google Rich Snippets Validator confirms valid JSON-LD LocalBusiness markup.\n"
     "  - Provides instant customer touchpoints, fulfilling the core objective of digital lead generation for the workshop.")
]

for idx, (q, a) in enumerate(pa_results):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q}\n"
    r_q.font.size = Pt(14)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = f"    {a}"
    r_a.font.size = Pt(11.5)
    r_a.font.color.rgb = COLOR_GRAY_DARK
    r_a.font.name = "Arial"
    p.space_before = Pt(14)


# ==============================================================================
# SLIDE 10: REQUIREMENT ENGINEERING (Client-side communication)
# ==============================================================================
slide10 = prs.slides.add_slide(blank_layout)
set_slide_background(slide10)
add_academic_header(slide10, "Requirement Engineering", "(Client-side communication)")

content_box = slide10.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

intro_req = tf.paragraphs[0]
r_intro = intro_req.add_run()
r_intro.text = "Requirements collected during stakeholder meetings with Mr. Murlidhar Chaudhari (Proprietor):"
r_intro.font.size = Pt(14)
r_intro.font.bold = True
r_intro.font.color.rgb = COLOR_BLACK
r_intro.font.name = "Arial"

client_reqs = [
    ("1. Highlighting Niche DJ & Pro Audio Specialization:",
     "Prioritize heavy equipment repair (voice coils, toroidal transformers, power MOSFETs) to distinguish the workshop from generic small appliance repairers."),
     
    ("2. Genuine Workshop & Precision Bench Imagery:",
     "Feature real photos of the Limbayat testing facility, repair workbench, and diagnostic tools to build authentic trust among touring sound engineers and event managers."),
     
    ("3. Direct Phone & Instant WhatsApp Helpline:",
     "Prominently display `+91 98254 85520` with one-tap dialer and WhatsApp chat for emergency event breakdowns (especially weekend weddings and music shows)."),
     
    ("4. Geolocation Accuracy in Limbayat, Surat:",
     "Integrate explicit postal address (Plot 199, Sarve No. 1-2, Sanjay Nagar, Udhna Yard, Limbayat) and an interactive map to prevent logistics confusion."),
     
    ("5. Dark-Themed Professional Visual Tone:",
     "Adopt high-contrast dark aesthetics (navy slate `#0B1120`, amber gold `#F59E0B`, electric cyan `#38BDF8`) reflecting pro-audio concert and lighting environments.")
]

for head, body in client_reqs:
    p = tf.add_paragraph()
    r1 = p.add_run()
    r1.text = f"•  {head} "
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_BLACK
    r1.font.name = "Arial"
    
    r2 = p.add_run()
    r2.text = body
    r2.font.size = Pt(11.5)
    r2.font.color.rgb = COLOR_GRAY_DARK
    r2.font.name = "Arial"
    p.space_before = Pt(10)


# ==============================================================================
# SLIDE 11: REQUIREMENT ENGINEERING (Functional Requirements Table)
# ==============================================================================
slide11 = prs.slides.add_slide(blank_layout)
set_slide_background(slide11)
add_academic_header(slide11, "Requirement Engineering", "(Functional Requirements)")

# Two Functional Requirement Tables formatted exactly like the PDF template
def draw_fr_table(slide, left, top, width, fr_id, title, desc, actor, inputs, outputs):
    rows = 5
    cols = 2
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, Inches(2.3))
    table = table_shape.table
    table.columns[0].width = Inches(1.3)
    table.columns[1].width = width - Inches(1.3)
    
    # Header Row
    cell_00 = table.cell(0, 0)
    cell_00.fill.solid()
    cell_00.fill.fore_color.rgb = COLOR_TABLE_HEADER
    cell_00.text = fr_id
    p = cell_00.text_frame.paragraphs[0]
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_BLACK
    
    cell_01 = table.cell(0, 1)
    cell_01.fill.solid()
    cell_01.fill.fore_color.rgb = COLOR_TABLE_HEADER
    cell_01.text = title
    p = cell_01.text_frame.paragraphs[0]
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_BLACK
    
    data = [
        ("Description", desc),
        ("Actor(s)", actor),
        ("Input(s)", inputs),
        ("Output(s)", outputs)
    ]
    
    for r_idx, (label, val) in enumerate(data, start=1):
        c0 = table.cell(r_idx, 0)
        c0.fill.solid()
        c0.fill.fore_color.rgb = COLOR_WHITE
        c0.text = label
        p0 = c0.text_frame.paragraphs[0]
        p0.font.bold = True
        p0.font.size = Pt(10)
        p0.font.color.rgb = COLOR_BLACK
        
        c1 = table.cell(r_idx, 1)
        c1.fill.solid()
        c1.fill.fore_color.rgb = COLOR_WHITE
        c1.text = val
        p1 = c1.text_frame.paragraphs[0]
        p1.font.size = Pt(10)
        p1.font.color.rgb = COLOR_GRAY_DARK

# Table 1: FR-1 Service Catalog
draw_fr_table(
    slide11, Inches(1.0), Inches(1.8), Inches(11.3),
    "FR-1", "Equipment Repair Service Catalog",
    "Allows users to browse categorized electronic repair capabilities with technical failure modes.",
    "Customer (DJ, Event Manager, Consumer)",
    "User scrolls to services or clicks category tabs (Audio, Lighting, Visual, General).",
    "Renders structured service cards with equipment images, failure symptoms, and component fixes."
)

# Table 2: FR-2 Direct Customer Contact
draw_fr_table(
    slide11, Inches(1.0), Inches(4.5), Inches(11.3),
    "FR-2", "Direct Emergency Communication (Call & WhatsApp)",
    "Provides instant telephonic and WhatsApp connection to the chief technician.",
    "Customer, Chief Technician (Murlidhar Chaudhari)",
    "User clicks on 'Call Now' (+91 98254 85520) or 'Inquire on WhatsApp' action buttons.",
    "Prompts mobile phone dialer or opens WhatsApp chat with pre-filled repair inquiry message."
)


# ==============================================================================
# SLIDE 12: SOFTWARE DEVELOPMENT MODEL (Agile Development)
# ==============================================================================
slide12 = prs.slides.add_slide(blank_layout)
set_slide_background(slide12)
add_academic_header(slide12, "Software Development Model", "(Agile Development Model)")

content_box = slide12.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

intro = tf.paragraphs[0]
r = intro.add_run()
r.text = "Model Chosen: Agile Methodology (Iterative & Incremental Development)"
r.font.size = Pt(15)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

agile_steps = [
    ("Why Agile for this Project?",
     "Agile was selected due to the evolving requirement specifications from the workshop client, requiring continuous feedback loops, rapid prototype validation, and progressive photo cataloging of specialized electronic gear."),
     
    ("Sprint 1 – Requirement Engineering & Wireframing:",
     "• On-site workshop interviews in Limbayat to capture hardware scope (Sharpy lights, line arrays, PA amps).\n• Low-fidelity wireframing of responsive layouts and user flow mapping."),
     
    ("Sprint 2 – Core Architectural Scaffolding:",
     "• Next.js 14 project initialization with TypeScript and Tailwind CSS configuration.\n• Establishment of centralized data schema in `src/data/business.ts` and component skeleton."),
     
    ("Sprint 3 – UI Implementation & Interactivity:",
     "• Development of modular UI sections (`Hero`, `ServiceCard`, `About`, `Contact`, `Footer`).\n• Implementation of Framer Motion micro-animations, responsive drawer, and Google Maps embed."),
     
    ("Sprint 4 – SEO Optimization, Quality Assurance & Deployment:",
     "• Embedding Schema.org LocalBusiness JSON-LD markup, OpenGraph social preview tags.\n• Lighthouse audit, cross-browser compatibility testing, and client sign-off.")
]

for head, body in agile_steps:
    p = tf.add_paragraph()
    r1 = p.add_run()
    r1.text = f"•  {head}\n"
    r1.font.size = Pt(12.5)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_BLACK
    r1.font.name = "Arial"
    
    r2 = p.add_run()
    r2.text = f"    {body}"
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = COLOR_GRAY_DARK
    r2.font.name = "Arial"
    p.space_before = Pt(8)


# ==============================================================================
# SLIDE 13: MODELING (Developing the use cases)
# ==============================================================================
slide13 = prs.slides.add_slide(blank_layout)
set_slide_background(slide13)
add_academic_header(slide13, "Modeling", "(Developing the use cases)")

content_box = slide13.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

use_case_q = [
    ("I. Primary & Secondary actors?", "Primary: Event Organizers, DJs, Sound Rental Vendors, Household Consumers. Secondary: Chief Technician (Murlidhar Chaudhari)."),
    ("II. Actor's goal(s)?", "Identify repair capabilities for faulty gear, verify workshop trustworthiness, and initiate instant repair inquiry."),
    ("III. Preconditions required?", "Active internet connection and modern web browser (mobile/desktop)."),
    ("IV. Main tasks performed?", "Browse repair categories, review diagnostic highlights, trigger direct phone calls, initiate WhatsApp messages, and locate workshop via map."),
    ("V. Exceptions considered?", "Network disconnect (served via service worker/cached assets); device lacks native dialer (fallback copy-to-clipboard)."),
    ("VI. Interaction variations?", "Direct mobile call vs. desktop WhatsApp Web redirect; list view vs. category-filtered view."),
    ("VII. System information handled?", "Service details, workshop GPS coordinates, business operating hours (Mon-Sat 9:30 AM - 9:00 PM)."),
    ("VIII. External environment updates?", "Dynamic business telephone numbers or new equipment specializations updated via centralized data layer."),
    ("IX. Information desired by actor?", "Estimated repair turnaround time, component-level repair feasibility, workshop physical location."),
    ("X. Informed of unexpected changes?", "Yes; critical workshop holiday schedules or emergency on-call notices displayed in the status banner.")
]

for idx, (q, a) in enumerate(use_case_q):
    p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
    r_q = p.add_run()
    r_q.text = f"{q} "
    r_q.font.size = Pt(10.5)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_BLACK
    r_q.font.name = "Arial"
    
    r_a = p.add_run()
    r_a.text = a
    r_a.font.size = Pt(10)
    r_a.font.color.rgb = COLOR_GRAY_DARK
    r_a.font.name = "Arial"
    p.space_before = Pt(4)


# ==============================================================================
# SLIDE 14: MODELING (Use cases Table)
# ==============================================================================
slide14 = prs.slides.add_slide(blank_layout)
set_slide_background(slide14)
add_academic_header(slide14, "Modeling", "(Use cases)")

def draw_use_case_table(slide, left, top, width, uc_name, actor, goal, precond, trigger):
    rows = 5
    cols = 2
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, Inches(2.3))
    table = table_shape.table
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = width - Inches(2.2)
    
    # Header
    cell_00 = table.cell(0, 0)
    cell_00.fill.solid()
    cell_00.fill.fore_color.rgb = COLOR_TABLE_HEADER
    cell_00.text = "<use_case_name>"
    p = cell_00.text_frame.paragraphs[0]
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_BLACK
    
    cell_01 = table.cell(0, 1)
    cell_01.fill.solid()
    cell_01.fill.fore_color.rgb = COLOR_TABLE_HEADER
    cell_01.text = uc_name
    p = cell_01.text_frame.paragraphs[0]
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_BLACK
    
    data = [
        ("Primary Actor(s)", actor),
        ("Goal of the use case", goal),
        ("Precondition(s)", precond),
        ("Trigger", trigger)
    ]
    
    for r_idx, (label, val) in enumerate(data, start=1):
        c0 = table.cell(r_idx, 0)
        c0.fill.solid()
        c0.fill.fore_color.rgb = COLOR_WHITE
        c0.text = label
        p0 = c0.text_frame.paragraphs[0]
        p0.font.bold = True
        p0.font.size = Pt(10)
        p0.font.color.rgb = COLOR_BLACK
        
        c1 = table.cell(r_idx, 1)
        c1.fill.solid()
        c1.fill.fore_color.rgb = COLOR_WHITE
        c1.text = val
        p1 = c1.text_frame.paragraphs[0]
        p1.font.size = Pt(10)
        p1.font.color.rgb = COLOR_GRAY_DARK

# Table 1: UC-1
draw_use_case_table(
    slide14, Inches(1.0), Inches(1.8), Inches(11.3),
    "UC-1: Equipment Repair Consultation & Booking",
    "Customer (DJ, Event Organizer, Sound Vendor)",
    "Inquire about diagnostic feasibility, repair estimation, and drop-off scheduling.",
    "User has navigated to the Vikash Electronics portal on mobile or desktop.",
    "User clicks on 'Call Now' (+91 98254 85520) or 'Inquire on WhatsApp' on a specific equipment service."
)

# Table 2: UC-2
draw_use_case_table(
    slide14, Inches(1.0), Inches(4.5), Inches(11.3),
    "UC-2: Workshop Geolocation & Physical Drop-off",
    "Customer (Equipment Owner)",
    "Locate the exact workshop address in Limbayat, Surat and obtain GPS turn-by-turn directions.",
    "Customer has identified the need for physical equipment drop-off.",
    "User scrolls to the Contact section and clicks 'Get Directions' on the interactive Google Map."
)


# ==============================================================================
# SLIDE 15: OTHERS (Snapshots & Tech Stack)
# ==============================================================================
slide15 = prs.slides.add_slide(blank_layout)
set_slide_background(slide15)
add_academic_header(slide15, "Others", "(Development Snapshots & Technology Stack)")

# Left Column: Tech Stack & Key Features
content_box = slide15.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.8), Inches(5.2))
tf = content_box.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
r = p.add_run()
r.text = "Technology Stack Implemented:"
r.font.size = Pt(14)
r.font.bold = True
r.font.color.rgb = COLOR_BLACK
r.font.name = "Arial"

tech_items = [
    "Next.js 14 (App Router, Server-Side Rendering & SEO optimization)",
    "TypeScript (Robust static typing and structured business data contracts)",
    "Tailwind CSS (Utility-first responsive design tokens & custom dark theme)",
    "Framer Motion (Hardware-accelerated UI micro-animations and transitions)",
    "Lucide React (Modern, accessible SVG iconography)",
    "Schema.org JSON-LD (Rich snippet LocalBusiness search integration)"
]

for t in tech_items:
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = f"•  {t}"
    r.font.size = Pt(11)
    r.font.color.rgb = COLOR_GRAY_DARK
    r.font.name = "Arial"
    p.space_before = Pt(4)

p_fut = tf.add_paragraph()
r_fut = p_fut.add_run()
r_fut.text = "\nFuture Scope & Enhancements:"
r_fut.font.size = Pt(14)
r_fut.font.bold = True
r_fut.font.color.rgb = COLOR_BLACK
r_fut.font.name = "Arial"

future_items = [
    "Online Repair Ticket Tracking system with automated WhatsApp status alerts.",
    "Customer portal for repair estimates, quotation approvals, and digital invoices.",
    "E-commerce spare parts inventory for pro-audio cables, connectors, and discharge lamps."
]

for f in future_items:
    p = tf.add_paragraph()
    r = p.add_run()
    r.text = f"•  {f}"
    r.font.size = Pt(10.5)
    r.font.color.rgb = COLOR_GRAY_DARK
    r.font.name = "Arial"
    p.space_before = Pt(3)

# Right Column: Visual Artifacts / Workshop & Service Images
workshop_img = os.path.join(BASE_DIR, "public", "images", "about", "about_workshop.jpg")
if os.path.exists(workshop_img):
    slide15.shapes.add_picture(workshop_img, Inches(7.0), Inches(1.8), width=Inches(5.5), height=Inches(2.5))

# Second row with 2 service snapshots
spk_img = os.path.join(BASE_DIR, "public", "images", "services", "speaker.jpg")
amp_img = os.path.join(BASE_DIR, "public", "images", "services", "amplifier.jpg")

if os.path.exists(spk_img):
    slide15.shapes.add_picture(spk_img, Inches(7.0), Inches(4.6), width=Inches(2.65), height=Inches(1.8))

if os.path.exists(amp_img):
    slide15.shapes.add_picture(amp_img, Inches(9.85), Inches(4.6), width=Inches(2.65), height=Inches(1.8))

# Save presentation
prs.save(OUTPUT_FILE)
print(f"Academic presentation successfully created at: {OUTPUT_FILE}")
